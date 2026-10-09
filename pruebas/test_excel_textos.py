from copy import deepcopy
import json

from openpyxl import load_workbook
import pytest

from aprendizaje import GestorAprendizaje
from configuracion import cargar_config
from generar_excel import generar_excel


@pytest.fixture
def excel_legacy_sintetico(monkeypatch, tmp_path):
    import pruebas_externas

    config = cargar_config()
    config["aprendizaje"] = {"directorio": str(tmp_path / "aprendizaje")}
    config["colores"] = {"verde": "#22c55e", "amarillo": "#f59e0b", "rojo": "#ef4444", "gris": "#9ca3af"}
    filas = []
    nombres = ["01_A1_variante2", "02_A1_variante3", "03_B7_variante1",
               "04_B7_variante2", "05_C2_variante1", "06_D4_variante1"]
    for i, nombre in enumerate(nombres):
        carpeta = tmp_path / "lote" / nombre
        carpeta.mkdir(parents=True)
        imagen = carpeta / "etiqueta.png"
        imagen.write_bytes(f"contenido sintetico {i}".encode())
        lineas = [{"texto": f"ETQ-{i}-{j}", "bbox": [1, 2 + j * 12, 40, 10], "confianza": .9}
                  for j in range(2)]
        filas.append({"nombre": nombre, "ruta": str(carpeta), "perfil": "legacy",
            "identificador": nombre, "confianza_ocr_pct": 90,
            "comparacion": {"resultado": "coincidencia_total"},
            "imagenes": [{"nombre": imagen.name, "ruta": str(imagen), "rol": "etiqueta",
                          "resultado_ocr": {"lineas_texto": lineas}}]})
    documento = {"perfil": "legacy", "raiz": str(tmp_path / "lote"), "resultados": filas}
    ruta_validacion = tmp_path / "validacion.json"
    ruta_validacion.write_text(json.dumps(documento), encoding="utf-8")
    externas = []
    for i in range(8):
        imagen = tmp_path / f"externa-{i}.png"
        imagen.write_bytes(f"externa sintetica {i}".encode())
        externas.append({"id": f"externa-{i}", "imagen": imagen.name, "ruta": str(imagen),
            "tipo": "codigo", "coincidencia_exacta": True,
            "lineas_texto": [{"texto": f"EXT-{i}", "bbox": [1, 2, 40, 10], "confianza": .9}]})
    monkeypatch.setattr(pruebas_externas, "cargar", lambda _config: {"resultados": deepcopy(externas)})
    return ruta_validacion, config


def test_excel_agrupa_textos_por_carpeta_imagen_y_region(tmp_path, excel_legacy_sintetico):
    ruta_validacion, config = excel_legacy_sintetico
    salida = generar_excel(ruta_validacion=ruta_validacion, config=config,
                           ruta_salida=tmp_path / "resultado.xlsx")
    libro = load_workbook(salida, read_only=False)

    assert "Textos por imagen" in libro.sheetnames
    hoja = libro["Textos por imagen"]
    encabezados = [celda.value for celda in hoja[1]]
    assert encabezados[:8] == [
        "Carpeta", "Ruta carpeta", "Imagen", "Rol", "Origen", "Revisión", "Orden", "Texto"]
    filas = list(hoja.iter_rows(min_row=2, values_only=True))
    assert len(filas) >= 13
    assert {fila[0] for fila in filas} >= {
        "01_A1_variante2", "02_A1_variante3", "03_B7_variante1"}
    assert all(fila[2] and fila[4] in {
        "OCR", "Manual", "OCR externo", "Manual externo"} for fila in filas)
    assert "Bandeja de revisión" in libro.sheetnames
    bandeja = libro["Bandeja de revisión"]
    assert bandeja.max_row >= 15  # encabezado + 6 carpetas + 8 pruebas externas


def test_excel_incluye_texto_manual_de_prueba_externa(monkeypatch, tmp_path, excel_legacy_sintetico):
    import pruebas_externas

    imagen = tmp_path / "externa.jpg"
    imagen.write_bytes(b"imagen-externa-controlada")
    ruta_validacion, config = excel_legacy_sintetico
    config["aprendizaje"] = {
        **config.get("aprendizaje", {}),
        "directorio": str(tmp_path / "aprendizaje"),
    }
    GestorAprendizaje(config).registrar_region(
        "SERIE EXTERNA 42", [4, 5, 60, 18], ruta_imagen=str(imagen),
        carpeta_id="externa-controlada", carpeta_nombre="Pruebas complejas",
        imagen_nombre=imagen.name, fuente="prueba")
    monkeypatch.setattr(pruebas_externas, "cargar", lambda _config: {"resultados": [{
        "id": "externa-controlada", "imagen": imagen.name, "ruta": str(imagen),
        "tipo": "codigo", "coincidencia_exacta": False,
        "lineas_texto": [{"texto": "SERIE EXTERNA", "bbox": [1, 2, 40, 10],
                           "confianza": 0.8}],
    }]})

    salida = generar_excel(ruta_validacion=ruta_validacion,
                           ruta_salida=tmp_path / "manual-externo.xlsx", config=config)
    hoja = load_workbook(salida)["Textos por imagen"]
    filas = list(hoja.iter_rows(min_row=2, values_only=True))

    manual = next(fila for fila in filas if fila[4] == "Manual externo")
    assert manual[2] == imagen.name
    assert manual[7] == "SERIE EXTERNA 42"
    assert manual[9:13] == (4, 5, 60, 18)
