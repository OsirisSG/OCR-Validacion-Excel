"""V15.9.2 (C7): identidad autoritativa del proyecto en los Excel de SALIDA.

Reglas que implementa este módulo:

- La plantilla original es INMUTABLE: sólo se escribe sobre la copia de
  salida ya cargada en memoria por el generador.
- El proyecto se escribe únicamente en CELDAS DE METADATOS IDENTIFICADAS:
  un bloque propio en la hoja de resumen y, si la plantilla declara una
  celda de proyecto mediante etiqueta de cabecera («Proyecto:», «Project:»,
  «Projekt:»), esa celda concreta. NUNCA se hace buscar-y-reemplazar sobre
  todo el libro (podría dañar fórmulas, etiquetas o evidencia histórica).
- Si la plantilla externa no tiene campo de proyecto identificable, se
  escribe una ADVERTENCIA verificable en el bloque de identidad y se marca
  ``identidad_verificada=False``: el reporte identificado queda bloqueado
  (no se suplanta el título antiguo).
"""
from __future__ import annotations

import re
import unicodedata
from datetime import datetime

ETIQUETAS_PROYECTO = re.compile(
    r"^(?:proyecto|project|projekt|pr[oö]jekt)\s*[:=]?$", re.IGNORECASE)

_TRADUCCIONES = {
    "es": {"bloque": "Identidad del proyecto (autoritativa)",
           "nombre": "nombre_proyecto", "id": "project_id",
           "origen": "Origen de la identidad", "estado": "Estado",
           "ejecucion": "ejecucion_id", "fecha": "Fecha de ejecución",
           "responsable": "Responsable", "sin_asignar": "Sin asignar (pendiente)",
           "pendiente": "PENDIENTE de confirmación del usuario",
           "confirmado": "Confirmado por el usuario en esta ejecución",
           "advertencia": ("ADVERTENCIA: esta plantilla no declara una celda de "
                           "proyecto identificable; se bloquea la publicación del "
                           "reporte identificado y NO se reemplaza ningún título "
                           "existente."),
           "borrador": ("BORRADOR NO PUBLICABLE: la identidad del proyecto no se "
                        "pudo verificar en esta plantilla; la salida oficial NO se "
                        "generó ni se reemplazó. Declara una celda de proyecto con "
                        "etiqueta («Proyecto:»/«Project:»/«Projekt:») o confirma el "
                        "proyecto y vuelve a procesar."),
           "contradiccion": ("REFERENCIA CONTRADICTORIA: {ref} muestra otro "
                             "proyecto («{texto}»). La salida oficial NO se "
                             "publica mientras persista esa referencia; "
                             "resuelve el título o su mapeo en la plantilla. "
                             "Ningún texto existente fue reemplazado."),
           "detalle": "Referencia no resoluble con certeza: {detalle}.",
           "verificacion": "Identidad del proyecto verificada",
           "decision": "Decisión del usuario confirmada",
           "afirmativo": "sí", "negativo": "no",
           "exentas": ("Hojas exentas del examen de identidad "
                       "(histórico/evidencia, por contrato)")},
    "en": {"bloque": "Authoritative project identity",
           "nombre": "nombre_proyecto", "id": "project_id",
           "origen": "Identity source", "estado": "Status",
           "ejecucion": "ejecucion_id", "fecha": "Execution date",
           "responsable": "Responsible", "sin_asignar": "Unassigned (pending)",
           "pendiente": "PENDING user confirmation",
           "confirmado": "Confirmed by the user for this run",
           "advertencia": ("WARNING: this template does not declare an "
                           "identifiable project cell; publishing the identified "
                           "report is blocked and no existing title is replaced."),
           "borrador": ("NON-PUBLISHABLE DRAFT: the project identity could not be "
                        "verified in this template; the official output was NOT "
                        "created or replaced. Declare a labelled project cell "
                        "(«Proyecto:»/«Project:»/«Projekt:») or confirm the project "
                        "and re-run."),
           "contradiccion": ("CONTRADICTORY REFERENCE: {ref} shows a different "
                             "project («{texto}»). The official output is NOT "
                             "published while that reference persists; resolve "
                             "the title or its mapping in the template. No "
                             "existing text was replaced."),
           "detalle": "Reference not resolvable with certainty: {detalle}.",
           "verificacion": "Project identity verified",
           "decision": "User decision confirmed",
           "afirmativo": "yes", "negativo": "no",
           "exentas": ("Sheets exempt from the identity examination "
                       "(historical/evidence, by contract)")},
    "de": {"bloque": "Autoritative Projektidentität",
           "nombre": "nombre_proyecto", "id": "project_id",
           "origen": "Identitätsquelle", "estado": "Status",
           "ejecucion": "ejecucion_id", "fecha": "Ausführungsdatum",
           "responsable": "Verantwortlich", "sin_asignar": "Nicht zugeordnet (ausstehend)",
           "pendiente": "AUSSTEHEND: Bestätigung des Benutzers",
           "confirmado": "Vom Benutzer für diesen Lauf bestätigt",
           "advertencia": ("WARNUNG: Diese Vorlage deklariert keine erkennbare "
                           "Projektzelle; die Veröffentlichung des identifizierten "
                           "Berichts ist gesperrt und kein vorhandener Titel wird "
                           "ersetzt."),
           "borrador": ("NICHT VERÖFFENTLICHUNGSFÄHIGER ENTWURF: Die "
                        "Projektidentität konnte in dieser Vorlage nicht verifiziert "
                        "werden; die offizielle Ausgabe wurde NICHT erstellt oder "
                        "ersetzt. Deklariere eine Projektzelle mit Beschriftung "
                        "(«Proyecto:»/«Project:»/«Projekt:») oder bestätige das "
                        "Projekt und starte erneut."),
           "contradiccion": ("WIDERSPRÜCHLICHER VERWEIS: {ref} zeigt ein "
                             "anderes Projekt («{texto}»). Die offizielle "
                             "Ausgabe wird NICHT veröffentlicht, solange dieser "
                             "Verweis besteht; kläre den Titel oder sein Mapping "
                             "in der Vorlage. Kein vorhandener Text wurde "
                             "ersetzt."),
           "detalle": "Verweis nicht sicher auflösbar: {detalle}.",
           "verificacion": "Projektidentität verifiziert",
           "decision": "Benutzerentscheidung bestätigt",
           "afirmativo": "ja", "negativo": "nein",
           "exentas": ("Vom Identitäts-Prüfung ausgenommene Blätter "
                       "(Historie/Evidenz, per Vertrag)")},
}


def textos_identidad(idioma: str = "es") -> dict:
    return _TRADUCCIONES.get((idioma or "es").lower(), _TRADUCCIONES["es"])


# V15.9.4 (CRIT-N1/N2): contrato de mapeo de títulos/encabezados de proyecto.
# Además de la celda etiquetada («Proyecto:»/«Project:»/«Projekt:»), se
# verifican las referencias visibles a proyecto en las hojas oficiales
# (resumen + hojas de búsqueda): una referencia a OTRO proyecto que la
# identidad declarada es una contradicción no resoluble con certeza ⇒ la
# salida no es publicable. NUNCA se hace reemplazo global ciego (podría dañar
# fórmulas o evidencia histórica); el título contradictorio permanece intacto
# y se lista en el bloque de identidad.
MARCADOR_BORRADOR = ".borrador_no_publicable"

_REFERENCIA_PROYECTO = re.compile(r"\b(?:proyecto|project|projekt|pr[oö]jekt)\b",
                                  re.IGNORECASE)
_TOKENS_INNOCUOS = {"del", "de", "la", "el", "los", "las", "y", "o", "en", "con",
                    "sin", "actual", "declarado", "identificable", "seleccionado",
                    "elegido", "este", "esta", "para", "por", "su", "un", "una"}

# V15.9.5 (MEDIA-ALTA-N8) / V15.9.6 (ALTA-N11, MEDIA-ALTA-N12): contrato de
# análisis del texto que sigue a una mención de proyecto. El texto de un
# título puede combinar el NOMBRE del proyecto con METADATOS de la prueba
# (año, semana, fecha, versión, etapa/fase…). Los metadatos se separan por
# segmentos (delimitadores) y por patrones; el NOMBRE se compara como unidad
# COMPLETA normalizada contra las representaciones de la identidad autoritativa
# (nombre_proyecto, project_id, nombre visible y alias registrados
# explícitamente). Una palabra compartida NO verifica un nombre compuesto:
# la coincidencia parcial se trata fail-closed. Una referencia SIN ningún
# nombre resoluble (p. ej. sólo un número) también es ambigua ⇒ fail-closed.
_DELIMITADORES_SEGMENTO = re.compile(
    r"\s*(?:\u2014|\u2013|\||\u00b7)\s*|\s+-\s+|\s*,\s*|[()]")


def _plegar(texto) -> str:
    """Normalización trivial (N11): mayúsculas, diacríticos, espaciado.

    ``casefold`` + NFD sin marcas combinantes + guiones/barras como espacio +
    espacios colapsados. No altera el contenido comparable del nombre.
    """
    plegado = unicodedata.normalize("NFD", str(texto or "").casefold())
    plegado = "".join(c for c in plegado if not unicodedata.combining(c))
    return re.sub(r"[\s\-_/]+", " ", plegado).strip()


def _palabras(texto) -> list[str]:
    return re.findall(r"[0-9a-z]+", _plegar(texto))


_TOKENES_METADATOS = {_plegar(t) for t in {
    # año / semana / mes / trimestre (ES/EN/DE)
    "año", "anio", "años", "year", "years", "jahr", "jahre",
    "semana", "semanas", "week", "weeks", "woche", "wochen", "kw", "cw",
    "mes", "meses", "month", "months", "monat", "monate",
    "trimestre", "trimestres", "quarter", "quarters", "quartal", "quartale",
    # versión / revisión / fecha
    "versión", "version", "versiones", "versions",
    "revisión", "revision", "revisiones", "revisions", "rev",
    "fecha", "fechas", "date", "dates", "datum",
    # nombres de mes (ES/EN/DE)
    "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
    "septiembre", "octubre", "noviembre", "diciembre",
    "january", "february", "march", "april", "may", "june", "july", "august",
    "september", "october", "november", "december",
    "januar", "jänner", "maerz", "mai", "juli", "oktober", "dezember",
}}

# V15.9.6 (MEDIA-ALTA-N12): reconocimiento CONTEXTUAL de etapas/fases del
# flujo V-Matrix. Los valores (1ST, 2ST, 3RD, NACH, VOR, LEGACY) sólo son
# metadato cuando el segmento incluye su clave (Etapa/Stage/Fase/Phase/
# Stufe…); NUNCA se aceptan de forma global, porque podrían formar parte
# legítima del nombre de OTRO proyecto.
_CLAVES_ETAPA_FASE = {_plegar(t) for t in {
    "etapa", "etapas", "stage", "stages", "fase", "fases",
    "phase", "phasen", "stufe", "stufen"}}
_VALORES_ETAPA_FASE = {"1st", "2st", "3rd", "nach", "vor", "legacy"}

# V15.9.9 (CRIT-N16): vocabulario de CABECERAS DE COLUMNA habituales en
# tablas de ensayos/reportes (ES/EN/DE). Un valor contiguo a una etiqueta
# desnuda de proyecto que pertenece a este vocabulario es una cabecera de
# tabla —aunque la tabla esté VACÍA (sin filas de datos)— y nunca se
# sobrescribe. Un proyecto cuyo nombre colisione de verdad con una
# cabecera («Etapa») exige asignación explícita (etiqueta con separador o
# contrato de plantilla); jamás se sobrescribe automáticamente.
_CABECERAS_COLUMNA = {_plegar(t) for t in {
    # etapa / fase / caso (ES/EN/DE)
    "etapa", "etapas", "stage", "stages", "fase", "fases",
    "phase", "phasen", "stufe", "stufen",
    "caso", "casos", "case", "cases", "fall", "fälle", "faelle",
    "prueba", "pruebas", "test", "tests", "versuch", "versuche",
    "ensayo", "ensayos", "experiment", "experiments",
    # identificadores y resultados de columna
    "id", "ids", "n", "nr", "num", "numero", "número", "number",
    "valor", "valores", "value", "values", "wert", "werte",
    "resultado", "resultados", "result", "results",
    "ergebnis", "ergebnisse",
    "estado", "status", "zustand",
    "criterio", "criterios", "criterion", "criteria",
    "kriterium", "kriterien",
    "fecha", "date", "datum",
    "camara", "cámara", "camera", "cameras", "kamera", "kameras",
    "observacion", "observación", "observaciones",
    "nota", "notas", "note", "notes", "bemerkung", "bemerkungen",
    "comentario", "comentarios", "comment", "comments",
    "kommentar", "kommentare",
}}

# Palabras clave de mención (plegadas): se retiran del inicio del nombre al
# comparar, de modo que «Project B» y «Proyecto B» tienen el mismo núcleo.
_CLAVES_MENCION = {_plegar(t) for t in {"proyecto", "project", "projekt",
                                        "pröjekt", "projeto"}}

_PATRONES_METADATO = (
    re.compile(r"^(?:19|20)\d{2}$"),                                # año
    re.compile(r"^\d{1,2}[./]\d{1,2}(?:[./](?:\d{2}|\d{4}))?$"),    # fecha
    re.compile(r"^v?\d+(?:[.,]\d+)*$"),                             # versión/número
)


def _tokens_identidad(texto) -> set[str]:
    return set(_palabras(texto))


def _token_es_metadato(token: str) -> bool:
    return (token in _TOKENES_METADATOS
            or any(patron.match(token) for patron in _PATRONES_METADATO))


def _segmento_es_metadato(segmento: str) -> bool:
    """¿El segmento completo es metadato (año/semana/fecha/versión/etapa)?

    - Sólo palabras funcionales ⇒ no aporta identidad.
    - Sólo patrones (fecha/versión/número/año) ⇒ metadato numérico.
    - Con clave de metadato (incluidas claves de etapa/fase, N12) y el resto
      de tokens numéricos o valores de etapa/fase ⇒ metadato contextual.
    Un valor de etapa suelto («1ST») SIN su clave NO es metadato: podría ser
    el nombre de otro proyecto (fail-closed).
    """
    tokens = [t for t in _palabras(segmento) if t not in _TOKENS_INNOCUOS]
    if not tokens:
        return True
    if all(any(patron.match(t) for patron in _PATRONES_METADATO)
           for t in tokens):
        return True
    claves = set(tokens) & (_TOKENES_METADATOS | _CLAVES_ETAPA_FASE)
    if not claves:
        return False
    resto = set(tokens) - claves
    return all(_token_es_metadato(t) or t in _VALORES_ETAPA_FASE
               for t in resto)


def _nucleo_nombre(frase):
    """Núcleo comparable de un nombre: sin funcionales ni clave de mención."""
    toks = [t for t in _palabras(frase) if t not in _TOKENS_INNOCUOS]
    while toks and toks[0] in _CLAVES_MENCION:
        toks = toks[1:]
    return tuple(toks)


def _representaciones_identidad(identidad: dict | None, idioma: str = "es"):
    """Representaciones COMPLETAS autorizadas de la identidad (N11).

    ``nombre_proyecto`` y ``project_id`` se tratan como unidades completas
    (no como conjuntos de palabras fusionados); se añade el nombre visible y
    los alias sólo si están registrados explícitamente en la identidad.
    """
    identidad = identidad or {}
    fuentes = [identidad.get("nombre_proyecto"), identidad.get("project_id"),
               nombre_visible(identidad, idioma)]
    alias = identidad.get("aliases") or identidad.get("alias") or []
    if isinstance(alias, (list, tuple, set)):
        fuentes.extend(alias)
    representaciones = set()
    for fuente in fuentes:
        if isinstance(fuente, str) and fuente.strip():
            nucleo = _nucleo_nombre(fuente)
            if nucleo:
                representaciones.add(nucleo)
    return representaciones


def _analizar_referencia(candidato: str, mencion: str, representaciones):
    """N8/N11/N12: nombre del proyecto (completo) vs metadatos del título.

    Construye la frase de nombre con la palabra de mención + los segmentos no
    metadatos y la compara como UNIDAD COMPLETA normalizada contra las
    representaciones autorizadas. Se aceptan dos niveles:

    - estricto: la frase tal cual aparece;
    - laxo: sin sufijos numéricos/versión sueltos dentro del segmento de
      nombre («B 2026» ≡ «B»).

    Coincidencia parcial de un nombre compuesto ⇒ contradicción. Sin nombre
    resoluble ⇒ ambigüedad fail-closed. Devuelve ``None`` si se resuelve o el
    ``detalle`` de lo que no se pudo verificar.
    """
    if not [t for t in _palabras(candidato) if t not in _TOKENS_INNOCUOS]:
        return None  # sólo palabras funcionales: no es una referencia de identidad
    segmentos = _DELIMITADORES_SEGMENTO.split(candidato)
    estricta = [mencion]
    laxa = [mencion]
    hay_nombre = False
    for segmento in segmentos:
        segmento = (segmento or "").strip(" ,;:\u2013\u2014")
        if not segmento:
            continue
        funcion = [t for t in _palabras(segmento) if t not in _TOKENS_INNOCUOS]
        if not funcion or _segmento_es_metadato(segmento):
            continue
        hay_nombre = True
        estricta.append(segmento)
        residual = [t for t in funcion
                    if not any(p.match(t) for p in _PATRONES_METADATO)]
        if residual:
            laxa.append(" ".join(residual))
    if not hay_nombre:
        return "no contiene ningún nombre de proyecto resoluble"
    nucleo = _nucleo_nombre(" ".join(estricta))
    if nucleo in representaciones:
        return None
    nucleo_laxo = _nucleo_nombre(" ".join(laxa))
    if nucleo_laxo in representaciones:
        return None
    leido = " ".join(nucleo_laxo or nucleo)
    return (f"el nombre leído «{leido}» no coincide con el nombre completo "
            "del proyecto declarado ni con su ID (la coincidencia parcial no "
            "verifica la identidad)")


def es_borrador_no_publicable(ruta) -> bool:
    """¿La ruta corresponde a un borrador bloqueado por identidad (N6)?"""
    return MARCADOR_BORRADOR in str(getattr(ruta, "name", ruta) or "")


class PublicacionExcelBloqueadaError(RuntimeError):
    """V15.9.4 (ALTA-N6): publicación oficial bloqueada por identidad.

    La usan los flujos de reexportación (cola durable, endpoints) para NO
    declarar éxito cuando el generador devolvió un borrador no publicable.
    """


def _textos_cabecera_pie(hoja):
    """V15.9.5 (CRIT-N7): textos de cabecera/pie de página de una hoja.

    Algunas plantillas declaran el proyecto en la cabecera o el pie
    imprimible; si el texto existe, forma parte del entregable visible.
    """
    for nombre in ("oddHeader", "evenHeader", "firstHeader",
                   "oddFooter", "evenFooter", "firstFooter"):
        bloque = getattr(hoja, nombre, None)
        if bloque is None:
            continue
        for posicion in ("left", "center", "right"):
            parte = getattr(bloque, posicion, None)
            texto = getattr(parte, "text", None)
            if isinstance(texto, str) and texto.strip():
                yield f"{nombre}.{posicion}", texto.strip()


def hojas_del_entregable(libro, hojas_exentas=None) -> list[str]:
    """V15.9.5 (CRIT-N7): contrato de hojas del entregable.

    TODAS las hojas VISIBLES del libro de salida forman parte del reporte
    exportado y deben examinarse, aunque el generador no las declare en
    ``hojas_busqueda``. Excepción explícita y documentada: hojas de
    histórico/ejemplo/evidencia declaradas por el generador en
    ``hojas_exentas``. Las hojas ocultas no son entregable visible (se
    conservan intactas en el archivo).
    """
    exentas = {str(nombre) for nombre in (hojas_exentas or ())}
    visibles = [hoja.title for hoja in libro.worksheets
                if hoja.sheet_state == "visible"]
    return [nombre for nombre in visibles if nombre not in exentas]


def detectar_referencias_contradictorias(libro, hojas: list[str],
                                         identidad: dict | None,
                                         idioma: str = "es",
                                         incluir_cabeceras: bool = True) -> list[dict]:
    """Referencias visibles a un proyecto DISTINTO del declarado.

    V15.9.5 (CRIT-N7): el llamador decide el contrato de hojas; por regla es
    ``hojas_del_entregable`` (todas las visibles), no sólo las mapeadas. Se
    extrae el texto que sigue a cada mención de «proyecto/project/projekt» en
    celdas y, si la plantilla los usa, en cabeceras/pies de página. Una celda
    que ES sólo la etiqueta (celda de metadatos mapeable) se omite.

    V15.9.5 (MEDIA-ALTA-N8): el texto tras la mención se analiza por
    segmentos: los metadatos de la prueba (año, semana, fecha, versión,
    etapa) se separan del nombre. Una referencia sin ningún nombre resoluble
    es ambigua ⇒ fail-closed con explicación (``detalle``).

    V15.9.6 (ALTA-N11): el nombre se compara como UNIDAD COMPLETA normalizada
    contra las representaciones autorizadas (``nombre_proyecto``,
    ``project_id``, nombre visible y alias registrados explícitamente); una
    coincidencia parcial NO verifica un nombre compuesto. (MEDIA-ALTA-N12):
    los valores de etapa/fase del flujo (1ST, 2ST, 3RD, NACH, VOR) sólo se
    aceptan como metadato junto a su clave contextual (Etapa/Stage/Fase/
    Phase/Stufe); sueltos pueden ser parte del nombre de otro proyecto.

    V15.9.6 (CRIT-N10): este análisis de textos LIBRES complementa, no
    sustituye, la verificación estructural de pares etiqueta/valor
    (``mapear_pares_etiqueta_valor``): una etiqueta «Proyecto:» cuyo valor
    está en la celda contigua no contiene texto tras la mención y por tanto
    es invisible para este escáner.
    """
    identidad = identidad or {}
    representaciones = _representaciones_identidad(identidad, idioma)
    contradictorias: list[dict] = []
    vistas: set[tuple[str, str]] = set()

    def _registrar(nombre_hoja: str, referencia: str, texto: str, detalle: str):
        clave = (nombre_hoja, referencia)
        if clave in vistas:
            return
        vistas.add(clave)
        contradictorias.append({"hoja": nombre_hoja, "celda": referencia,
                                "texto": texto[:180], "detalle": detalle})

    def _examinar_texto(nombre_hoja: str, referencia: str, texto: str):
        if ETIQUETAS_PROYECTO.match(texto):
            return  # etiqueta pura de celda de metadatos
        for mencion in _REFERENCIA_PROYECTO.finditer(texto):
            resto = texto[mencion.end():].split("\n")[0]
            resto = resto.lstrip(":=").strip()
            if not resto:
                continue
            candidato = re.split(r"[.;:!?]\s", resto)[0].strip(" ,;-\u2013\u2014()")
            if not candidato:
                continue
            mencion_texto = texto[mencion.start():mencion.end()]
            detalle = _analizar_referencia(candidato, mencion_texto,
                                           representaciones)
            if detalle:
                _registrar(nombre_hoja, referencia, texto, detalle)
                return

    for nombre_hoja in hojas:
        if nombre_hoja not in libro.sheetnames:
            continue
        hoja = libro[nombre_hoja]
        for fila in hoja.iter_rows():
            for celda in fila:
                valor = celda.value
                if not isinstance(valor, str) or not valor.strip():
                    continue
                _examinar_texto(nombre_hoja, celda.coordinate, valor.strip())
        if incluir_cabeceras:
            for referencia, texto in _textos_cabecera_pie(hoja):
                _examinar_texto(nombre_hoja, referencia, texto)
    return contradictorias


def nombre_visible(identidad: dict | None, idioma: str = "es") -> str:
    """Texto visible coherente en los tres idiomas (P07)."""
    t = textos_identidad(idioma)
    identidad = identidad or {}
    if identidad.get("sin_asignar") or not identidad.get("nombre_proyecto"):
        return t["sin_asignar"]
    return str(identidad["nombre_proyecto"])


def buscar_celdas_proyecto(libro, hojas: list[str], fila_max: int = 4) -> list[dict]:
    """Localiza CANDIDATOS de proyecto por etiqueta de cabecera.

    Búsqueda acotada (primeras ``fila_max`` filas de las hojas indicadas) a
    celdas cuyo texto sea una etiqueta de proyecto; la celda de valor es la
    contigua a la derecha. No es un buscar-y-reemplazar global. V15.9.8
    (N15): encontrar la etiqueta NO autoriza la escritura — cada candidato
    se clasifica después con ``_clasificar_par`` (campo único, cabecera
    tabular, estructura ambigua, destino no editable) antes de modificar
    ninguna celda.
    """
    hallazgos = []
    for nombre_hoja in hojas:
        if nombre_hoja not in libro.sheetnames:
            continue
        hoja = libro[nombre_hoja]
        for fila in hoja.iter_rows(min_row=1, max_row=min(fila_max, hoja.max_row or 1)):
            for celda in fila:
                if isinstance(celda.value, str) and ETIQUETAS_PROYECTO.match(celda.value.strip()):
                    valor = hoja.cell(celda.row, celda.column + 1)
                    hallazgos.append({"hoja": nombre_hoja,
                                      "etiqueta": str(celda.value).strip(),
                                      "celda_etiqueta": celda.coordinate,
                                      "celda_valor": valor.coordinate,
                                      "valor_anterior": valor.value})
    return hallazgos


def _celda_valor_editable(celda) -> bool:
    """¿Puede sustituirse el valor de esta celda sin riesgo (N10/N13)?

    Editable: texto plano o vacía (campo de proyecto sin rellenar). No
    editable: fórmula (no se borran fórmulas), celda combinada (ni la inicial
    ni una ``MergedCell`` de sólo lectura: escribir ahí lanza
    ``AttributeError`` o rompe el rango) u otro tipo de valor (números,
    fechas, booleanos: estructura ambigua ⇒ bloqueo en vez de sustitución).
    """
    from openpyxl.cell.cell import MergedCell
    if isinstance(celda, MergedCell) or celda.data_type == "f":
        return False
    for rango in getattr(celda.parent, "merged_cells", None) or ():
        if celda.coordinate in rango:
            return False
    return celda.value is None or isinstance(celda.value, str)


def _motivo_no_editable(celda):
    """Motivo legible por el que una celda de valor no es editable (N13)."""
    from openpyxl.cell.cell import MergedCell
    if isinstance(celda, MergedCell):
        return "celda combinada de sólo lectura"
    for rango in getattr(celda.parent, "merged_cells", None) or ():
        if celda.coordinate in rango:
            return f"celda combinada ({rango.coord})"
    if celda.data_type == "f":
        return "fórmula (no se borra)"
    return f"valor no textual ({type(celda.value).__name__})"


def _registro_campo_no_editable(hoja_nombre: str, celda, etiqueta: str,
                                clase: str) -> dict:
    """Contradicción estructural uniforme para campos no editables (N13).

    La misma política rige el mapeo DECLARADO (``hojas_busqueda``) y el
    COMPLEMENTARIO (resto de hojas elegibles): nada se sobrescribe y el
    bloqueo se registra con hoja/celda/motivo.
    """
    motivo = _motivo_no_editable(celda)
    return {"hoja": hoja_nombre, "celda": celda.coordinate,
            "texto": f"{etiqueta} → {celda.value!r}",
            "detalle": (f"campo de proyecto {clase} no editable: {motivo}; "
                        "el valor quedó intacto y la identidad no se puede "
                        "verificar en este campo"),
            "estructural": True}


# Etiqueta EXPLÍCITA de campo de proyecto (con separador «:»/«=»). El mapeo
# estructural complementario (N10) exige el separador para no confundir un
# campo etiqueta/valor con la cabecera de una tabla de datos («Proyecto» como
# columna de trazabilidad/evidencia, cuyas filas llevan valores por caso).
_ETIQUETA_EXPLICITA = re.compile(
    r"^(?:proyecto|project|projekt|pr[oö]jekt|projeto)\s*[:=]$", re.IGNORECASE)

# V15.9.7 (ALTA-N14): etiqueta DESNUDA (sin separador): «Proyecto» como único
# contenido de la celda. Sólo es campo de proyecto con evidencia estructural
# segura (campo único); si puede ser cabecera de tabla, sus columnas son
# datos de trazabilidad y NO se mapean ni reemplazan (lección V15.9.6: las
# filas de cabecera de Trazabilidad/Evidencia/BMG llevan «Proyecto»/
# «Project» desnudo con ≥3 celdas). Cubre ES/EN/DE/PT como la explícita.
_ETIQUETA_DESNUDA = re.compile(
    r"^(?:proyecto|project|projekt|pr[oö]jekt|projeto)$", re.IGNORECASE)

# V15.9.10 (CRIT-N17): mención de proyecto AL INICIO del texto. Un valor
# contiguo que empieza por «Proyecto/Project/Projekt/…» («Proyecto A»,
# «Project North») es evidencia positiva de que la celda contiene un nombre
# de proyecto (campo histórico/migrado); una palabra desconocida
# («Temperatura», «Color», «A») no lo es.
_MENCION_INICIAL = re.compile(
    r"^(?:proyecto|project|projekt|pr[oö]jekt|projeto)\b", re.IGNORECASE)


def _celda_en_cabecera_estructural(hoja, fila: int, columnas) -> bool:
    """¿La celda cae en la fila de cabecera de una estructura FORMAL (N17)?

    Prueba estructural definitiva, independiente del texto, del idioma y
    del número de filas de datos: las tablas Excel registradas
    (``ws.tables`` con su ``Table.ref``) y los autofiltros
    (``ws.auto_filter.ref``) declaran explícitamente qué fila contiene
    cabeceras de columna. Ninguna celda de esa fila es un campo de proyecto
    editable, aunque su texto parezca un valor («Temperatura» junto a
    «Proyecto»). Cubre tablas vacías y pobladas, posiciones arbitrarias y
    varias tablas por hoja.
    """
    from openpyxl.utils import range_boundaries
    rangos = []
    for tabla in (getattr(hoja, "tables", None) or {}).values():
        if not getattr(tabla, "headerRowCount", 1):
            continue  # tabla registrada SIN fila de cabecera
        ref = getattr(tabla, "ref", None)
        if ref:
            rangos.append(ref)
    filtro = getattr(hoja, "auto_filter", None)
    ref_filtro = getattr(filtro, "ref", None) if filtro is not None else None
    if ref_filtro:
        rangos.append(ref_filtro)
    for ref in rangos:
        try:
            limites = range_boundaries(ref)
        except (ValueError, TypeError):
            continue
        if not limites:
            continue
        min_col, min_row, max_col, _max_row = limites
        if fila == min_row and any(min_col <= c <= max_col for c in columnas):
            return True
    return False


def _evidencia_valor_proyecto(valor, representaciones, visible) -> bool:
    """¿Hay evidencia POSITIVA de que el texto es un valor de proyecto (N17)?

    Sí cuando ya coincide con la identidad (núcleo normalizado o nombre
    visible) o cuando empieza por una palabra de mención seguida de nombre
    («Proyecto A», «Project North»): campo histórico/migrado. Una palabra
    desconocida («Temperatura», «Color», «A», «PRJ-001») NO es evidencia:
    sin tabla registrada ni contrato explícito, el par desnudo no se
    escribe por suposición (fail-closed).
    """
    if not isinstance(valor, str) or not valor.strip():
        return False
    if _nucleo_nombre(valor) in representaciones or valor.strip() == visible:
        return True
    return bool(_MENCION_INICIAL.match(valor.strip()))


def _evidencia_campo_unico(hoja, celda) -> str:
    """Evidencia estructural de una etiqueta desnuda de proyecto (N14).

    - ``"tabular"``: la fila tiene ≥3 celdas no vacías ⇒ cabecera de tabla;
      sus columnas son datos de trazabilidad, jamás se mapean.
    - ``"unico"``: la fila contiene sólo la etiqueta y, como mucho, su valor
      contiguo, SIN contenido debajo en ninguna de las dos columnas ⇒ campo
      único inequívoco: el par etiqueta/valor puede mapearse.
    - ``"ambigua"``: hay datos debajo ⇒ no se puede distinguir campo de
      cabecera de tabla de dos columnas sin verificar la coherencia de la
      columna con la identidad (fail-closed si diverge).
    """
    no_vacias = sum(1 for c in hoja[celda.row] if c.value not in (None, ""))
    if no_vacias > 2:
        return "tabular"
    for fila in range(celda.row + 1, (hoja.max_row or celda.row) + 1):
        for columna in (celda.column, celda.column + 1):
            if hoja.cell(fila, columna).value not in (None, ""):
                return "ambigua"
    return "unico"


def _es_cabecera_reconocible(texto) -> bool:
    """¿El valor contiguo es una cabecera de columna reconocible?

    V15.9.9 (CRIT-N16): todo su contenido son palabras clave de metadatos,
    de etapa/fase o del vocabulario de cabeceras de columna
    («Etapa», «Stage», «Stufe», «Caso», «Case», «Fall», «Valor», «Datum»,
    …): es cabecera de tabla —aunque la tabla esté vacía—, no un valor de
    proyecto. Clasificador centralizado: lo consultan AMBAS rutas
    (declarada y complementaria) antes de cualquier sustitución.
    """
    tokens = [t for t in _palabras(texto) if t not in _TOKENS_INNOCUOS]
    return bool(tokens) and all(
        t in _TOKENES_METADATOS or t in _CLAVES_ETAPA_FASE
        or t in _CABECERAS_COLUMNA for t in tokens)


def _coherente_con_identidad(valor, representaciones, visible) -> bool:
    """Un valor visible sólo es coherente si ES la identidad (o está vacío)."""
    if valor in (None, ""):
        return True
    if not isinstance(valor, str):
        return False
    return _nucleo_nombre(valor) in representaciones or valor.strip() == visible


def _valores_ajenos_columna(hoja, columna, fila_desde, representaciones,
                            visible) -> list[str]:
    """Coordenadas bajo la columna de la etiqueta con valores ajenos (N14)."""
    ajenos = []
    for fila in range(fila_desde, (hoja.max_row or fila_desde - 1) + 1):
        celda = hoja.cell(fila, columna)
        if celda.value in (None, ""):
            continue
        if isinstance(celda.value, str):
            texto = celda.value.strip()
            if _ETIQUETA_EXPLICITA.match(texto) or _ETIQUETA_DESNUDA.match(texto):
                # V15.9.8 (N15): una etiqueta de proyecto anidada más abajo
                # es otra cabecera de campo, no un dato de la columna; se
                # clasifica por su cuenta en su propia fila.
                continue
        if not _coherente_con_identidad(celda.value, representaciones, visible):
            ajenos.append(celda.coordinate)
    return ajenos


def _clasificar_par(hoja, nombre_hoja, celda, destino, *, explicita, clase,
                    representaciones, visible):
    """V15.9.8 (CRIT-N15): política estructural ÚNICA por par etiqueta/valor.

    Clasifica el par ANTES de tocar ninguna celda, venga la etiqueta de la
    ruta DECLARADA (``hojas_busqueda``) o de la COMPLEMENTARIA: un
    «Proyecto» desnudo puede ser campo único o cabecera de tabla y sólo la
    estructura permite distinguirlo. No basta con que el destino contenga
    texto.

    - Estructura FORMAL (V15.9.10, N17): celda en la fila de cabecera de una
      tabla Excel registrada (``ws.tables``/``Table.ref``) o de un
      autofiltro ⇒ nunca es campo de proyecto; se omite en silencio, con la
      tabla y sus nombres de columna intactos.
    - Fila tabular (≥3 celdas no vacías): etiqueta desnuda ⇒ cabecera de
      tabla (trazabilidad/BMG), no se mapea ni se bloquea y sus columnas
      quedan intactas; etiqueta explícita ⇒ estructura ambigua, sólo se
      admite si el valor YA coincide con la identidad y en caso contrario
      se registra contradicción estructural.
    - Fila de ≤2 celdas con etiqueta desnuda: evidencia estructural con
      ``_evidencia_campo_unico``; campo único ⇒ mapeable, SALVO que el
      contiguo sea una cabecera de columna reconocible (V15.9.9, N16:
      tabla vacía «Proyecto | Etapa» sin filas de datos) ⇒ se omite sin
      tocarla; ambigua (datos debajo) ⇒ verificación de coherencia del
      contiguo y de la columna bajo la etiqueta; cualquier valor ajeno ⇒
      contradicción estructural y ningún reemplazo.
    - Destino no editable (fórmula, valor no textual o celda combinada,
      N13) ⇒ contradicción estructural con hoja/celda/motivo; nada se
      escribe.

    Devuelve ``(mapear, bloqueo)``; ``mapear=True`` sólo para campo único
    (o etiqueta explícita de ≤2 celdas) con destino editable.
    """
    etiqueta = str(celda.value).strip()
    # V15.9.10 (CRIT-N17): la estructura FORMAL manda. Una celda en la fila
    # de cabecera de una tabla Excel registrada (ws.tables / Table.ref) o de
    # un autofiltro NUNCA es campo de proyecto editable, sin importar su
    # texto, idioma o número de filas de datos. Ambas rutas comparten esta
    # comprobación (función estructural única).
    if _celda_en_cabecera_estructural(hoja, celda.row,
                                      (celda.column, destino.column)):
        return False, None
    ocupadas = sum(1 for c in hoja[celda.row] if c.value not in (None, ""))
    if ocupadas > 2:
        if not explicita:
            # Cabecera tabular desnuda (N14/N15): las columnas de casos y
            # datos históricos no se tocan.
            return False, None
        actual = destino.value
        coherente = (isinstance(actual, str)
                     and (_nucleo_nombre(actual) in representaciones
                          or actual.strip() == visible))
        if coherente:
            return False, None
        return False, {
            "hoja": nombre_hoja, "celda": destino.coordinate,
            "texto": f"{etiqueta} → {actual!r}",
            "detalle": ("fila tabular con etiqueta de proyecto: la "
                        "estructura no permite resolver un campo único sin "
                        "riesgo; no se reemplaza nada"),
            "estructural": True}
    if not explicita:
        evidencia = _evidencia_campo_unico(hoja, celda)
        if evidencia == "tabular":
            return False, None
        if evidencia == "unico":
            # V15.9.9 (CRIT-N16) + V15.9.10 (CRIT-N17): dos textos contiguos
            # sin filas debajo NO son garantía de campo único — una tabla
            # vacía también cumple esa forma. Política:
            # - cabecera reconocible por vocabulario ⇒ omisión silenciosa;
            # - valor con evidencia POSITIVA de proyecto (coincide con la
            #   identidad o empieza por mención: «Proyecto A») ⇒ campo
            #   editable (migración/normalización);
            # - contiguo desconocido («Temperatura», «Color», «A») ⇒
            #   fail-closed: contradicción estructural, no se escribe por
            #   suposición (revisión o mapeo explícito requerido).
            actual = destino.value
            if isinstance(actual, str) and actual.strip():
                if _es_cabecera_reconocible(actual):
                    return False, None
                if not _evidencia_valor_proyecto(actual, representaciones,
                                                 visible):
                    return False, {
                        "hoja": nombre_hoja,
                        "celda": destino.coordinate,
                        "texto": f"{etiqueta} → {actual!r}",
                        "detalle": ("par desnudo de proyecto cuyo valor "
                                    "contiguo no tiene evidencia positiva de "
                                    "ser un nombre de proyecto (ni cabecera "
                                    "de tabla registrada, ni cabecera "
                                    "reconocible, ni contrato explícito): no "
                                    "se escribe por suposición (revisión o "
                                    "mapeo explícito requerido)"),
                        "estructural": True}
            # Valor vacío (campo por rellenar) o con evidencia positiva ⇒
            # sigue la política de sustitución segura (N13) de más abajo.
        if evidencia == "ambigua":
            # Fila de 2 celdas con datos debajo: no se puede distinguir con
            # certeza campo de cabecera de tabla de dos columnas. Se valida
            # la coherencia; ante cualquier duda, bloqueo fail-closed.
            actual = destino.value
            ajenos = []
            if not (isinstance(actual, str)
                    and _es_cabecera_reconocible(actual)) \
                    and not _coherente_con_identidad(
                        actual, representaciones, visible):
                ajenos.append(destino.coordinate)
            ajenos.extend(_valores_ajenos_columna(
                hoja, celda.column, celda.row + 1, representaciones, visible))
            if ajenos:
                return False, {
                    "hoja": nombre_hoja,
                    "celda": destino.coordinate,
                    "texto": f"{etiqueta} → {actual!r}",
                    "detalle": ("etiqueta de proyecto sin separador con "
                                "datos debajo: no se puede distinguir con "
                                "certeza un campo de una cabecera de tabla; "
                                "las celdas " + ", ".join(ajenos)
                                + " muestran valores ajenos a la identidad; "
                                "no se reemplaza nada (revisión requerida)"),
                    "estructural": True}
            return False, None
    if not _celda_valor_editable(destino):
        return False, _registro_campo_no_editable(
            nombre_hoja, destino, etiqueta, clase)
    return True, None


def mapear_pares_etiqueta_valor(libro, hojas: list[str], identidad: dict | None,
                                idioma: str = "es",
                                excluidas: set[tuple[str, str]] | None = None):
    """V15.9.6 (CRIT-N10): mapeo ESTRUCTURAL de pares etiqueta/valor.

    Complementa (no sustituye) la revisión de textos libres: recorre TODAS
    las hojas elegibles del entregable —no sólo ``hojas_busqueda``— buscando
    etiquetas EXPLÍCITAS de proyecto («Proyecto:»/«Project:»/«Projekt:», con
    separador) y su celda de valor contigua a la derecha. Política por par:

    - fila de campo único (etiqueta + valor, sin más contenido) y valor
      editable (texto plano o vacío) ⇒ se sustituye SÓLO esa celda por el
      nombre visible autoritativo y queda registrado (trazabilidad);
    - valor no resoluble sin riesgo (fórmula, celda combinada, valor no
      textual) ⇒ contradicción ESTRUCTURAL que bloquea la publicación,
      señalando hoja/celda; nada se reemplaza;
    - fila tabular con etiqueta explícita (más contenido a la derecha):
      estructura ambigua ⇒ bloquea, salvo que el valor ya coincida con la
      identidad (coherente, no se toca).

    Las cabeceras de tabla SIN separador («Proyecto» como nombre de columna,
    ≥3 celdas en la fila) no son campos de identidad y no se mapean; sus
    valores se verifican por el escáner de textos cuando contienen
    referencias visibles. V15.9.7 (ALTA-N14): una etiqueta DESNUDA («Proyecto»
    sin «:») en fila de ≤2 celdas SÍ puede mapearse cuando la estructura da
    evidencia segura de campo único (sin datos debajo); con datos debajo se
    verifica la coherencia de la columna y cualquier divergencia genera
    contradicción estructural y bloqueo (nunca publicación en silencio con un
    valor ajeno visible). Las
    excepciones de histórico/evidencia rigen igual que para el escáner: sólo
    hojas declaradas por el generador en ``hojas_exentas``. La plantilla
    original nunca se toca (se trabaja sobre la copia de salida) y no se hace
    ningún reemplazo masivo.

    Devuelve ``(corregidas, bloqueos)``.
    """
    excluidas = set(excluidas or ())
    visible = nombre_visible(identidad, idioma)
    representaciones = _representaciones_identidad(identidad, idioma)
    corregidas: list[dict] = []
    bloqueos: list[dict] = []
    for nombre_hoja in hojas:
        if nombre_hoja not in libro.sheetnames:
            continue
        hoja = libro[nombre_hoja]
        for fila in hoja.iter_rows():
            for celda in fila:
                valor = celda.value
                if not isinstance(valor, str):
                    continue
                etiqueta = valor.strip()
                explicita = bool(_ETIQUETA_EXPLICITA.match(etiqueta))
                if not explicita and not _ETIQUETA_DESNUDA.match(etiqueta):
                    continue
                if (nombre_hoja, celda.coordinate) in excluidas:
                    # V15.9.8 (N15): par ya CLASIFICADO en la ruta declarada
                    # con esta misma política centralizada; la exclusión sólo
                    # evita repetir el registro, nunca oculta una validación.
                    continue
                destino = hoja.cell(celda.row, celda.column + 1)
                # V15.9.8 (CRIT-N15): la clasificación estructural única
                # (cabecera tabular / campo único / ambigua / no editable)
                # decide antes de escribir cualquier celda.
                mapear, bloqueo = _clasificar_par(
                    hoja, nombre_hoja, celda, destino, explicita=explicita,
                    clase="complementario", representaciones=representaciones,
                    visible=visible)
                if bloqueo is not None:
                    bloqueos.append(bloqueo)
                if not mapear:
                    continue
                anterior = destino.value
                destino.value = visible
                registro = {"hoja": nombre_hoja, "etiqueta": etiqueta,
                            "celda_etiqueta": celda.coordinate,
                            "celda_valor": destino.coordinate,
                            "valor_anterior": anterior}
                if not explicita:
                    registro["nota"] = "etiqueta desnuda (campo único)"
                corregidas.append(registro)
    return corregidas, bloqueos


def escribir_bloque_identidad(hoja, identidad: dict | None, *, idioma: str = "es",
                              celdas_mapeadas: list[dict] | None = None,
                              contradicciones: list[dict] | None = None,
                              columna: int = 1,
                              hojas_exentas: list[str] | None = None) -> dict:
    """Escribe el bloque de identidad bajo el contenido existente del resumen.

    Devuelve metadatos verificables: ``{"identidad_verificada",
    "decision_usuario_confirmada", "estado_identidad", "advertencias",
    "fila_inicio"}``. Sólo toca celdas nuevas del bloque identificado.
    V15.9.4 (CRIT-N1/N2): lista las referencias contradictorias visibles
    (títulos/encabezados de OTRO proyecto) sin reemplazarlas nunca.
    V15.9.5 (ALTA-N9): separa TRES estados que antes se confundían:

    - ``decision_usuario_confirmada``: el usuario declaró un proyecto o
      eligió expresamente «Sin asignar».
    - ``identidad_verificada``: hay celda de proyecto mapeada, NO es una
      decisión «Sin asignar» y no queda ninguna contradicción visible.
      «Sin asignar» es pendiente, jamás identidad confirmada.
    - ``estado_identidad``: ``verificada`` / ``pendiente_sin_asignar`` /
      ``bloqueada_contradiccion`` / ``no_verificada``.
    """
    t = textos_identidad(idioma)
    identidad = identidad or {}
    advertencias = []
    mapeadas = celdas_mapeadas or []
    contradicciones = contradicciones or []
    sin_asignar = bool(identidad.get("sin_asignar"))
    origen = str(identidad.get("origen") or "")
    # P08: sin celda identificable en la plantilla ⇒ reporte identificado
    # bloqueado con advertencia visible; no se suplanta ningún título.
    if not mapeadas and not sin_asignar and identidad:
        advertencias.append(t["advertencia"])
    identidad_verificada = bool(mapeadas) and not sin_asignar and not contradicciones
    decision_usuario_confirmada = bool(identidad) and (
        sin_asignar or origen not in _ORIGENES_LEGACY)
    estado_identidad = ("bloqueada_contradiccion" if contradicciones
                        else "verificada" if identidad_verificada
                        else "pendiente_sin_asignar" if sin_asignar
                        else "no_verificada")

    fila = max((hoja.max_row or 1), 1) + 2
    inicio = fila

    def _escribir(etiqueta: str, valor) -> None:
        nonlocal fila
        hoja.cell(fila, columna, etiqueta)
        hoja.cell(fila, columna + 1, valor)
        fila += 1

    from openpyxl.styles import Font
    celda_titulo = hoja.cell(fila, columna, t["bloque"])
    celda_titulo.font = Font(bold=True)
    fila += 1
    _escribir(t["nombre"], nombre_visible(identidad, idioma))
    _escribir(t["id"], identidad.get("project_id"))
    _escribir(t["origen"], identidad.get("origen"))
    _escribir(t["estado"], (t["sin_asignar"] if identidad.get("sin_asignar")
                            else t["pendiente"] if identidad.get("pendiente")
                            else t["confirmado"]))
    # V15.9.5 (ALTA-N9): los tres estados quedan por escrito en el bloque.
    _escribir(t["verificacion"],
              t["afirmativo"] if identidad_verificada else t["negativo"])
    _escribir(t["decision"],
              t["afirmativo"] if decision_usuario_confirmada else t["negativo"])
    if hojas_exentas:
        _escribir(t["exentas"], ", ".join(hojas_exentas))
    if identidad.get("ejecucion_id"):
        _escribir(t["ejecucion"], identidad.get("ejecucion_id"))
    _escribir(t["fecha"], identidad.get("fecha")
              or datetime.now().isoformat(timespec="seconds"))
    if identidad.get("responsable"):
        _escribir(t["responsable"], identidad.get("responsable"))
    for celda in mapeadas:
        _escribir(f"{t['id']} · {celda['hoja']}!{celda['celda_valor']}",
                  f"{celda.get('valor_anterior')!r} → "
                  f"{nombre_visible(identidad, idioma)!r}")
    for aviso in advertencias:
        aviso_celda = hoja.cell(fila, columna, aviso)
        aviso_celda.font = Font(bold=True, color="9C0006")
        fila += 1
    for referencia in contradicciones:
        texto_aviso = t["contradiccion"].format(
            ref=f"{referencia['hoja']}!{referencia['celda']}",
            texto=referencia["texto"])
        if referencia.get("detalle"):
            # V15.9.5 (MEDIA-ALTA-N8): explicar exactamente qué no se pudo
            # verificar (token de otro nombre o referencia ambigua).
            texto_aviso += " " + t["detalle"].format(detalle=referencia["detalle"])
        celda_aviso = hoja.cell(fila, columna, texto_aviso)
        celda_aviso.font = Font(bold=True, color="9C0006")
        fila += 1
    return {"identidad_verificada": identidad_verificada,
            "decision_usuario_confirmada": decision_usuario_confirmada,
            "estado_identidad": estado_identidad,
            "advertencias": advertencias, "fila_inicio": inicio}


def aplicar_identidad_a_salida(libro, identidad: dict | None, *, idioma: str = "es",
                               hojas_busqueda: list[str], hoja_resumen: str,
                               fila_max: int = 4,
                               hojas_exentas: list[str] | None = None) -> dict:
    """Paso completo sobre el libro de SALIDA (nunca sobre la plantilla).

    1. Localiza celdas de proyecto identificadas por etiqueta.
    2. Sustituye SU VALOR por la identidad autoritativa (celdas concretas).
    3. Escribe el bloque de identidad/advertencias en la hoja resumen.

    V15.9.3 (CRIT-03): añade ``publicable`` a los metadatos. Con
    ``publicable=False`` el generador debe guardar un borrador no publicable
    en ubicación diferenciada y NO tocar el destino oficial.

    V15.9.4 (CRIT-N1/N2): verificar una celda etiquetada NO basta — se
    escanean las hojas en busca de referencias visibles a OTRO proyecto.
    Los textos contradictorios NUNCA se reemplazan; se listan en el bloque.

    V15.9.5 (CRIT-N7): el examen cubre TODAS las hojas visibles del
    entregable (``hojas_del_entregable``), no sólo ``hojas_busqueda``; el
    generador puede declarar ``hojas_exentas`` (histórico/evidencia) y quedan
    documentadas en el bloque. Se examinan también cabeceras/pies de página.

    V15.9.5 (ALTA-N9): política de publicación explícita con tres estados:

    - identidad VERIFICADA (celda mapeada + sin contradicciones) ⇒ publica;
    - «Sin asignar» expreso sin contradicciones ⇒ publica como PENDIENTE
      (P02) con estado diferenciado (``estado_identidad=pendiente_sin_asignar``,
      ``identidad_verificada=False``): jamás como proyecto verificado;
    - contradicción o plantilla sin celda identificable ⇒ borrador bloqueado.
    Los orígenes legacy conservan el flujo P11 de V15.9.2.

    V15.9.6 (CRIT-N10): además del mapeo declarado, se verifica
    ESTRUCTURALMENTE cada par etiqueta/valor de proyecto en TODAS las hojas
    elegibles (``mapear_pares_etiqueta_valor``): los pares editables se
    actualizan a la identidad de la ejecución (registrados en
    ``celdas_complementarias``) y los no resolubles sin riesgo bloquean la
    publicación señalando hoja/celda. El resumen y las hojas auxiliares
    quedan con identidad coherente.
    """
    if not identidad:
        return {"identidad_verificada": None, "publicable": True,
                "decision_usuario_confirmada": False,
                "estado_identidad": "sin_identidad",
                "advertencias": [], "celdas": [], "celdas_complementarias": [],
                "contradicciones": []}
    celdas = buscar_celdas_proyecto(libro, hojas_busqueda, fila_max=fila_max)
    visible = nombre_visible(identidad, idioma)
    representaciones = _representaciones_identidad(identidad, idioma)
    # V15.9.7 (CRIT-N13) + V15.9.8 (CRIT-N15): política segura ÚNICA para
    # todo mapeo (declarado y complementario). Cada candidato se CLASIFICA
    # con `_clasificar_par` ANTES de escribir: una cabecera tabular
    # («Proyecto | Etapa [| Caso]») no es campo de proyecto y no se mapea;
    # la estructura ambigua se valida contra la identidad; fórmulas, valores
    # no textuales y celdas combinadas NO se sobrescriben nunca. El campo
    # queda intacto y los bloqueos se registran como contradicciones
    # estructurales (hoja/celda/motivo) que impiden publicar; sin
    # excepciones de bajo nivel (AttributeError de MergedCell) ni éxitos
    # falsos. Una celda no validada no cuenta como identidad verificada.
    mapeadas: list[dict] = []
    bloqueos_declarados: list[dict] = []
    for celda in celdas:
        hoja = libro[celda["hoja"]]
        destino = hoja[celda["celda_valor"]]
        explicita = bool(_ETIQUETA_EXPLICITA.match(celda["etiqueta"]))
        mapear, bloqueo = _clasificar_par(
            hoja, celda["hoja"], hoja[celda["celda_etiqueta"]], destino,
            explicita=explicita, clase="declarado",
            representaciones=representaciones, visible=visible)
        if bloqueo is not None:
            bloqueos_declarados.append(bloqueo)
        if not mapear:
            continue
        destino.value = visible
        mapeadas.append(celda)
    # El escaneo va DESPUÉS de sustituir las celdas mapeadas (ya contienen la
    # identidad declarada) y ANTES de escribir el bloque (sus textos propios
    # no deben auto-detectarse).
    # V15.9.5 (CRIT-N7): contrato de hojas del entregable — todas las
    # visibles, más las declaradas por el generador aunque estuvieran
    # ocultas, menos las exentas documentadas.
    exentas = sorted({str(nombre) for nombre in (hojas_exentas or ())})
    declaradas = list(dict.fromkeys([*hojas_busqueda, hoja_resumen]))
    examinar = [nombre for nombre
                in dict.fromkeys(declaradas + hojas_del_entregable(libro, exentas))
                if nombre not in set(exentas)]
    # V15.9.6 (CRIT-N10): verificación estructural de pares etiqueta/valor
    # en TODAS las hojas elegibles (complementa la revisión de textos).
    # V15.9.8 (N15): la ruta DECLARADA clasifica cada etiqueta encontrada
    # con la MISMA política centralizada (`_clasificar_par`) antes de esta
    # exclusión: `excluidas` sólo evita reclasificar pares cuyo resultado
    # (mapeado, omitido como cabecera tabular o bloqueado) ya quedó
    # registrado arriba; nunca suprime la validación estructural.
    excluidas = {(c["hoja"], c["celda_etiqueta"]) for c in celdas}
    corregidas, bloqueos = mapear_pares_etiqueta_valor(
        libro, examinar, identidad, idioma=idioma, excluidas=excluidas)
    # V15.9.7 (CRIT-N13): los campos declarados NO editables suman su bloqueo
    # estructural a las contradicciones (nunca se cuentan como identificadas).
    contradicciones = bloqueos_declarados + bloqueos + \
        detectar_referencias_contradictorias(
            libro, examinar, identidad, idioma=idioma)
    identificadas = [*mapeadas, *corregidas]
    meta = {"celdas": mapeadas, "celdas_complementarias": corregidas,
            "contradicciones": contradicciones,
            "hojas_examinadas": examinar, "hojas_exentas": exentas}
    sin_asignar = bool(identidad.get("sin_asignar"))
    origen = str(identidad.get("origen") or "")
    if hoja_resumen in libro.sheetnames:
        meta.update(escribir_bloque_identidad(libro[hoja_resumen], identidad,
                                              idioma=idioma,
                                              celdas_mapeadas=identificadas,
                                              contradicciones=contradicciones,
                                              hojas_exentas=exentas))
    else:
        meta["identidad_verificada"] = (bool(identificadas) and not sin_asignar
                                        and not contradicciones)
        meta["decision_usuario_confirmada"] = (sin_asignar
                                               or origen not in _ORIGENES_LEGACY)
        meta["estado_identidad"] = (
            "bloqueada_contradiccion" if contradicciones
            else "verificada" if meta["identidad_verificada"]
            else "pendiente_sin_asignar" if sin_asignar
            else "no_verificada")
        meta["advertencias"] = []
    # V15.9.5 (ALTA-N9): «Sin asignar» NO suma a identidad_verificada; su
    # publicación como pendiente es una rama explícita de la política.
    meta["publicable"] = (bool(meta.get("identidad_verificada"))
                          or (sin_asignar and not contradicciones)
                          or origen in _ORIGENES_LEGACY)
    if not meta["publicable"] and hoja_resumen in libro.sheetnames:
        marcar_borrador_no_publicable(libro, hoja_resumen, idioma=idioma)
    return meta


# Orígenes de identidad que NO constituyen una declaración auditada del
# usuario: conservan el flujo legacy (publicación con advertencia visible).
_ORIGENES_LEGACY = {"nombre_carpeta_legacy", "sin_datos", ""}


def marcar_borrador_no_publicable(libro, hoja_resumen: str, *, idioma: str = "es") -> None:
    """V15.9.3 (CRIT-03): marca visible de borrador bloqueado en el resumen."""
    if hoja_resumen not in libro.sheetnames:
        return
    from openpyxl.styles import Font
    hoja = libro[hoja_resumen]
    fila = max((hoja.max_row or 1), 1) + 1
    celda = hoja.cell(fila, 1, textos_identidad(idioma)["borrador"])
    celda.font = Font(bold=True, color="9C0006")
