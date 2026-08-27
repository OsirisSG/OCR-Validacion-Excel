"""
Analizador estructural y semántico de libros Excel.

Recorre todas las hojas de un archivo .xlsx/.xlsm sin modificarlo y genera un
JSON con encabezados, tipos, valores faltantes, fórmulas, cardinalidad,
categorías observadas y posibles nombres canónicos para cada columna.

El análisis usa lectura incremental para mantener acotado el consumo de memoria.
Las fórmulas se cuentan, pero no se evalúan: openpyxl solo puede leer el último
valor almacenado por Excel u otra aplicación compatible.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import statistics
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from itertools import islice, zip_longest
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


EXTENSIONES_ADMITIDAS = {".xlsx", ".xlsm"}
VALORES_BOOLEANOS = {
    "si", "sí", "no", "true", "false", "verdadero", "falso", "yes", "y", "n",
    "activo", "inactivo",
}
VALORES_ESTADO = {
    "verde", "amarillo", "rojo", "pendiente", "aprobado", "rechazado",
    "completado", "cancelado", "activo", "inactivo", "abierto", "cerrado",
}

# Categoría semántica -> nombres canónicos y términos que suelen aparecer en el
# encabezado. Las coincidencias se puntúan; no se presentan como certeza.
REGLAS_SEMANTICAS = (
    ("identificador", "identificador", {"id", "identificador", "folio", "codigo", "clave", "sku", "uuid"}),
    ("nombre", "nombre", {"nombre", "name", "razon social", "cliente", "proveedor", "producto"}),
    ("fecha_hora", "fecha", {"fecha", "date", "hora", "time", "creado", "actualizado", "timestamp"}),
    ("estado_resultado", "estado", {"estado", "status", "estatus", "semaforo", "resultado", "veredicto"}),
    ("categoria_clasificacion", "categoria", {"categoria", "category", "tipo", "clase", "grupo", "familia", "segmento"}),
    ("descripcion", "descripcion", {"descripcion", "description", "detalle", "observacion", "comentario", "notas"}),
    ("ruta_archivo", "ruta", {"ruta", "path", "archivo", "file", "imagen", "documento", "carpeta", "directorio"}),
    ("confianza_porcentaje", "confianza_pct", {"confianza", "confidence", "score", "probabilidad", "porcentaje", "pct"}),
    ("cantidad", "cantidad", {"cantidad", "conteo", "unidades", "qty", "total", "numero", "num"}),
    ("importe_moneda", "importe", {"precio", "costo", "importe", "monto", "subtotal", "venta", "saldo"}),
    ("contacto", "contacto", {"correo", "email", "telefono", "phone", "movil", "contacto"}),
    ("ubicacion", "ubicacion", {"direccion", "address", "ciudad", "pais", "estado geografico", "cp", "postal"}),
)


def _texto_simple(valor: Any) -> str:
    """Representación estable y serializable para muestras y categorías."""
    if isinstance(valor, (datetime, date, time)):
        return valor.isoformat()
    return str(valor).strip()


def normalizar_nombre(valor: Any, respaldo: str = "columna") -> str:
    """Convierte un encabezado a snake_case ASCII sin perder el original."""
    texto = unicodedata.normalize("NFKD", _texto_simple(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c)).lower()
    texto = re.sub(r"[^a-z0-9]+", "_", texto).strip("_")
    return texto or respaldo


def _es_vacio(valor: Any) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _tipo_valor(valor: Any) -> str:
    if _es_vacio(valor):
        return "vacio"
    if isinstance(valor, bool):
        return "booleano"
    if isinstance(valor, datetime):
        return "fecha_hora"
    if isinstance(valor, date):
        return "fecha"
    if isinstance(valor, time):
        return "hora"
    if isinstance(valor, int):
        return "entero"
    if isinstance(valor, float):
        return "decimal" if math.isfinite(valor) else "numero_no_finito"
    if isinstance(valor, str) and valor.startswith("="):
        return "formula"
    if isinstance(valor, str) and valor.startswith("#"):
        return "error_excel"
    return "texto"


def _firma(valor: Any) -> tuple[str, str]:
    """Clave acotada para contar valores heterogéneos sin confundir 1 con '1'."""
    return (_tipo_valor(valor), _texto_simple(valor))


def _valor_json(valor: Any) -> Any:
    if isinstance(valor, (datetime, date, time)):
        return valor.isoformat()
    if isinstance(valor, float) and not math.isfinite(valor):
        return str(valor)
    return valor


def _posibles_significados(encabezado: str, tipos: Counter, categorias: list[str]) -> list[dict]:
    normalizado = normalizar_nombre(encabezado)
    palabras = set(normalizado.split("_"))
    frase = normalizado.replace("_", " ")
    propuestas: list[dict] = []

    for categoria, canonico, terminos in REGLAS_SEMANTICAS:
        coincidencias = []
        for termino in terminos:
            termino_normalizado = normalizar_nombre(termino)
            tokens = set(termino_normalizado.split("_"))
            if termino_normalizado == normalizado:
                coincidencias.append(1.0)
            elif termino_normalizado in frase or tokens <= palabras:
                coincidencias.append(0.86)
        if coincidencias:
            propuestas.append({
                "categoria": categoria,
                "nombre_canonico": canonico,
                "confianza": max(coincidencias),
                "motivo": "coincidencia con el encabezado",
            })

    tipos_datos = {tipo for tipo, cantidad in tipos.items() if cantidad and tipo not in {"vacio", "formula"}}
    categorias_norm = {normalizar_nombre(v) for v in categorias}
    inferencias: list[tuple[str, str, float, str]] = []
    if tipos_datos and tipos_datos <= {"fecha", "fecha_hora", "hora"}:
        inferencias.append(("fecha_hora", "fecha", 0.82, "tipo de datos observado"))
    if categorias_norm and categorias_norm <= {normalizar_nombre(v) for v in VALORES_BOOLEANOS}:
        inferencias.append(("indicador_booleano", "es_activo", 0.78, "valores binarios observados"))
    if categorias_norm & {normalizar_nombre(v) for v in VALORES_ESTADO}:
        inferencias.append(("estado_resultado", "estado", 0.76, "valores de estado observados"))
    if tipos_datos and tipos_datos <= {"entero", "decimal"} and any(
        token in palabras for token in {"porcentaje", "pct", "confianza", "probabilidad"}
    ):
        inferencias.append(("confianza_porcentaje", "confianza_pct", 0.9, "encabezado y valores numéricos"))

    existentes = {(p["categoria"], p["nombre_canonico"]) for p in propuestas}
    for categoria, canonico, confianza, motivo in inferencias:
        if (categoria, canonico) not in existentes:
            propuestas.append({
                "categoria": categoria,
                "nombre_canonico": canonico,
                "confianza": confianza,
                "motivo": motivo,
            })

    propuestas.sort(key=lambda p: (-p["confianza"], p["categoria"]))
    return propuestas[:4]


def _puntuar_fila_encabezado(filas: list[tuple[Any, ...]], indice: int) -> float:
    fila = filas[indice]
    no_vacios = [v for v in fila if not _es_vacio(v)]
    if not no_vacios:
        return float("-inf")
    ancho = max(1, len(fila))
    if ancho > 1 and len(no_vacios) == 1:
        return -1.0 - indice * 0.05

    textos = sum(isinstance(v, str) and not str(v).startswith("=") for v in no_vacios)
    unicos = len({normalizar_nombre(v) for v in no_vacios})
    densidad = len(no_vacios) / ancho
    score = len(no_vacios) + 2.2 * (textos / len(no_vacios)) + densidad + (unicos / len(no_vacios))

    siguientes = filas[indice + 1: indice + 4]
    if siguientes:
        densidades = [sum(not _es_vacio(v) for v in f) / max(1, ancho) for f in siguientes]
        score += 1.5 * statistics.fmean(densidades)
    # Ante filas igualmente densas, el encabezado suele ser la primera. Esta
    # penalización evita confundir la primera fila de datos con el encabezado,
    # incluso si este último contiene nombres repetidos.
    score -= indice * 0.3
    return score


def detectar_fila_encabezado(filas: list[tuple[Any, ...]]) -> tuple[int | None, float]:
    """Retorna índice base cero dentro del buffer y una confianza heurística."""
    if not filas or not any(any(not _es_vacio(v) for v in fila) for fila in filas):
        return None, 0.0
    puntajes = [_puntuar_fila_encabezado(filas, i) for i in range(len(filas))]
    mejor = max(range(len(puntajes)), key=puntajes.__getitem__)
    score = puntajes[mejor]
    confianza = max(0.0, min(1.0, (score - 1.5) / 8.0))
    return mejor, round(confianza, 3)


@dataclass
class PerfilColumna:
    indice: int
    encabezado_original: str
    nombre_normalizado: str
    max_muestras: int
    max_categorias: int
    total_filas: int = 0
    vacios: int = 0
    formulas: int = 0
    tipos: Counter = field(default_factory=Counter)
    frecuencias: Counter = field(default_factory=Counter)
    muestras: list[Any] = field(default_factory=list)
    valores_numericos: list[float] = field(default_factory=list)
    unicos_truncados: bool = False

    def agregar(self, valor: Any, formula: Any = None) -> None:
        self.total_filas += 1
        es_formula = isinstance(formula, str) and formula.startswith("=")
        if es_formula:
            self.formulas += 1
            self.tipos["formula"] += 1
        if _es_vacio(valor):
            if not es_formula:
                self.vacios += 1
                self.tipos["vacio"] += 1
            return

        tipo = _tipo_valor(valor)
        self.tipos[tipo] += 1
        if len(self.muestras) < self.max_muestras:
            serializable = _valor_json(valor)
            if serializable not in self.muestras:
                self.muestras.append(serializable)

        firma = _firma(valor)
        limite_frecuencias = max(200, self.max_categorias + 1)
        if firma in self.frecuencias or len(self.frecuencias) < limite_frecuencias:
            self.frecuencias[firma] += 1
        else:
            self.unicos_truncados = True

        if tipo in {"entero", "decimal"} and len(self.valores_numericos) < 10_000:
            self.valores_numericos.append(float(valor))

    def resultado(self) -> dict:
        no_vacios = self.total_filas - self.vacios
        unicos = len(self.frecuencias)
        categorias_ordenadas = [
            (firma[1], cantidad)
            for firma, cantidad in self.frecuencias.most_common(self.max_categorias)
        ]
        es_categorica = (
            not self.unicos_truncados
            and no_vacios >= 2
            and 0 < unicos <= self.max_categorias
            and (unicos / no_vacios <= 0.5 or set(self.tipos) <= {"booleano", "texto", "vacio", "formula"})
        )
        categorias = [valor for valor, _ in categorias_ordenadas] if es_categorica else []
        tipo_dominante = "vacio"
        candidatos_tipo = {k: v for k, v in self.tipos.items() if k not in {"vacio", "formula"}}
        if candidatos_tipo:
            tipo_dominante = max(candidatos_tipo, key=candidatos_tipo.get)
        elif self.formulas:
            tipo_dominante = "formula"

        numerico = None
        if self.valores_numericos:
            numerico = {
                "minimo": min(self.valores_numericos),
                "maximo": max(self.valores_numericos),
                "promedio": statistics.fmean(self.valores_numericos),
            }

        return {
            "indice_excel": self.indice,
            "letra_excel": get_column_letter(self.indice),
            "encabezado_original": self.encabezado_original,
            "nombre_normalizado": self.nombre_normalizado,
            "tipo_dominante": tipo_dominante,
            "distribucion_tipos": dict(sorted(self.tipos.items())),
            "filas_analizadas": self.total_filas,
            "valores_no_vacios": no_vacios,
            "valores_vacios": self.vacios,
            "porcentaje_vacio": round(self.vacios / self.total_filas * 100, 2) if self.total_filas else 0.0,
            "formulas": self.formulas,
            "valores_unicos_observados": unicos,
            "conteo_unicos_truncado": self.unicos_truncados,
            "es_posible_categoria": es_categorica,
            "categorias_posibles": categorias,
            "frecuencias_categorias": (
                [{"valor": valor, "conteo": cantidad} for valor, cantidad in categorias_ordenadas]
                if es_categorica else []
            ),
            "muestras": self.muestras,
            "resumen_numerico": numerico,
            "posibles_significados": _posibles_significados(
                self.encabezado_original, self.tipos, categorias
            ),
        }


def _hacer_encabezados_unicos(valores: Iterable[Any], ancho: int) -> list[tuple[str, str]]:
    usados: Counter = Counter()
    salida = []
    valores = list(valores)
    for indice in range(ancho):
        original = _texto_simple(valores[indice]) if indice < len(valores) and not _es_vacio(valores[indice]) else f"Columna {get_column_letter(indice + 1)}"
        base = normalizar_nombre(original, f"columna_{indice + 1}")
        usados[base] += 1
        unico = base if usados[base] == 1 else f"{base}_{usados[base]}"
        salida.append((original, unico))
    return salida


def _fila_combinada(fila_formula: tuple[Any, ...], fila_valor: tuple[Any, ...]) -> tuple[Any, ...]:
    ancho = max(len(fila_formula), len(fila_valor))
    return tuple(
        (fila_valor[i] if i < len(fila_valor) and fila_valor[i] is not None else
         fila_formula[i] if i < len(fila_formula) else None)
        for i in range(ancho)
    )


def _analizar_hoja(
    hoja_formula,
    hoja_valores,
    max_filas_encabezado: int,
    max_muestras: int,
    max_categorias: int,
) -> dict:
    iter_formula = hoja_formula.iter_rows(values_only=True)
    iter_valores = hoja_valores.iter_rows(values_only=True)
    buffer_pares = list(islice(
        zip_longest(iter_formula, iter_valores, fillvalue=()),
        max_filas_encabezado,
    ))
    filas_combinadas = [_fila_combinada(f, v) for f, v in buffer_pares]
    indice_encabezado, confianza = detectar_fila_encabezado(filas_combinadas)

    if indice_encabezado is None:
        return {
            "nombre": hoja_formula.title,
            "estado": hoja_formula.sheet_state,
            "vacia": True,
            "fila_encabezado": None,
            "confianza_encabezado": 0.0,
            "filas_con_datos": 0,
            "columnas_con_datos": 0,
            "formulas": 0,
            "columnas": [],
            "advertencias": ["La hoja no contiene valores."],
        }

    ancho_buffer = max((len(f) for f in filas_combinadas), default=0)
    encabezados = _hacer_encabezados_unicos(filas_combinadas[indice_encabezado], ancho_buffer)
    perfiles = [
        PerfilColumna(i + 1, original, normalizado, max_muestras, max_categorias)
        for i, (original, normalizado) in enumerate(encabezados)
    ]

    ultima_fila = indice_encabezado + 1
    ultima_columna = 0
    filas_datos_no_vacias = 0

    def procesar(numero_fila: int, fila_formula: tuple[Any, ...], fila_valor: tuple[Any, ...]) -> None:
        nonlocal ultima_fila, ultima_columna, filas_datos_no_vacias, perfiles
        ancho = max(len(fila_formula), len(fila_valor))
        if ancho > len(perfiles):
            nuevos = _hacer_encabezados_unicos([], ancho)
            existentes = {p.nombre_normalizado for p in perfiles}
            for i in range(len(perfiles), ancho):
                original, normalizado = nuevos[i]
                while normalizado in existentes:
                    normalizado += "_2"
                existentes.add(normalizado)
                perfiles.append(PerfilColumna(i + 1, original, normalizado, max_muestras, max_categorias))

        fila_tiene_datos = False
        for i, perfil in enumerate(perfiles):
            formula = fila_formula[i] if i < len(fila_formula) else None
            es_formula = isinstance(formula, str) and formula.startswith("=")
            valor = fila_valor[i] if i < len(fila_valor) else None
            if not es_formula:
                valor = formula
            if not _es_vacio(formula) or not _es_vacio(valor):
                fila_tiene_datos = True
                ultima_columna = max(ultima_columna, i + 1)
            perfil.agregar(valor, formula)
        if fila_tiene_datos:
            filas_datos_no_vacias += 1
            ultima_fila = numero_fila

    for offset, (fila_formula, fila_valor) in enumerate(buffer_pares[indice_encabezado + 1:], start=indice_encabezado + 2):
        procesar(offset, fila_formula, fila_valor)
    for numero_fila, (fila_formula, fila_valor) in enumerate(
        zip_longest(iter_formula, iter_valores, fillvalue=()), start=len(buffer_pares) + 1
    ):
        procesar(numero_fila, fila_formula, fila_valor)

    columnas_resultado = [p.resultado() for p in perfiles[:ultima_columna or len(perfiles)]]
    advertencias = []
    if confianza < 0.5:
        advertencias.append("La fila de encabezado fue inferida con confianza baja; conviene revisarla.")
    normalizados = [c["nombre_normalizado"] for c in columnas_resultado]
    if len(normalizados) != len(set(normalizados)):
        advertencias.append("Persisten nombres de columna duplicados después de normalizar.")

    return {
        "nombre": hoja_formula.title,
        "estado": hoja_formula.sheet_state,
        "vacia": False,
        "fila_encabezado": indice_encabezado + 1,
        "confianza_encabezado": confianza,
        "filas_antes_del_encabezado": indice_encabezado,
        "filas_de_datos_no_vacias": filas_datos_no_vacias,
        "ultima_fila_con_datos": ultima_fila,
        "columnas_con_datos": ultima_columna or len(perfiles),
        "formulas": sum(c["formulas"] for c in columnas_resultado),
        "columnas": columnas_resultado,
        "advertencias": advertencias,
    }


def analizar_libro(
    ruta_excel: str | Path,
    *,
    max_filas_encabezado: int = 25,
    max_muestras: int = 5,
    max_categorias: int = 20,
) -> dict:
    """Analiza todas las hojas de un libro y retorna un diccionario JSON-safe."""
    ruta = Path(ruta_excel).expanduser().resolve()
    if not ruta.is_file():
        raise FileNotFoundError(f"No existe el archivo Excel: {ruta}")
    if ruta.suffix.lower() not in EXTENSIONES_ADMITIDAS:
        raise ValueError(
            f"Formato no admitido: {ruta.suffix or '(sin extensión)'}. "
            "Usa .xlsx o .xlsm; convierte primero los archivos .xls heredados."
        )
    if max_filas_encabezado < 1 or max_muestras < 0 or max_categorias < 1:
        raise ValueError("Los límites de análisis deben ser valores positivos.")

    keep_vba = ruta.suffix.lower() == ".xlsm"
    libro_formula = load_workbook(ruta, read_only=True, data_only=False, keep_vba=keep_vba)
    libro_valores = load_workbook(ruta, read_only=True, data_only=True, keep_vba=keep_vba)
    try:
        hojas = []
        for nombre in libro_formula.sheetnames:
            hojas.append(_analizar_hoja(
                libro_formula[nombre],
                libro_valores[nombre],
                max_filas_encabezado,
                max_muestras,
                max_categorias,
            ))
        propiedades = libro_formula.properties
        definidos = sorted(str(nombre) for nombre in libro_formula.defined_names)
        return {
            "archivo": str(ruta),
            "nombre_archivo": ruta.name,
            "tamano_bytes": ruta.stat().st_size,
            "formato": ruta.suffix.lower().lstrip("."),
            "propiedades": {
                "titulo": propiedades.title,
                "asunto": propiedades.subject,
                "creador": propiedades.creator,
                "ultima_modificacion_por": propiedades.lastModifiedBy,
                "creado": _valor_json(propiedades.created),
                "modificado": _valor_json(propiedades.modified),
            },
            "nombres_definidos": definidos,
            "total_hojas": len(hojas),
            "hojas_visibles": sum(h["estado"] == "visible" for h in hojas),
            "hojas_ocultas": [h["nombre"] for h in hojas if h["estado"] != "visible"],
            "total_formulas": sum(h["formulas"] for h in hojas),
            "hojas": hojas,
            "metodologia": {
                "solo_lectura": True,
                "formulas_evaluadas": False,
                "max_filas_busqueda_encabezado": max_filas_encabezado,
                "max_muestras_por_columna": max_muestras,
                "max_categorias_por_columna": max_categorias,
            },
        }
    finally:
        libro_formula.close()
        libro_valores.close()


def guardar_analisis(resultado: dict, ruta_salida: str | Path) -> Path:
    ruta = Path(ruta_salida).expanduser().resolve()
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(json.dumps(resultado, ensure_ascii=False, indent=2), encoding="utf-8")
    return ruta


def _resumen_terminal(resultado: dict) -> str:
    lineas = [
        f"Libro: {resultado['nombre_archivo']}",
        f"Hojas: {resultado['total_hojas']} | Fórmulas: {resultado['total_formulas']}",
    ]
    for hoja in resultado["hojas"]:
        if hoja["vacia"]:
            lineas.append(f"  - {hoja['nombre']}: vacía ({hoja['estado']})")
        else:
            lineas.append(
                f"  - {hoja['nombre']}: encabezado fila {hoja['fila_encabezado']}, "
                f"{hoja['columnas_con_datos']} columnas, "
                f"{hoja['filas_de_datos_no_vacias']} filas de datos"
            )
    return "\n".join(lineas)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Analiza todas las hojas, columnas, tipos y categorías de un Excel."
    )
    parser.add_argument("excel", help="Ruta del archivo .xlsx o .xlsm.")
    parser.add_argument("--salida", help="JSON de salida; por defecto: <archivo>_analisis_excel.json")
    parser.add_argument("--max-filas-encabezado", type=int, default=25)
    parser.add_argument("--max-muestras", type=int, default=5)
    parser.add_argument("--max-categorias", type=int, default=20)
    parser.add_argument("--sin-muestras", action="store_true", help="No incluye valores de ejemplo.")
    args = parser.parse_args()

    ruta = Path(args.excel)
    salida = Path(args.salida) if args.salida else ruta.with_name(f"{ruta.stem}_analisis_excel.json")
    try:
        resultado = analizar_libro(
            ruta,
            max_filas_encabezado=args.max_filas_encabezado,
            max_muestras=0 if args.sin_muestras else args.max_muestras,
            max_categorias=args.max_categorias,
        )
        guardado = guardar_analisis(resultado, salida)
    except (FileNotFoundError, ValueError, OSError) as exc:
        parser.error(str(exc))
    print(_resumen_terminal(resultado))
    print(f"Análisis JSON: {guardado}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
