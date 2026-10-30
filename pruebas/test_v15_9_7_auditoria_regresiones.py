"""Pruebas adicionales independientes de integridad de identidad del proyecto.
Todos los libros son sintéticos y existen solo en memoria."""
import pytest
from openpyxl import Workbook
from utilidades.identidad_excel import aplicar_identidad_a_salida

IDENTIDAD_B = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b", "origen": "declarado_usuario", "sin_asignar": False}

def preparar():
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    resumen["A2"] = "Proyecto:"
    resumen["B2"] = "Proyecto A"
    return libro

@pytest.mark.parametrize("valor", ['=CONCAT("Proyecto ","A")', 2025])
def test_n13_no_sobrescribir_formula_ni_numero_en_celda_declarada(valor):
    libro = preparar()
    resumen = libro["Resumen"]
    resumen["B2"] = valor
    resultado = aplicar_identidad_a_salida(libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert resumen["B2"].value == valor, "Una fórmula/un dato numérico no se debe destruir para identificar el proyecto"
    assert resultado["publicable"] is False, "Una identidad cuyo campo oficial contiene una fórmula o número debe bloquearse"


def test_n13_celda_combinada_de_destino_no_debe_provocar_excepcion_no_controlada():
    libro = preparar()
    resumen = libro["Resumen"]
    resumen["B2"] = None
    resumen.merge_cells("A2:B2")
    resultado = aplicar_identidad_a_salida(libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert resultado["publicable"] is False
    assert resultado["contradicciones"], "El mapa de celdas debe señalar la estructura no editable"


def test_n14_otro_proyecto_en_hoja_visible_sin_dos_puntos_no_puede_publicarse_silenciosamente():
    libro = preparar()
    hoja = libro.create_sheet("Indicadores")
    hoja["A1"] = "Proyecto"
    hoja["B1"] = "A"
    resultado = aplicar_identidad_a_salida(libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert resultado["publicable"] is False or hoja["B1"].value == "Proyecto B", (
        "Si una hoja adicional exhibe Proyecto/A, debe mapearse con prueba o bloquearse; nunca publicar B con A visible")
