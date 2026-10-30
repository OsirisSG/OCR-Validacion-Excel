"""V15.9.10 · Ampliación de la corrección CRIT-N17.

Protección ESTRUCTURAL de cabeceras: tablas Excel registradas
(``ws.tables``/``Table.ref``) y autofiltros mandan sobre cualquier
heurística de texto; un par desnudo sin tabla formal ni evidencia positiva
de que el contiguo sea un nombre de proyecto se trata fail-closed (no se
escribe por suposición). Cobertura exigida por
`CORRECCIONES_OBLIGATORIAS_V_MATRIX_V15_9_10.md`: tablas vacías y pobladas,
posiciones arbitrarias, varias tablas por hoja, 3 idiomas, combinadas y
colisiones, campo explícito en cuerpo de tabla, persistencia con Excel
guardados/releídos (nombres de columna y ``Table.ref``), exportadores
empresarial y BMG reales con hash de plantilla, y controles legítimos
(campo único con mención, explícito, valor vacío).
"""
from __future__ import annotations

import pathlib
import sys

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "pruebas"))

from test_plantilla_empresarial import _plantilla
from test_v15_9_6_ampliacion import _hash
from utilidades.identidad_excel import aplicar_identidad_a_salida, nombre_visible

IDENTIDAD_B = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b",
               "origen": "declarado_usuario", "sin_asignar": False,
               "pendiente": False, "ejecucion_id": "ej-v15910"}
IDENTIDAD_TEMPERATURA = {"nombre_proyecto": "Temperatura",
                         "project_id": "temperatura",
                         "origen": "declarado_usuario", "sin_asignar": False}


def _libro_base():
    libro = Workbook()
    libro.active.title = "Resumen"
    return libro


def _aplicar(libro, identidad=None, idioma="es", **extra):
    return aplicar_identidad_a_salida(
        libro, identidad or IDENTIDAD_B, hojas_busqueda=["Resumen"],
        hoja_resumen="Resumen", idioma=idioma, **extra)


# ---------------------------------------------------------------------------
# Tablas Excel registradas: estructura formal por encima del texto
# ---------------------------------------------------------------------------


def test_n17_tabla_formal_poblada_preserva_cabeceras_y_no_publica(tmp_path):
    libro = _libro_base()
    resumen = libro["Resumen"]
    resumen["A2"], resumen["B2"] = "Proyecto", "Color"   # fuera de vocabulario
    resumen["A3"], resumen["B3"] = "Proyecto B", "rojo"
    resumen.add_table(Table(displayName="TablaPoblada", ref="A2:B3"))
    temporal = tmp_path / "n17_poblada.xlsx"
    libro.save(temporal)
    releido = load_workbook(temporal)
    meta = _aplicar(releido)
    hoja = releido["Resumen"]
    assert (hoja["A2"].value, hoja["B2"].value) == ("Proyecto", "Color")
    assert (hoja["A3"].value, hoja["B3"].value) == ("Proyecto B", "rojo")
    assert meta["celdas"] == []
    assert meta["publicable"] is False        # sin campo verificable (P08)
    assert not [c for c in meta["contradicciones"] if c.get("estructural")]
    assert hoja.tables["TablaPoblada"].ref == "A2:B3"


def test_n17_varias_tablas_en_posiciones_arbitrarias_preservadas(tmp_path):
    libro = _libro_base()
    ficha = libro.create_sheet("Ficha")
    ficha["D5"], ficha["E5"] = "Proyecto", "Humedad"
    ficha.add_table(Table(displayName="TablaUno", ref="D5:E5"))
    ficha["B9"], ficha["C9"] = "Project", "Size"
    ficha.add_table(Table(displayName="TablaDos", ref="B9:C9"))
    temporal = tmp_path / "n17_varias.xlsx"
    libro.save(temporal)
    releido = load_workbook(temporal)
    meta = _aplicar(releido)
    hoja = releido["Ficha"]
    assert (hoja["D5"].value, hoja["E5"].value) == ("Proyecto", "Humedad")
    assert (hoja["B9"].value, hoja["C9"].value) == ("Project", "Size")
    assert hoja.tables["TablaUno"].ref == "D5:E5"
    assert hoja.tables["TablaDos"].ref == "B9:C9"
    assert meta["celdas_complementarias"] == []
    assert not [c for c in meta["contradicciones"] if c.get("estructural")]


def test_n17_autofiltro_protege_cabecera_y_permite_publicar_con_explicito():
    libro = _libro_base()
    resumen = libro["Resumen"]
    resumen["A1"], resumen["B1"] = "Proyecto:", "Proyecto A"   # campo válido
    resumen["A2"], resumen["B2"] = "Proyecto", "Temperatura"   # cabecera
    resumen["A3"], resumen["B3"] = "Proyecto B", 21.5
    resumen["A4"], resumen["B4"] = "Proyecto B", 22.0
    resumen.auto_filter.ref = "A2:B4"
    meta = _aplicar(libro)
    assert resumen["B1"].value == "Proyecto B"     # campo explícito mapeado
    assert resumen["B2"].value == "Temperatura"    # cabecera protegida
    assert meta["publicable"] is True              # sin contradicción falsa
    assert not any(c["celda_valor"] == "B2" for c in meta["celdas"])


def test_n17_etiqueta_explicita_en_cabecera_de_tabla_tampoco_se_sustituye():
    """«Ninguna cabecera de tabla declarada se sustituye, importe el texto»."""
    libro = _libro_base()
    resumen = libro["Resumen"]
    resumen["A2"], resumen["B2"] = "Proyecto:", "X-9"
    resumen.add_table(Table(displayName="TablaExplicita", ref="A2:B2"))
    meta = _aplicar(libro)
    assert resumen["B2"].value == "X-9"            # intacta pese al «:»
    assert meta["celdas"] == []
    assert meta["publicable"] is False


def test_n17_par_explicito_en_cuerpo_de_tabla_se_mapea():
    """El cuerpo (filas de datos) no es cabecera: rige la política normal."""
    libro = _libro_base()
    resumen = libro["Resumen"]
    resumen["A2"], resumen["B2"] = "Caso", "Valor"
    resumen["A3"], resumen["B3"] = "Proyecto:", "Proyecto A"
    resumen.add_table(Table(displayName="TablaCuerpo", ref="A2:B3"))
    meta = _aplicar(libro)
    assert (resumen["A2"].value, resumen["B2"].value) == ("Caso", "Valor")
    assert resumen["B3"].value == "Proyecto B"     # cuerpo: contrato explícito
    assert [c["celda_valor"] for c in meta["celdas"]] == ["B3"]
    assert meta["publicable"] is True


# ---------------------------------------------------------------------------
# Sin tabla formal: evidencia positiva o fail-closed
# ---------------------------------------------------------------------------


def test_n17_tabla_visual_desconocida_con_campo_explicito_bloquea_fail_closed():
    """Visual «Proyecto|Temperatura» sin Table: ambigua, bloqueo fail-closed.

    A diferencia de la tabla REGISTRADA (omisión silenciosa), sin estructura
    formal no se puede probar que sea cabecera: se bloquea la publicación
    señalando la celda (revisión requerida), aunque exista otro campo válido.
    """
    libro = _libro_base()
    resumen = libro["Resumen"]
    resumen["A1"], resumen["B1"] = "Proyecto:", "Proyecto A"   # campo válido
    resumen["A3"], resumen["B3"] = "Proyecto", "Temperatura"   # sin Table
    meta = _aplicar(libro)
    assert resumen["B1"].value == "Proyecto B"
    assert resumen["B3"].value == "Temperatura"    # nada se escribe
    assert meta["publicable"] is False
    bloqueos = [c for c in meta["contradicciones"] if c.get("estructural")]
    assert bloqueos and bloqueos[0]["celda"] == "B3"


@pytest.mark.parametrize("etiqueta,valor", [
    ("Proyecto", "Proyecto A"),        # mención inicial (histórico/migración)
    ("Project", "Project North"),      # mención inicial EN
    ("Proyecto", "Proyecto B"),        # ya coincide con la identidad
    ("Proyecto", "proyecto-b"),        # coincide con el project_id
])
def test_n17_evidencia_positiva_del_contiguo_si_se_mapea(etiqueta, valor):
    libro = _libro_base()
    hoja = libro.create_sheet("Ficha")
    hoja["A1"], hoja["B1"] = etiqueta, valor
    meta = _aplicar(libro)
    assert hoja["B1"].value == "Proyecto B"
    assert meta["publicable"] is True
    assert any(c["hoja"] == "Ficha" and c["valor_anterior"] == valor
               for c in meta["celdas_complementarias"])


def test_n17_valor_vacio_contiguo_se_rellena():
    """Campo vacío junto a etiqueta desnuda: se rellena (comportamiento N10)."""
    libro = _libro_base()
    libro["Resumen"]["A2"], libro["Resumen"]["B2"] = "Proyecto:", "Proyecto A"
    hoja = libro.create_sheet("Ficha")
    hoja["A1"] = "Proyecto"            # contiguo vacío, sin datos debajo
    meta = _aplicar(libro)
    assert hoja["B1"].value == "Proyecto B"
    assert meta["publicable"] is True


def test_n17_colision_proyecto_llamado_temperatura():
    """Identidad coincidente con palabra de cabecera: sólo vía evidencia."""
    # Desnuda: el contiguo coincide EXACTAMENTE con la identidad declarada
    # (evidencia positiva) ⇒ coherente, se publica sin alterar nada.
    libro = _libro_base()
    libro["Resumen"]["A2"], libro["Resumen"]["B2"] = "Proyecto", "Temperatura"
    meta = _aplicar(libro, identidad=IDENTIDAD_TEMPERATURA)
    assert libro["Resumen"]["B2"].value == "Temperatura"
    assert meta["publicable"] is True
    # Explícita: contrato inequívoco ⇒ mapea y publica.
    libro2 = _libro_base()
    libro2["Resumen"]["A2"], libro2["Resumen"]["B2"] = "Proyecto:", "viejo"
    meta2 = _aplicar(libro2, identidad=IDENTIDAD_TEMPERATURA)
    assert libro2["Resumen"]["B2"].value == "Temperatura"
    assert meta2["publicable"] is True


@pytest.mark.parametrize("idioma,etiqueta", [
    ("es", "Proyecto:"), ("en", "Project:"), ("de", "Projekt:"),
])
def test_n14_campo_explicito_en_tres_idiomas_se_mapea(idioma, etiqueta):
    """Cobertura positiva ES/EN/DE conservada tras el SUPERSEDE de N17."""
    libro = _libro_base()
    hoja = libro.create_sheet("Ficha")
    hoja["A1"], hoja["B1"] = etiqueta, "Old"
    meta = _aplicar(libro, idioma=idioma)
    assert hoja["B1"].value == nombre_visible(IDENTIDAD_B, idioma)
    assert meta["publicable"] is True, idioma


# ---------------------------------------------------------------------------
# Exportadores reales (copias temporales + hashes + releídos)
# ---------------------------------------------------------------------------


def test_n17_empresarial_tabla_formal_vacia_genera_borrador_intacto(tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = tmp_path / "base_n17.xlsx"
    _plantilla(plantilla)
    libro = load_workbook(plantilla)
    resumen = libro["Resumen"]
    resumen["H1"], resumen["I1"] = "Proyecto", "Temperatura"
    resumen.add_table(Table(displayName="TablaVacia", ref="H1:I1"))
    libro.save(plantilla)
    hash_plantilla = _hash(plantilla)
    oficial = tmp_path / "salida_n17.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla, oficial, [], identidad_proyecto=IDENTIDAD_B))
    assert not oficial.exists()
    assert "borrador_no_publicable" in resultado.name
    borrador = load_workbook(resultado)          # releer el Excel temporal
    assert (borrador["Resumen"]["H1"].value, borrador["Resumen"]["I1"].value) \
        == ("Proyecto", "Temperatura")
    assert borrador["Resumen"].tables["TablaVacia"].ref == "H1:I1"
    assert _hash(plantilla) == hash_plantilla    # plantilla original intacta


def test_n17_empresarial_tabla_formal_con_campo_explicito_publica(tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = tmp_path / "base_n17b.xlsx"
    _plantilla(plantilla)
    libro = load_workbook(plantilla)
    resumen = libro["Resumen"]
    resumen["H1"], resumen["I1"] = "Proyecto", "Temperatura"
    resumen.add_table(Table(displayName="TablaVacia", ref="H1:I1"))
    resumen["A2"], resumen["B2"] = "Proyecto:", "Proyecto A"
    libro["Captura_pruebas"]["A1"] = "Reporte del Proyecto B"
    libro.save(plantilla)
    hash_plantilla = _hash(plantilla)
    oficial = tmp_path / "salida_n17b.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla, oficial, [], identidad_proyecto=IDENTIDAD_B))
    assert resultado == oficial and oficial.is_file()
    salida = load_workbook(oficial)
    assert (salida["Resumen"]["H1"].value, salida["Resumen"]["I1"].value) \
        == ("Proyecto", "Temperatura")           # cabecera intacta
    assert salida["Resumen"].tables["TablaVacia"].ref == "H1:I1"
    assert salida["Resumen"]["B2"].value == "Proyecto B"
    assert _hash(plantilla) == hash_plantilla


def test_n17_bmg_real_tabla_formal_publica_con_seleccion_usuario_intacta(tmp_path):
    from reporte_bmg import generar_reporte_bmg
    from configuracion import cargar_config, cargar_reglas
    from test_reporte_bmg_v7 import caso
    real = RAIZ / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    hash_real = _hash(real)
    copia = tmp_path / "bmg_n17.xlsx"
    copia.write_bytes(real.read_bytes())
    libro = load_workbook(copia)
    rangos = libro["Rangos"]
    rangos["A1"], rangos["B1"] = "Projekt", "Temperatur"
    rangos.add_table(Table(displayName="TablaRangos", ref="A1:B1"))
    libro.save(copia)
    oficial = tmp_path / "bmg_oficial.xlsx"
    resultados = [caso(1001, "1ST", "HT", "UGL", 5.8)]
    salida = pathlib.Path(generar_reporte_bmg(
        copia, oficial, resultados, cargar_config(), cargar_reglas(),
        idioma="de", identidad_proyecto=IDENTIDAD_B))
    # V15.10.1: una selección explícita del proyecto puede publicar cuando la
    # plantilla carece de una celda de identidad utilizable y no hay contradicción.
    assert oficial.exists()
    assert salida == oficial
    publicado = load_workbook(salida)
    assert (publicado["Rangos"]["A1"].value, publicado["Rangos"]["B1"].value) \
        == ("Projekt", "Temperatur")
    assert publicado["Rangos"].tables["TablaRangos"].ref == "A1:B1"
    assert _hash(real) == hash_real              # plantilla del repo intacta
