"""
verificar_excel.py — Lee resultado_maestro.xlsx y verifica hojas, valores y
colores del semáforo (evidencia para docs/fase3_excel.md).
"""
from openpyxl import load_workbook

wb = load_workbook("resultado_maestro.xlsx")
print("Hojas:", wb.sheetnames)

ws = wb["Estructura completa"]
print("\n=== Estructura completa (encabezados) ===")
print([c.value for c in ws[1]])
for fila in ws.iter_rows(min_row=2, values_only=True):
    print(" | ".join(str(v)[:22] for v in fila))

wm = wb["Matriz de cumplimiento"]
print("\n=== Matriz de cumplimiento ===")
for fila in wm.iter_rows(min_row=1, max_row=wm.max_row):
    valores = [c.value for c in fila]
    semaforo = fila[4]
    rgb = semaforo.fill.start_color.rgb if semaforo.fill and semaforo.fill.fill_type == "solid" else None
    print(f"{valores[1]:16} conf={str(valores[2]):7} coinc={str(valores[3]):20} "
          f"semaforo={str(valores[4]):14} fill={rgb}")

print("\nReglas condicionales en matriz:")
for rango, reglas in wm.conditional_formatting._cf_rules.items():
    for r in reglas:
        print(f"  {rango.sqref} -> {type(r).__name__}", getattr(r, "operator", ""),
              getattr(r, "formula", ""))
