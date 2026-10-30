"""V15.9.1 · B3: paridad de la matriz entre dashboard, Excel y Listado.

Reproduce el bloqueante de la auditoría: un caso ``tipo_st='3RD'`` con
temperatura HT y generador UGL aparecía en ``sin_clasificar`` del dashboard
pero NO en la hoja Resumen del Excel. Ahora existe UN bloque único global de
casos fuera de etapa/categorización (3RD, LEGACY, sin etapa, combinaciones
fuera de columnas) alimentado por la MISMA agrupación que el dashboard.
"""
from pathlib import Path

from openpyxl import load_workbook

import reporte_bmg
from configuracion import cargar_config, cargar_reglas
from dashboard.backend.app import _fila_publica, _matriz_apertura
from utilidades.matriz_apertura import agrupar_casos, causa_sin_clasificar

RAIZ = Path(__file__).resolve().parents[1]
PLANTILLA = RAIZ / "plantillas" / "reporte_bmg_1st_2st.xlsx"


def fila(numero, tipo_st, temp, inflador, apertura):
    return {
        "ruta": f"/tmp/{tipo_st or 'sin'}/{numero}", "nombre": str(numero),
        "identificador": str(numero), "test_number": str(numero),
        "tipo_st": tipo_st, "temperature_condition": temp,
        "inflator_type": inflador,
        "comparacion": {"resultado": "coincidencia_total"},
        "confianza_ocr_pct": 95.0, "campos": {
            "opening_time_ms": {"valor": apertura, "estado": "confirmado",
                                "requiere_revision": False},
            "test_number": {"valor": str(numero)},
        },
    }


CASOS = [
    fila(101, "1ST", "HT", "UGL", 5.8),     # dentro de matriz
    fila(102, "2ST", "NT", "NOM", 6.4),     # dentro de matriz
    fila(103, "3RD", "HT", "UGL", 5.5),     # B3: caso de la auditoría
    fila(104, "LEGACY", "RT", "NOM", 6.9),  # etapa fuera de matriz
    fila(105, None, "HT", "UGL", 6.1),      # sin etapa declarada
    fila(106, "1ST", "XX", "UGL", 5.9),     # combinación fuera de columnas
    fila(107, "2ST", "RT", None, 6.2),      # sin generador
]

FUERA_ESPERADOS = {"103", "104", "105", "106", "107"}
DENTRO_ESPERADOS = {"101", "102"}


def _ids_dashboard(matriz):
    dentro = set()
    for etapa in matriz["etapas"]:
        for columna in etapa["columnas"]:
            for celda in columna["celdas"]:
                if celda["tipo"] == "caso":
                    dentro.add(str(celda["test_number"]))
    fuera = {str(c["test_number"]) for c in matriz["sin_clasificar"]}
    return dentro, fuera


def test_dashboard_clasifica_igual_que_el_motor_compartido():
    publicas = [_fila_publica(dict(f)) for f in CASOS]
    matriz = _matriz_apertura(publicas)
    dentro, fuera = _ids_dashboard(matriz)
    assert dentro == DENTRO_ESPERADOS
    assert fuera == FUERA_ESPERADOS
    # La causa es trazable en el dashboard.
    causas = {str(c["test_number"]): c["causa"] for c in matriz["sin_clasificar"]}
    assert "3RD" in causas["103"]
    assert "LEGACY" in causas["104"]
    assert "etapa" in causas["105"].lower()
    assert "XX" in causas["106"]
    assert "generador" in causas["107"].lower()


def test_excel_resume_incluye_bloque_unico_sin_clasificar(tmp_path):
    salida = tmp_path / "matriz_b3.xlsx"
    reporte_bmg.generar_reporte_bmg(PLANTILLA, salida, [dict(f) for f in CASOS],
                                    cargar_config(), cargar_reglas(), idioma="es")
    libro = load_workbook(salida)
    resumen = libro[reporte_bmg.HOJA_RESUMEN]
    titulo = reporte_bmg.TEXTOS["es"]["unclassified"]
    # Localiza el bloque único global (ya no hay bloques intra-etapa).
    celdas_titulo = [celda.row for fila_ws in resumen.iter_rows(min_col=2, max_col=2)
                     for celda in fila_ws if celda.value == titulo]
    assert len(celdas_titulo) == 1, f"Debe existir UN solo bloque '{titulo}'"
    fila_bloque = celdas_titulo[0]
    # El conteo acompaña al título.
    assert "5" in str(resumen.cell(fila_bloque, 3).value)
    # Fila de encabezados y luego los casos ordenados por número.
    ids_excel, causas_excel = [], {}
    r = fila_bloque + 2
    while resumen.cell(r, 2).value not in (None, ""):
        ids_excel.append(str(resumen.cell(r, 2).value))
        causas_excel[str(resumen.cell(r, 2).value)] = str(resumen.cell(r, 6).value)
        r += 1
    assert set(ids_excel) == FUERA_ESPERADOS
    assert "3RD" in causas_excel["103"]
    assert "LEGACY" in causas_excel["104"]
    # El ID real y el vínculo al Listado existen (nunca IDs inventados).
    assert resumen.cell(fila_bloque + 2, 2).hyperlink is not None
    # Los casos dentro de matriz NO están en el bloque sin clasificar.
    assert not (DENTRO_ESPERADOS & set(ids_excel))


def test_paridad_total_ids_dashboard_excel_listado(tmp_path):
    salida = tmp_path / "matriz_paridad.xlsx"
    reporte_bmg.generar_reporte_bmg(PLANTILLA, salida, [dict(f) for f in CASOS],
                                    cargar_config(), cargar_reglas(), idioma="es")
    libro = load_workbook(salida)
    # IDs del dashboard.
    publicas = [_fila_publica(dict(f)) for f in CASOS]
    matriz = _matriz_apertura(publicas)
    dentro_dash, fuera_dash = _ids_dashboard(matriz)
    ids_dashboard = dentro_dash | fuera_dash
    # IDs del Listado (columna C, destino de los vínculos de la matriz).
    lista = libro[reporte_bmg.HOJA_LISTA]
    ids_listado = {str(c.value) for fila_ws in lista.iter_rows(min_col=3, max_col=3)
                   for c in fila_ws if c.value not in (None, "")}
    assert ids_dashboard <= ids_listado, "Todo caso del dashboard debe estar en el Listado"
    # IDs del bloque Excel «Sin clasificar» == sin_clasificar del dashboard.
    resumen = libro[reporte_bmg.HOJA_RESUMEN]
    titulo = reporte_bmg.TEXTOS["es"]["unclassified"]
    fila_bloque = next(celda.row for fila_ws in resumen.iter_rows(min_col=2, max_col=2)
                       for celda in fila_ws if celda.value == titulo)
    ids_excel_fuera = set()
    r = fila_bloque + 2
    while resumen.cell(r, 2).value not in (None, ""):
        ids_excel_fuera.add(str(resumen.cell(r, 2).value))
        r += 1
    assert ids_excel_fuera == fuera_dash
    assert ids_dashboard == DENTRO_ESPERADOS | FUERA_ESPERADOS


def test_causa_compartida_entre_modulos():
    bruta = fila(200, "3RD", "HT", "UGL", 5.0)
    assert causa_sin_clasificar(bruta) == "Etapa fuera de matriz (3RD)"
    por_columna, fuera = agrupar_casos([bruta])
    assert len(fuera) == 1 and all(len(v) == 0 for v in por_columna.values())


def test_sin_casos_fuera_no_aparece_bloque(tmp_path):
    salida = tmp_path / "matriz_limpia.xlsx"
    casos = [dict(fila(301, "1ST", "HT", "UGL", 5.7))]
    reporte_bmg.generar_reporte_bmg(PLANTILLA, salida, casos,
                                    cargar_config(), cargar_reglas(), idioma="es")
    libro = load_workbook(salida)
    resumen = libro[reporte_bmg.HOJA_RESUMEN]
    titulo = reporte_bmg.TEXTOS["es"]["unclassified"]
    apariciones = [celda.row for fila_ws in resumen.iter_rows(min_col=2, max_col=2)
                   for celda in fila_ws if celda.value == titulo]
    assert apariciones == []


# ------------------------------------------------ B1/B2 end-to-end hasta el Excel


def _fila_con_videos(numero, tipo_st, temp, inflador, items):
    return {
        "ruta": f"/tmp/{tipo_st}/{numero}", "nombre": str(numero),
        "identificador": str(numero), "test_number": str(numero),
        "tipo_st": tipo_st, "temperature_condition": temp, "inflator_type": inflador,
        "comparacion": {"resultado": "coincidencia_total"}, "confianza_ocr_pct": 95.0,
        "campos": {"test_number": {"valor": str(numero)}},
        "videos": items,
    }


def _item_camara(nombre, tiempo, **extra):
    it = {"nombre": nombre, "estado": "candidato" if tiempo is not None else "no_event",
          "tiempo_ms": tiempo, "tiempo_fps_ms": tiempo, "resolucion_nativa_ms": 0.1,
          "precision_fuente_ms": 0.1, "fuente_tiempo": "fps_subframe"}
    it.update(extra)
    return it


def _celda_del_numero(resumen, numero):
    for fila in resumen.iter_rows():
        for c in fila:
            if str(c.value) == str(numero):
                return c
    return None


def test_e2e_propuesto_y_causa_grave_quedan_grises_en_el_excel(tmp_path):
    import importlib
    backend = importlib.import_module("dashboard.backend.app")
    # Caso 1: consenso PROPUESTO (2 de 3 coincidentes) → B1.
    fila_propuesta = _fila_con_videos(111, "1ST", "HT", "UGL", [
        _item_camara("09_DETAIL.AVI", 5000.0), _item_camara("09_FRONTAL.AVI", 5000.05),
        _item_camara("09_LATERAL.AVI", None)])
    backend._aplicar_consenso_video_a_fila(
        fila_propuesta, __import__("video_airbag").consenso_videos(
            fila_propuesta["videos"], {}))
    # Caso 2: 3 coincidentes pero una con FPS no corroborados → B2.
    fila_grave = _fila_con_videos(112, "1ST", "RT", "NOM", [
        _item_camara("09_DETAIL.AVI", 5000.0, advertencia_fps="12 fps declarados"),
        _item_camara("09_FRONTAL.AVI", 5000.05), _item_camara("09_LATERAL.AVI", 5000.1)])
    backend._aplicar_consenso_video_a_fila(
        fila_grave, __import__("video_airbag").consenso_videos(fila_grave["videos"], {}))
    # Caso 3: confirmado limpio → se colorea.
    fila_ok = _fila_con_videos(113, "2ST", "HT", "UGL", [
        _item_camara("09_DETAIL.AVI", 5000.0), _item_camara("09_FRONTAL.AVI", 5000.05),
        _item_camara("09_LATERAL.AVI", 5000.1)])
    backend._aplicar_consenso_video_a_fila(
        fila_ok, __import__("video_airbag").consenso_videos(fila_ok["videos"], {}))

    salida = tmp_path / "consenso_e2e.xlsx"
    reporte_bmg.generar_reporte_bmg(
        PLANTILLA, salida, [dict(f) for f in (fila_propuesta, fila_grave, fila_ok)],
        cargar_config(), cargar_reglas(), idioma="es")
    resumen = load_workbook(salida)[reporte_bmg.HOJA_RESUMEN]
    config = cargar_config()
    colores = config.get("colores", {})

    def rgb(celda):
        return str(celda.fill.fgColor.rgb or "").upper()

    gris = (colores.get("gris") or "#9ca3af").lstrip("#").upper()
    celda_propuesta = _celda_del_numero(resumen, 111)
    assert celda_propuesta is not None, "El caso propuesto debe aparecer en la matriz"
    assert rgb(celda_propuesta).endswith(gris), \
        "B1: la celda propuesta debe ser GRIS en el Excel"
    comentario = str(celda_propuesta.comment.text if celda_propuesta.comment else "")
    assert "Pendiente de revisión" in comentario
    assert "5000" in comentario  # la medida pendiente es visible, nunca coloreada

    celda_grave = _celda_del_numero(resumen, 112)
    assert rgb(celda_grave).endswith(gris), \
        "B2: la causa grave debe dejar la celda GRIS en el Excel"
    comentario_grave = str(celda_grave.comment.text if celda_grave.comment else "")
    assert "Pendiente de revisión" in comentario_grave

    celda_ok = _celda_del_numero(resumen, 113)
    rojo = (colores.get("rojo") or "#FF0000").lstrip("#").upper()
    assert rgb(celda_ok).endswith(rojo), \
        "El consenso confirmado limpio sí se colorea (5000 ms > 7 ms → rojo)"
