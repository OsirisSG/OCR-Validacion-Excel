from pathlib import Path

from configuracion import cargar_config, cargar_reglas
from flujo_empresarial import consolidar_caso, extraer_candidatos_texto
from reporte_bmg import generar_reporte_bmg
from utilidades.cola_excel import ColaExcel


def _caso():
    return {
        "project_id": "proyecto-a", "case_key": "proyecto-a|1ST|232541",
        "tipo_st": "1ST", "ruta_relativa": "1ST/232541", "ruta": "/tmp/232541",
        "metadata_ruta": {"test_number": "232541", "temperature_condition": "HT",
                          "inflator_type": "UGL"},
    }


def test_opening_automatico_queda_candidato_pero_no_se_consolida():
    caso = _caso()
    contexto = {"project_id": caso["project_id"], "case_key": caso["case_key"],
                "tipo_st": "1ST", "test_number": "232541", "version_procesamiento": "v8"}
    evidencias = extraer_candidatos_texto(
        "OPENING 5.8 ms", {"ruta_relativa": "IMG_1.jpg", "confianza_ocr": .95}, contexto)
    apertura = next(e for e in evidencias if e["clave"] == "opening_time_ms")
    assert apertura["requiere_confirmacion_manual"] is True
    consolidado = consolidar_caso(caso, evidencias, requeridos={"opening_time_ms"})
    assert consolidado["campos"]["opening_time_ms"]["valor"] is None


def test_cola_expone_estado_persistencia_sin_romper_estado_legacy(tmp_path):
    cola = ColaExcel(tmp_path / "cola.sqlite3", lambda: "maestro.xlsx")
    assert cola.estado()["estado_persistencia"] == "guardado"
    cola.solicitar()
    assert cola.estado()["estado"] == "pendiente"
    assert cola.estado()["estado_persistencia"] == "pendiente"
    cola.procesar()
    assert cola.estado()["estado_persistencia"] == "guardado"


def test_reporte_incluye_diccionario_y_trazabilidad_multilingue(tmp_path):
    raiz = Path(__file__).resolve().parents[1]
    plantilla = raiz / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    salida = tmp_path / "v8_en.xlsx"
    fila = {
        "project_id": "P", "case_key": "P|1ST|1", "tipo_st": "1ST",
        "test_number": "1", "temperature_condition": "HT", "inflator_type": "UGL",
        "campos": {"test_number": {"valor": "1"}, "temperature_condition": {"valor": "HT"},
                   "inflator_type": {"valor": "UGL"}, "opening_time_ms": {"valor": None}},
        "candidatos": [{"proyecto":"P","etapa":"1ST","caso":"P|1ST|1",
                        "componente":"bolsa_aire","clave":"module_part_number",
                        "imagen":"IMG_1.jpg","imagen_absoluta":"/tmp/IMG_1.jpg",
                        "bbox_original":[1,2,3,4],"valor_original":"2GJ 880 204 J",
                        "valor_normalizado":"2GJ.880.204.J","orientacion_grados":0,
                        "regla_aplicada":"numero_parte","fuente":"ocr","version_procesamiento":"v8"}],
    }
    generar_reporte_bmg(plantilla, salida, [fila], cargar_config(), cargar_reglas(), idioma="en")
    from openpyxl import load_workbook
    libro = load_workbook(salida, read_only=False)
    assert "Traceability" in libro.sheetnames
    assert "Data dictionary" in libro.sheetnames
    dic = libro["Data dictionary"]
    assert dic["A1"].value == "Field"
    filas = {dic.cell(r, 1).value: dic.cell(r, 3).value for r in range(2, dic.max_row + 1)}
    assert "opening_time_ms" in filas
    assert "manual" in filas["opening_time_ms"].lower()


def test_frontend_conserva_v6_y_separa_correccion_de_entrenamiento():
    raiz = Path(__file__).resolve().parents[1]
    js = (raiz / "dashboard" / "frontend" / "app.js").read_text(encoding="utf-8")
    assert "variantes-siempre-visibles" in js
    assert "Comparar lecturas de esta zona" not in js
    assert 'useState("guardar_sin_entrenar")' in js
    assert "versionBorradorRef" in js
    assert "idiomas_exportacion: idiomasExportacion" in js
