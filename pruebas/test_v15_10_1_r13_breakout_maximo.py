from __future__ import annotations

import math

import numpy as np

from utilidades.estandar_apertura import (
    analizar_secuencia_despliegue,
    evaluar_evidencia_estandar,
)
from utilidades.matriz_apertura import celda_caso, criterio_desde_limites
from utilidades.resultado_caso import resolver_apertura
from video_airbag import _metricas_primer_blanco


def _m(area=0, ruptura=0, blanco=0, grieta=0, borde=0.0, deform=None, prox=None):
    return {
        "area_bolsa_visible_estimada_px": area,
        "pixeles_ruptura": ruptura,
        "pixeles_blanco_nuevo": blanco,
        "pixeles_grieta_nueva": grieta,
        "componente_ruptura_px": ruptura,
        "componente_blanco_px": max(ruptura, blanco),
        "min_pixeles_debil": 1,
        "min_componente_debil": 1,
        "fraccion_borde_nuevo": borde,
        "deformacion_confiable": deform,
        "proximidad_deformacion": prox,
        "scores_fuentes": {"bordes_nuevos": borde, "reflejo": 0.0, "abombamiento": 0.0},
    }


def test_secuencia_r13_breakout_y_maximo_a_5000fps():
    metricas = [_m() for _ in range(5)] + [
        _m(area=1, ruptura=1, blanco=1, borde=0.002, deform=True, prox=True),
        _m(area=8, ruptura=4, blanco=4, borde=0.003, deform=True, prox=True),
        _m(area=25, ruptura=8, blanco=8, borde=0.003, deform=True, prox=True),
        _m(area=60, ruptura=10, blanco=10, borde=0.003, deform=True, prox=True),
        _m(area=100, ruptura=12, blanco=12, borde=0.003, deform=True, prox=True),
        _m(area=99, ruptura=12, blanco=12, borde=0.003, deform=True, prox=True),
        _m(area=97, ruptura=11, blanco=11, borde=0.003, deform=True, prox=True),
        _m(area=85, ruptura=10, blanco=10, borde=0.002, deform=True, prox=True),
    ]
    r = analizar_secuencia_despliegue(metricas, 5, 5000)
    assert r["frame_breakout"] == 5
    assert r["frame_max_inflado"] in {9, 10}
    assert math.isclose(r["tiempo_breakout_ms"], 1.0, abs_tol=1e-6)
    assert r["duracion_expansion_ms"] in {0.8, 1.0}
    assert r["tiempo_relativo_al_inicio_video"] is True
    assert r["trigger_electrico_disponible"] is False


def test_maximo_en_borde_clip_se_reporta_como_anomalia():
    metricas = [_m() for _ in range(3)] + [_m(area=x, ruptura=max(1, x//10), blanco=max(1, x//10)) for x in (2, 5, 15, 30, 60)]
    r = analizar_secuencia_despliegue(metricas, 3, 5000)
    assert r["frame_max_inflado"] is not None
    assert "maximo_en_borde_final_del_clip" in r["anomalías"]
    assert r["confianza_maximo"] < 1.0


def test_un_video_defendible_pinta_sin_confirmacion_humana():
    evaluacion = {
        "auto_colorable": True,
        "score_evidencia": 0.82,
        "nivel": "alta",
        "puertas_decision": {
            "onset_local_nuevo": True,
            "previo_limpio": True,
            "continuidad_posterior": True,
            "crecimiento_o_inicio_fuerte": True,
            "ruta_fisica_defendible": True,
            "sin_riesgo_grave": True,
        },
        "puertas_fallidas": [],
        "riesgos": [],
        "razones": ["breakout_defendible"],
        "version_estandar": "apertura-estandar-r13-v5",
    }
    fila = {
        "id": "c1", "test_number": "1", "tipo_st": "1ST",
        "temperature_condition": "HT", "inflator_type": "UGL",
        "estado_proceso": "procesada",
        "videos": [{
            "id": "cam1", "estado": "candidato", "frame_evento": 29,
            "evaluacion_apertura_estandar": evaluacion,
            "tiempo_color_estimado_ms": 5.8,
            "referencia_color_valida": True,
            "resolucion_nativa_ms": 0.2,
            "fuente_tiempo_color": "overlay_rec",
            "secuencia_despliegue_r13": {
                # R15.1 (MEDIO-09): la mediana confirmada del caso exige el
                # estado explícito de la secuencia; fixture actualizada.
                "estado": "maximo_inflado_estimado",
                "frame_breakout": 29, "frame_max_inflado": 70,
                "tiempo_max_inflado_ms": 14.0, "duracion_expansion_ms": 8.2,
                "confianza_maximo": 0.9, "anomalías": [],
            },
        }],
    }
    apertura = resolver_apertura(fila)
    assert apertura["estado"] == "PROPUESTA"
    assert apertura["auto_color_elegible"] is True
    assert apertura["tiempo_estimado_auto_ms"] == 5.8
    assert apertura["tiempo_max_inflado_estimado_ms"] == 14.0
    assert apertura["duracion_hasta_max_inflado_ms"] == 8.2
    celda = celda_caso(fila, criterio_desde_limites(6, 7))
    assert celda["estado"] == "verde"
    assert celda["tiempo_es_estimado"] is True


def test_conflicto_entre_dos_videos_defendibles_se_abstiene_de_color():
    def vid(nombre, tiempo):
        return {
            "id": nombre, "estado": "candidato", "frame_evento": 1,
            "evaluacion_apertura_estandar": {
                "auto_colorable": True, "score_evidencia": 0.9, "nivel": "alta",
                "puertas_decision": {"ruta_fisica_defendible": True},
                "puertas_fallidas": [], "riesgos": [], "razones": [],
            },
            "tiempo_color_estimado_ms": tiempo,
            "referencia_color_valida": True,
            "resolucion_nativa_ms": 0.2,
            "fuente_tiempo_color": "overlay_rec",
        }
    fila = {
        "id":"x", "test_number":"1", "tipo_st":"1ST",
        "temperature_condition":"HT", "inflator_type":"UGL",
        "videos":[vid("a",5.8), vid("b",7.2)],
    }
    apertura = resolver_apertura(fila)
    assert apertura["auto_color_conflicto"] is True
    assert apertura["auto_color_elegible"] is False
    celda = celda_caso(fila, criterio_desde_limites(6,7))
    assert celda["estado"] == "gris"


def test_ruta_temporal_sin_dic_puede_ser_alta_si_crece_y_tiene_dos_senales():
    cand = {
        "score_candidato": 0.6,
        "continuidad_frames": 3,
        "frames_previos_limpios": 2,
        "tipo_senales": ["material_claro_nuevo", "borde_local_nuevo"],
        "razones_riesgo": [],
    }
    metrica = _m(area=1, ruptura=1, blanco=1, borde=0.003, deform=False, prox=None)
    ventana = [metrica, _m(area=4, ruptura=4, blanco=4, borde=0.003), _m(area=8, ruptura=8, blanco=8, borde=0.003)]
    r = evaluar_evidencia_estandar(cand, {"razones_en_contra": []}, metrica, metricas_ventana=ventana)
    assert r["ruta_defendible"] in {"trayectoria_temporal", "fisura_a_bolsa"}
    assert r["defendible_un_video"] is True
    assert r["auto_colorable"] is True


def test_pixel_aislado_sin_crecimiento_no_es_auto_colorable():
    cand = {
        "score_candidato": 0.95,
        "continuidad_frames": 3,
        "frames_previos_limpios": 2,
        "tipo_senales": ["material_claro_nuevo"],
        "razones_riesgo": [],
    }
    metrica = _m(area=1, ruptura=1, blanco=1, borde=0.0, deform=False, prox=None)
    r = evaluar_evidencia_estandar(cand, {"razones_en_contra": []}, metrica, metricas_ventana=[metrica,metrica,metrica])
    assert r["auto_colorable"] is False
    assert "crecimiento_o_inicio_fuerte" in r["puertas_fallidas"]


def test_area_bolsa_amplia_crece_aunque_breakout_use_contorno_local():
    base = np.full((80, 120), 60, dtype=np.uint8)
    actual1 = base.copy(); actual1[38:43, 56:61] = 210
    actual2 = base.copy(); actual2[30:55, 45:80] = 210
    cfg = {
        "umbral_luminancia_blanco": 135,
        "umbral_incremento_blanco": 18,
        "umbral_luminancia_bolsa": 120,
        "umbral_incremento_bolsa": 12,
        "umbral_incremento_bolsa_max": 42,
        "umbral_luminancia_grieta": 120,
        "umbral_discontinuidad_nueva": 6.0,
        "residuo_min_compensado": 18,
        "compensar_movimiento": False,
        "min_fraccion_blanco_nuevo": 0.00003,
        "min_fraccion_componente_blanco": 0.000015,
        "min_pixeles_indicio": 1,
        "max_componentes_indicio": 64,
        "radio_indicio_px": 3.0,
        "estandar_apertura_activo": True,
        "umbral_iluminacion_global": 25.0,
    }
    m1 = _metricas_primer_blanco(actual1, base, cfg)
    m2 = _metricas_primer_blanco(actual2, base, cfg)
    assert m1["area_bolsa_visible_estimada_px"] > 0
    assert m2["area_bolsa_visible_estimada_px"] > m1["area_bolsa_visible_estimada_px"] * 20


def test_excel_apertura_expone_maximo_automatico_sin_mover_columnas_legacy(tmp_path):
    from pathlib import Path
    from openpyxl import load_workbook
    from configuracion import cargar_config, cargar_reglas
    from reporte_bmg import generar_reporte_bmg

    fila = {
        "project_id": "P", "case_key": "P|1ST|HT|UGL|42", "tipo_st": "1ST",
        "test_number": "42", "temperature_condition": "HT", "inflator_type": "UGL",
        "campos": {"test_number": {"valor": "42"}},
        "videos": [{
            "id": "cam1", "ruta": "cam1.avi",
            "secuencia_despliegue_r13": {
                "estado": "maximo_inflado_estimado",  # R15.1 (MEDIO-09)
                "frame_breakout": 29, "frame_max_inflado": 70,
                "tiempo_breakout_ms": 5.8, "tiempo_max_inflado_ms": 14.0,
                "duracion_expansion_ms": 8.2, "confianza_maximo": "alta",
                "anomalías": [],
            },
            "evidencia_apertura_r9": {"auto_colorable": True, "nivel_evidencia": "alta", "score_evidencia": .9},
            "tiempo_color_estimado_ms": 5.8, "referencia_color_valida": True,
        }],
    }
    plantilla = Path(__file__).resolve().parents[1] / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    salida = tmp_path / "r13_maximo.xlsx"
    generar_reporte_bmg(plantilla, salida, [fila], cargar_config(), cargar_reglas(), idioma="es")
    ws = load_workbook(salida)["Apertura V15.10"]
    headers = [c.value for c in ws[1]]
    assert headers[:19][-1] == "Ir al detalle"  # contrato previo intacto
    assert headers[19:23] == [
        "Máximo inflado estimado ms", "Duración breakout→máximo ms",
        "Soporte máximo", "Anomalías máximo",
    ]
    assert ws.cell(2, 20).value == 14.0
    assert ws.cell(2, 21).value == 8.2
    assert ws.cell(2, 22).value == 1
