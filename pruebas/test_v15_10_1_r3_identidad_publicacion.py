import sys
import types

from openpyxl import Workbook

from utilidades.identidad_excel_publicacion import (
    aplicar_identidad_a_salida_publicable,
)


def _instalar_base_falsa(monkeypatch, meta):
    modulo = types.ModuleType("utilidades.identidad_excel")
    traducciones = {
        "advertencia": "ADVERTENCIA-BASE",
        "borrador": "BORRADOR-BASE",
        "verificacion": "Identidad verificada",
        "afirmativo": "sí",
    }
    modulo.textos_identidad = lambda idioma="es": traducciones
    modulo.aplicar_identidad_a_salida = lambda *a, **k: dict(meta)
    monkeypatch.setitem(sys.modules, "utilidades.identidad_excel", modulo)


def test_usuario_explicito_publica_si_solo_falta_celda(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "ADVERTENCIA-BASE"
    ws["A2"] = "Identidad verificada"
    ws["B2"] = "no"
    ws["A3"] = "BORRADOR-BASE"
    _instalar_base_falsa(monkeypatch, {
        "publicable": False,
        "decision_usuario_confirmada": True,
        "identidad_verificada": False,
        "estado_identidad": "no_verificada",
        "contradicciones": [],
        "celdas": [],
        "celdas_complementarias": [],
        "advertencias": ["ADVERTENCIA-BASE"],
    })
    meta = aplicar_identidad_a_salida_publicable(
        wb, {"nombre_proyecto": "P1", "project_id": "p1",
             "origen": "declarado_usuario", "sin_asignar": False},
        hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is True
    assert meta["estado_identidad"] == "verificada_por_seleccion_usuario"
    assert ws["A1"].value is None
    assert ws["B2"].value == "sí"
    assert ws["A3"].value is None


def test_contradiccion_real_sigue_bloqueando(monkeypatch):
    wb = Workbook()
    wb.active.title = "Resumen"
    _instalar_base_falsa(monkeypatch, {
        "publicable": False,
        "decision_usuario_confirmada": True,
        "identidad_verificada": False,
        "estado_identidad": "bloqueada_contradiccion",
        "contradicciones": [{"hoja": "Resumen", "celda": "A1", "texto": "Proyecto viejo"}],
        "celdas": [],
        "celdas_complementarias": [],
        "advertencias": [],
    })
    meta = aplicar_identidad_a_salida_publicable(
        wb, {"nombre_proyecto": "P1", "project_id": "p1",
             "origen": "declarado_usuario", "sin_asignar": False},
        hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is False
    assert meta["contradicciones"]


def test_identidad_legacy_no_activa_fallback(monkeypatch):
    wb = Workbook()
    wb.active.title = "Resumen"
    _instalar_base_falsa(monkeypatch, {
        "publicable": False,
        "decision_usuario_confirmada": False,
        "identidad_verificada": False,
        "estado_identidad": "no_verificada",
        "contradicciones": [],
        "celdas": [],
        "celdas_complementarias": [],
        "advertencias": [],
    })
    meta = aplicar_identidad_a_salida_publicable(
        wb, {"nombre_proyecto": "P1", "project_id": "p1",
             "origen": "nombre_carpeta_legacy", "sin_asignar": False},
        hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is False
