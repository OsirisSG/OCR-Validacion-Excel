"""N15: la ruta DECLARADA de proyecto no puede tratar cabeceras tabulares como campo.
Ejecutar en la raíz del repositorio V-Matrix: PYTHONPATH=. pytest /ruta/PRUEBAS_ROJAS_V15_9_7_N15.py -q
Datos sintéticos en memoria, no modifica documentos de producción.
"""
from openpyxl import Workbook
from utilidades.identidad_excel import aplicar_identidad_a_salida

IDENTIDAD = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b", "origen": "declarado_usuario", "sin_asignar": False}

def aplicar(libro):
    return aplicar_identidad_a_salida(libro, IDENTIDAD, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")

def libro_base():
    wb = Workbook()
    wb.active.title = "Resumen"
    return wb, wb.active


def test_n15_cabecera_desnuda_dos_columnas_no_debe_desaparecer():
    wb, ws = libro_base()
    ws["A2"], ws["B2"] = "Proyecto", "Etapa"
    ws["A3"], ws["B3"] = "Proyecto B", "1ST"
    meta = aplicar(wb)
    assert ws["B2"].value == "Etapa", "Se reemplazó una cabecera válida de tabla por el nombre de proyecto"
    assert not [c for c in meta["celdas"] if c["celda_valor"] == "B2"]


def test_n15_cabecera_desnuda_tres_columnas_no_debe_desaparecer():
    wb, ws = libro_base()
    ws["A1"], ws["B1"], ws["C1"] = "Proyecto", "Etapa", "Caso"
    ws["A2"], ws["B2"], ws["C2"] = "Proyecto B", "2ST", 1001
    meta = aplicar(wb)
    assert ws["B1"].value == "Etapa", "Se alteró la estructura de la tabla"
    assert ws["C1"].value == "Caso"
    assert not [c for c in meta["celdas"] if c["celda_valor"] == "B1"]


def test_n15_campo_unico_desnudo_se_puede_actualizar():
    wb, ws = libro_base()
    ws["A2"], ws["B2"] = "Proyecto", "Proyecto A"
    meta = aplicar(wb)
    assert ws["B2"].value == "Proyecto B"
    assert meta["publicable"] is True


def test_n15_campo_explicito_se_puede_actualizar():
    wb, ws = libro_base()
    ws["A2"], ws["B2"] = "Proyecto:", "Proyecto A"
    meta = aplicar(wb)
    assert ws["B2"].value == "Proyecto B"
    assert meta["publicable"] is True
