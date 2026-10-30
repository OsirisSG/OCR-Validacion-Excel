"""N17: proteger cabeceras de tablas reales de Excel sin depender de un diccionario de palabras.

Ejecutar desde la raiz del repo V-Matrix:
  python -m pytest -q /ruta/PRUEBAS_ROJAS_V15_9_9_N17.py

Estos casos trabajan exclusivamente con XLSX temporales y no tocan datos de usuario.
"""
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import pytest
from utilidades.identidad_excel import aplicar_identidad_a_salida

IDENTIDAD = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b",
             "origen": "declarado_usuario", "sin_asignar": False}


def _libro_tabla(etiqueta="Proyecto", segunda="Temperatura", explicito=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A2"] = etiqueta
    ws["B2"] = segunda
    ws.add_table(Table(displayName="TablaEnsayos", ref="A2:B2"))
    if explicito:
        ws["A5"] = "Proyecto:"
        ws["B5"] = "Proyecto A"
    return wb


def _aplicar(wb, idioma):
    return aplicar_identidad_a_salida(wb, IDENTIDAD, idioma=idioma,
                                      hojas_busqueda=["Resumen"],
                                      hoja_resumen="Resumen")


@pytest.mark.parametrize("idioma,etiqueta,segunda", [
    ("es", "Proyecto", "Temperatura"),
    ("en", "Project", "Temperature"),
    ("de", "Projekt", "Temperatur"),
])
def test_n17_excel_table_formal_debe_preservar_encabezado_y_no_publicar(
        tmp_path, idioma, etiqueta, segunda):
    """Una tabla Excel Table registrada es una prueba estructural de cabecera."""
    original = tmp_path / "plantilla.xlsx"
    _libro_tabla(etiqueta, segunda).save(original)
    copia = load_workbook(original)
    meta = _aplicar(copia, idioma)
    salida = tmp_path / "salida.xlsx"
    copia.save(salida)
    final = load_workbook(salida)
    assert final["Resumen"]["B2"].value == segunda, "N17: se corrompio una cabecera de tabla real"
    assert final["Resumen"].tables["TablaEnsayos"].ref == "A2:B2"
    assert final["Resumen"]["A2"].value == etiqueta
    assert meta["identidad_verificada"] is False
    assert meta["publicable"] is False
    assert not any(c["celda_valor"] == "B2" for c in meta["celdas"])
    assert load_workbook(original)["Resumen"]["B2"].value == segunda


def test_n17_tabla_real_y_campo_explicito_se_publica_sin_romper_cabecera(tmp_path):
    original = tmp_path / "plantilla.xlsx"
    _libro_tabla(explicito=True).save(original)
    copia = load_workbook(original)
    meta = _aplicar(copia, "es")
    salida = tmp_path / "salida.xlsx"
    copia.save(salida)
    final = load_workbook(salida)["Resumen"]
    assert final["B2"].value == "Temperatura"
    assert final["B5"].value == "Proyecto B"
    assert meta["publicable"] is True
    assert not any(c["celda_valor"] == "B2" for c in meta["celdas"])
    assert load_workbook(original)["Resumen"]["B5"].value == "Proyecto A"


def test_n17_tabla_visual_vacia_con_columna_no_en_vocabulario_no_se_corrompe():
    """La cabecera puede no pertenecer a un Table formal (plantilla externa)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A2"], ws["B2"] = "Proyecto", "Temperatura"
    meta = _aplicar(wb, "es")
    assert ws["B2"].value == "Temperatura"
    assert not meta["publicable"]


def test_n17_control_campo_unico_real_no_debe_dejar_de_funcionar():
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A2"], ws["B2"] = "Proyecto", "Proyecto A"
    meta = _aplicar(wb, "es")
    assert ws["B2"].value == "Proyecto B"
    assert meta["publicable"] is True


def test_n17_control_explicitamente_etiquetado_sigue_funcionando():
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A2"], ws["B2"] = "Proyecto:", "Proyecto A"
    meta = _aplicar(wb, "es")
    assert ws["B2"].value == "Proyecto B"
    assert meta["publicable"] is True
