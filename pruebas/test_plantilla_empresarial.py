from openpyxl import Workbook, load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from flujo_empresarial import CLAVES_PLANTILLA
from plantilla_empresarial import (cargar_contrato, detectar_plantilla_automatica,
                                   generar_desde_plantilla)


def _plantilla(ruta):
    libro = Workbook()
    instrucciones = libro.active
    instrucciones.title = "Instrucciones"
    captura = libro.create_sheet("Captura_pruebas")
    for columna, clave in enumerate(CLAVES_PLANTILLA, 1):
        captura.cell(4, columna, f"Encabezado {clave}")
        captura.cell(5, columna, clave)
    captura.add_table(Table(displayName="TablaCaptura", ref="A4:AJ205"))
    dv = DataValidation(type="list", formula1='"HT,NT,RT"')
    dv.add("F6:F205")
    captura.add_data_validation(dv)
    diccionario = libro.create_sheet("Diccionario_campos")
    headers = ["posicion", "columna_excel", "encabezado_original", "clave_estable",
               "seccion", "tipo_esperado", "requerido", "valores_o_catalogo",
               "formato_excel", "ejemplo", "significado", "regla_ocr_parser",
               "confianza_inferencia", "minimo", "maximo_regex"]
    for col, header in enumerate(headers, 1):
        diccionario.cell(4, col, header)
    for fila, clave in enumerate(CLAVES_PLANTILLA, 5):
        tipo = "identifier" if clave == "test_number" else "enum" if clave in {
            "temperature_condition", "module_version", "inflator_type"} else "string"
        catalogo = {"temperature_condition": "CAT_TEMPERATURE",
                    "module_version": "CAT_VERSION", "inflator_type": "CAT_INFLATOR"}.get(clave, "Texto")
        valores = [fila - 4, chr(64 + fila - 4) if fila - 4 <= 26 else "AA", clave, clave,
                   "Prueba", tipo, "Sí" if clave == "test_number" else "No", catalogo,
                   "@", None, clave, "", "alta", None, None]
        for col, valor in enumerate(valores, 1):
            diccionario.cell(fila, col, valor)
    catalogos = libro.create_sheet("Catalogos")
    for col, (nombre, valores) in enumerate({
        "CAT_TEMPERATURE": ["HT", "NT", "RT"], "CAT_VERSION": ["RDW", "NAR"],
        "CAT_INFLATOR": ["NOM", "OGL", "UGL"]}.items(), 1):
        catalogos.cell(1, col, nombre)
        catalogos.cell(2, col, "Valor")
        for fila, valor in enumerate(valores, 3):
            catalogos.cell(fila, col, valor)
    resumen = libro.create_sheet("Resumen")
    resumen["B5"] = "=COUNTA(Captura_pruebas!$C$6:$C$205)"
    libro.save(ruta)


def test_detecta_plantilla_compatible_por_estructura(tmp_path):
    raiz = tmp_path / "Proyecto_1ST"
    plantillas = raiz / "plantillas"
    plantillas.mkdir(parents=True)
    esperada = plantillas / "plantilla_captura_1st.xlsx"
    _plantilla(esperada)
    config = {"fase3": {"deteccion_plantillas": {
        "activar": True,
        "estrategias": [{
            "nombre": "estilo_empresa", "perfiles": ["empresarial"],
            "tipos_st": ["1ST"], "patrones_archivo": ["*plantilla*.xlsx"],
        }],
    }}}
    estructura = {"casos_empresariales": [{}], "tipo_st": "1ST", "tipos_st": ["1ST"]}

    resultado = detectar_plantilla_automatica(raiz, estructura, config)

    assert resultado["ruta"] == str(esperada.resolve())
    assert resultado["fuente"] == "estructura"
    assert resultado["estrategia"] == "estilo_empresa"


def test_sin_plantilla_compatible_usa_generador_integrado(tmp_path):
    estructura = {"casos_empresariales": [{}], "tipo_st": "2ST", "tipos_st": ["2ST"]}
    resultado = detectar_plantilla_automatica(
        tmp_path, estructura, {"fase3": {"deteccion_plantillas": {"activar": True}}})
    assert resultado["ruta"] is None
    assert resultado["estilo"] == "generador_estandar"


def _resultado(numero):
    campos = {clave: {"valor": None, "estado": "faltante"} for clave in CLAVES_PLANTILLA}
    campos.update({"test_number": {"valor": str(numero), "estado": "extraido_ruta"},
                   "temperature_condition": {"valor": "HT", "estado": "extraido_ruta"},
                   "module_version": {"valor": "RDW", "estado": "extraido_ruta"},
                   "inflator_type": {"valor": "NOM", "estado": "extraido_ruta"}})
    return {"campos": campos, "trazabilidad": [{
        "test_number": str(numero), "tipo_st": "1ST", "fase": "VOR",
        "ruta_absoluta": f"/tmp/{numero}.jpg", "ruta_relativa": f"PHOTOS/VOR/{numero}.jpg",
        "estado_imagen": "descartada_sin_texto", "campos_detectados": [],
        "requiere_revision": False,
    }]}


def test_carga_claves_diccionario_y_catalogos(tmp_path):
    ruta = tmp_path / "plantilla.xlsx"
    _plantilla(ruta)
    contrato = cargar_contrato(ruta)
    assert len(contrato["claves"]) == 36
    assert contrato["claves"]["test_number"] == 3
    assert contrato["catalogos"]["CAT_TEMPERATURE"] == ["HT", "NT", "RT"]


def test_una_fila_por_id_trazabilidad_y_formula_mas_alla_205(tmp_path):
    plantilla = tmp_path / "plantilla.xlsx"
    salida = tmp_path / "salida.xlsx"
    _plantilla(plantilla)
    resultados = [_resultado(numero) for numero in range(1, 202)]
    generar_desde_plantilla(plantilla, salida, resultados)
    libro = load_workbook(salida, data_only=False)
    captura = libro["Captura_pruebas"]
    assert captura["C6"].value == "1"
    assert captura["C206"].value == "201"
    assert captura["C6"].number_format == "@"
    assert "206" in libro["Resumen"]["B5"].value
    assert libro["Trazabilidad_OCR"].max_row == 202
    assert next(iter(captura.tables.values())).ref.endswith("AJ206")
    assert any("206" in str(dv.sqref) for dv in captura.data_validations.dataValidation)
