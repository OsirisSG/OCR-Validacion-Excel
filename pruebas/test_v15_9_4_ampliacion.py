"""V15.9.4 · Ampliación de pruebas exigida por la auditoría (N1-N6).

Complementa `test_v15_9_4_auditoria_regresiones.py` (los 7 casos literales)
con los escenarios sintéticos adicionales de
`CORRECCIONES_OBLIGATORIAS_V_MATRIX_V15_9_4.md`:

- N1/N2: plantillas empresarial y BMG real, título en A1 + etiqueta en A2,
  Excel ES/EN/DE, «Sin asignar», proyectos A/B consecutivos, reanudación con
  proyecto B y plantilla con A (sin mezclas de project_id).
- N3: registro legible pero con elementos inválidos, duplicados, null,
  objetos incompletos y Unicode válido; ninguna pérdida de bytes.
- N4: matriz de consensos 3/4 limpio, 3/4 con revisión obligatoria, 2/4,
  1/1, 2/2 históricos y manual auditado.
- N5: fallo de escritura de resultados con invalidación de caché, reintento y
  recuperación.
- N6: cola Excel conserva el estado ante publicación bloqueada, borradores
  con PermissionError inyectado y descargas diferenciadas.

Todo sintético (tmp_path); no toca datos reales ni de red.
"""
from __future__ import annotations

import hashlib
import importlib
import json
import pathlib
import sys

import pytest
from openpyxl import Workbook, load_workbook

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "pruebas"))

import utilidades.proyectos as registro
import video_airbag as va
from utilidades.persistencia import ResultadoPersistencia
from test_plantilla_empresarial import _plantilla


def _backend():
    return importlib.import_module("dashboard.backend.app")


def _hash(ruta) -> str:
    return hashlib.sha256(pathlib.Path(ruta).read_bytes()).hexdigest()


def _item(nombre, tiempo, *, estado="candidato", revision=None, **extra):
    it = {"nombre": nombre, "estado": estado, "tiempo_ms": tiempo,
          "tiempo_fps_ms": tiempo, "resolucion_nativa_ms": 0.1,
          "precision_fuente_ms": 0.1, "fuente_tiempo": "fps_subframe"}
    if revision:
        it["revision_manual"] = revision
    it.update(extra)
    return it


CAM_A, CAM_B, CAM_C, CAM_D = ("09_DETAIL.AVI", "09_FRONTAL.AVI",
                              "09_LATERAL.AVI", "09_SCHEIBE.AVI")

IDENTIDAD_A = {"nombre_proyecto": "Proyecto A", "project_id": "proyecto-a",
               "origen": "declarado_usuario", "sin_asignar": False,
               "pendiente": False, "ejecucion_id": "ej-a"}
IDENTIDAD_B = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b",
               "origen": "declarado_usuario", "sin_asignar": False,
               "pendiente": False, "ejecucion_id": "ej-b"}
SIN_ASIGNAR = {"nombre_proyecto": None, "project_id": "sin-asignar",
               "origen": "sin_asignar_expreso", "sin_asignar": True,
               "pendiente": True}

MARCADORES_IDIOMA = {
    "es": ("REFERENCIA CONTRADICTORIA", "BORRADOR NO PUBLICABLE"),
    "en": ("CONTRADICTORY REFERENCE", "NON-PUBLISHABLE DRAFT"),
    "de": ("WIDERSPRÜCHLICHER VERWEIS", "NICHT VERÖFFENTLICHUNGSFÄHIGER ENTWURF"),
}


def _resumen_texto(libro, hoja="Resumen"):
    return "\n".join(str(c.value) for fila in libro[hoja].iter_rows()
                     for c in fila if c.value)


@pytest.fixture
def plantilla_titulo_y_etiqueta(tmp_path):
    """Título histórico en A1 + celda de proyecto etiquetada A2/B2."""
    ruta = tmp_path / "externa_titulo.xlsx"
    _plantilla(ruta)
    libro = load_workbook(ruta)
    captura = libro["Captura_pruebas"]
    captura["A1"] = "Reporte del Proyecto A"
    captura["A2"] = "Proyecto:"
    captura["B2"] = "Proyecto A"
    libro.save(ruta)
    return ruta


# ---------------------------------------------------------------------------
# N1/N2 · identidad visual contradictoria en Excel
# ---------------------------------------------------------------------------


def test_n1_titulo_contradictorio_bloquea_en_los_tres_idiomas(
        plantilla_titulo_y_etiqueta, tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    hash_plantilla = _hash(plantilla_titulo_y_etiqueta)
    for idioma in ("es", "en", "de"):
        oficial = tmp_path / f"salida_{idioma}.xlsx"
        resultado = pathlib.Path(generar_desde_plantilla(
            plantilla_titulo_y_etiqueta, oficial, [],
            identidad_proyecto=IDENTIDAD_B, idioma=idioma))
        assert not oficial.exists(), idioma
        assert "borrador_no_publicable" in resultado.name and resultado.is_file()
        libro = load_workbook(resultado)
        # El título ajeno NUNCA se reemplaza; la celda etiquetada sí se mapea.
        assert libro["Captura_pruebas"]["A1"].value == "Reporte del Proyecto A"
        assert libro["Captura_pruebas"]["B2"].value == "Proyecto B"
        resumen = _resumen_texto(libro)
        referencia, marcador = MARCADORES_IDIOMA[idioma]
        assert referencia in resumen, idioma
        assert marcador in resumen.upper(), idioma
    assert _hash(plantilla_titulo_y_etiqueta) == hash_plantilla


def test_n1_titulo_coherente_con_identidad_publica(plantilla_titulo_y_etiqueta, tmp_path):
    """Si el título visible concuerda con la identidad declarada, se publica."""
    from plantilla_empresarial import generar_desde_plantilla
    libro = load_workbook(plantilla_titulo_y_etiqueta)
    libro["Captura_pruebas"]["A1"] = "Reporte del Proyecto B"
    libro.save(plantilla_titulo_y_etiqueta)
    oficial = tmp_path / "salida_coherente.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla_titulo_y_etiqueta, oficial, [], identidad_proyecto=IDENTIDAD_B))
    assert resultado == oficial and oficial.is_file()
    assert load_workbook(oficial)["Captura_pruebas"]["B2"].value == "Proyecto B"


def test_n1_bmg_plantilla_real_con_titulo_ajeno_bloquea(tmp_path):
    from reporte_bmg import generar_reporte_bmg
    from configuracion import cargar_config, cargar_reglas
    from test_reporte_bmg_v7 import caso
    real = RAIZ / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    hash_real = _hash(real)
    resultados = [caso(1001, "1ST", "HT", "UGL", 5.8)]
    # Copia contaminada con un título de otro proyecto en una celda libre.
    sucia = tmp_path / "bmg_sucia.xlsx"
    sucia.write_bytes(real.read_bytes())
    libro = load_workbook(sucia)
    libro["Übersicht Standversuche"]["Z1"] = "Reporte del Proyecto A"
    libro.save(sucia)
    oficial = tmp_path / "bmg_oficial.xlsx"
    salida = pathlib.Path(generar_reporte_bmg(
        sucia, oficial, resultados, cargar_config(), cargar_reglas(),
        idioma="es", identidad_proyecto=IDENTIDAD_B))
    assert not oficial.exists()
    assert "borrador_no_publicable" in salida.name and salida.is_file()
    assert "REFERENCIA CONTRADICTORIA" in _resumen_texto(
        load_workbook(salida), "Übersicht Standversuche")
    # La plantilla real del repositorio no se modificó.
    assert _hash(real) == hash_real
    # Control actual: la MISMA plantilla limpia sí publica porque la selección
    # explícita del usuario es la fuente autoritativa. La ausencia de una celda
    # histórica de proyecto no es una contradicción.
    limpia = tmp_path / "bmg_limpia.xlsx"
    limpia.write_bytes(real.read_bytes())
    oficial2 = tmp_path / "bmg_oficial2.xlsx"
    salida2 = pathlib.Path(generar_reporte_bmg(
        limpia, oficial2, resultados, cargar_config(), cargar_reglas(),
        idioma="es", identidad_proyecto=IDENTIDAD_B))
    assert salida2 == oficial2 and oficial2.exists()
    libro2 = load_workbook(salida2)
    assert libro2["Übersicht Standversuche"]["A1"].value == "Proyecto B"


def test_n2_sin_asignar_con_titulo_historico_bloquea(plantilla_titulo_y_etiqueta, tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    oficial = tmp_path / "salida_sin_asignar.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla_titulo_y_etiqueta, oficial, [], identidad_proyecto=SIN_ASIGNAR))
    assert not oficial.exists()
    assert "borrador_no_publicable" in resultado.name
    libro = load_workbook(resultado)
    # El título del Proyecto A sigue intacto (no se suplanta ni se autoriza).
    assert libro["Captura_pruebas"]["A1"].value == "Reporte del Proyecto A"
    resumen = _resumen_texto(libro)
    assert "REFERENCIA CONTRADICTORIA" in resumen
    assert "Sin asignar" in resumen  # la decisión humana queda en trazabilidad


def test_n2_sin_asignar_con_plantilla_limpia_publica_pendiente(tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = tmp_path / "limpia.xlsx"
    _plantilla(plantilla)
    oficial = tmp_path / "salida_pendiente.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla, oficial, [], identidad_proyecto=SIN_ASIGNAR))
    assert resultado == oficial and oficial.is_file()
    assert "Sin asignar" in _resumen_texto(load_workbook(oficial))


def test_n1_proyectos_consecutivos_a_y_b_sin_fugas(plantilla_titulo_y_etiqueta, tmp_path):
    """Dos proyectos seguidos con la misma plantilla titulada «Proyecto A»."""
    from plantilla_empresarial import generar_desde_plantilla
    # Corrida 1: declara A, plantilla coherente (título y valor A) ⇒ publica.
    libro = load_workbook(plantilla_titulo_y_etiqueta)
    libro["Captura_pruebas"]["A1"] = "Reporte del Proyecto A"
    libro["Captura_pruebas"]["B2"] = "Proyecto A"
    libro.save(plantilla_titulo_y_etiqueta)
    oficial_a = tmp_path / "corrida_a.xlsx"
    salida_a = pathlib.Path(generar_desde_plantilla(
        plantilla_titulo_y_etiqueta, oficial_a, [], identidad_proyecto=IDENTIDAD_A))
    assert salida_a == oficial_a and oficial_a.is_file()
    hash_oficial_a = _hash(oficial_a)
    # Corrida 2: declara B con la misma plantilla (título A) ⇒ bloqueada.
    oficial_b = tmp_path / "corrida_b.xlsx"
    salida_b = pathlib.Path(generar_desde_plantilla(
        plantilla_titulo_y_etiqueta, oficial_b, [], identidad_proyecto=IDENTIDAD_B))
    assert "borrador_no_publicable" in salida_b.name and not oficial_b.exists()
    # La salida oficial de A queda byte-idéntica (no contaminada por B).
    assert _hash(oficial_a) == hash_oficial_a
    texto_a = _resumen_texto(load_workbook(oficial_a))
    assert "proyecto-b" not in texto_a and "REFERENCIA CONTRADICTORIA" not in texto_a
    libro_b = load_workbook(salida_b)
    texto_b = _resumen_texto(libro_b)
    # Sin mezclas de project_id: el borrador expone sólo la identidad B.
    assert "proyecto-b" in texto_b and "proyecto-a" not in texto_b
    assert libro_b["Captura_pruebas"]["A1"].value == "Reporte del Proyecto A"


def test_n2_reanudacion_proyecto_b_con_plantilla_a_sin_mezclas(
        plantilla_titulo_y_etiqueta, tmp_path):
    """Reanudación migrada a B: la identidad B nunca se mezcla con A."""
    from plantilla_empresarial import generar_desde_plantilla
    from utilidades.identidad import resolver_identidad_proyecto
    identidad = resolver_identidad_proyecto(
        declarado={"nombre_proyecto": "Proyecto B", "proyecto_id": "proyecto-b",
                   "migracion_explicita": True},
        raiz=tmp_path,
        checkpoint={"project_id": "proyecto-a", "nombre_proyecto": "Proyecto A"})
    # Con checkpoint vigente la identidad conservadora es A (P04); la
    # migración explícita se resuelve en el backend. Aquí se prueba el caso
    # declarado sin checkpoint y el caso checkpoint sin migrar.
    assert identidad["project_id"] == "proyecto-a"  # checkpoint manda (P04)
    identidad_b = resolver_identidad_proyecto(
        declarado={"nombre_proyecto": "Proyecto B", "proyecto_id": "proyecto-b"},
        raiz=tmp_path)
    assert identidad_b["project_id"] == "proyecto-b"
    # Plantilla con título A + identidad B ⇒ borrador sin mezclas de ID.
    oficial = tmp_path / "reanudado.xlsx"
    salida = pathlib.Path(generar_desde_plantilla(
        plantilla_titulo_y_etiqueta, oficial, [], identidad_proyecto=identidad_b))
    assert "borrador_no_publicable" in salida.name
    texto = _resumen_texto(load_workbook(salida))
    assert "proyecto-b" in texto and "proyecto-a" not in texto


# ---------------------------------------------------------------------------
# N3 · registro parcialmente corrupto
# ---------------------------------------------------------------------------


CASOS_CORRUPTOS = [
    ("entrada no objeto", {"version": 1, "proyectos": [
        {"proyecto_id": "proyecto-a", "nombre_proyecto": "Proyecto A"},
        "entrada_anterior_no_parseable"]}),
    ("entrada null", {"version": 1, "proyectos": [None]}),
    ("objeto incompleto", {"version": 1, "proyectos": [{"nombre_proyecto": "A"}]}),
    ("proyecto_id vacio", {"version": 1, "proyectos": [
        {"proyecto_id": "   ", "nombre_proyecto": "A"}]}),
    ("proyecto_id no texto", {"version": 1, "proyectos": [
        {"proyecto_id": 123, "nombre_proyecto": "A"}]}),
    ("nombre no texto", {"version": 1, "proyectos": [
        {"proyecto_id": "a", "nombre_proyecto": 123}]}),
    ("id duplicado", {"version": 1, "proyectos": [
        {"proyecto_id": "a", "nombre_proyecto": "A"},
        {"proyecto_id": "A", "nombre_proyecto": "A2"}]}),
]


@pytest.mark.parametrize("motivo,contenido", CASOS_CORRUPTOS)
def test_n3_estructura_parcialmente_corrupta_bloquea_sin_perder_bytes(
        tmp_path, motivo, contenido):
    ruta = tmp_path / "proyectos.json"
    ruta.write_text(json.dumps(contenido, ensure_ascii=False), encoding="utf-8")
    antes = ruta.read_bytes()
    with pytest.raises(registro.RegistroProyectosCorruptoError):
        registro.listar_proyectos(ruta)
    with pytest.raises(registro.RegistroProyectosCorruptoError):
        registro.registrar_proyecto("Proyecto C", ruta=ruta)
    assert ruta.read_bytes() == antes, motivo  # cero bytes perdidos
    # Recuperación explícita con respaldo íntegro de TODO el documento.
    info = registro.recuperar_registro(ruta)
    assert info["recuperado"] is True
    assert pathlib.Path(info["respaldo"]).read_bytes() == antes
    assert registro.listar_proyectos(ruta) == []


def test_n3_unicode_valido_no_bloquea_ni_pierde_informacion(tmp_path):
    ruta = tmp_path / "proyectos.json"
    contenido = {"version": 1, "proyectos": [
        {"proyecto_id": "proyecto-ñ", "nombre_proyecto": "Proyecto Ñ Ünicode Ω",
         "creado_en": "2026-10-12T00:00:00"}]}
    ruta.write_text(json.dumps(contenido, ensure_ascii=False), encoding="utf-8")
    assert [p["nombre_proyecto"] for p in registro.listar_proyectos(ruta)] \
        == ["Proyecto Ñ Ünicode Ω"]
    creado = registro.registrar_proyecto("Proyecto Ω", ruta=ruta)
    assert creado["proyecto_id"]
    nombres = {p["nombre_proyecto"] for p in registro.listar_proyectos(ruta)}
    assert nombres == {"Proyecto Ñ Ünicode Ω", "Proyecto Ω"}  # nada se perdió


# ---------------------------------------------------------------------------
# N4 · política única de publicación de consensos
# ---------------------------------------------------------------------------


def test_n4_matriz_de_publicacion_automatica():
    base = {"estado": "primera_aparicion", "tiempo_ms": 5.0,
            "requiere_revision": False}
    assert va.consenso_valido({**base, "soporte": 3, "total_principales": 4}) is True
    assert va.consenso_valido({**base, "soporte": 4, "total_principales": 4}) is True
    assert va.consenso_valido({**base, "soporte": 3, "total_principales": 3}) is True
    assert va.consenso_valido({**base, "soporte": 2, "total_principales": 4}) is False
    assert va.consenso_valido({**base, "soporte": 2, "total_principales": 2}) is False
    assert va.consenso_valido({**base, "soporte": 1, "total_principales": 1}) is False
    assert va.consenso_valido({**base, "soporte": 3, "total_principales": 4,
                               "requiere_revision": True}) is False
    # Manual auditado: confirmación humana real con evidencia, siempre manda.
    assert va.consenso_valido({"estado": "confirmado_manual", "tiempo_ms": 5.0,
                               "soporte": 1, "total_principales": 4}) is True


def test_n4_tres_de_4_con_revision_obligatoria_queda_pendiente():
    items = [_item(CAM_A, 5.0),
             _item(CAM_B, 5.0, advertencia_fps="12 fps declarados, 1000 reales"),
             _item(CAM_C, 5.0), _item(CAM_D, 90.0)]
    consenso = va.consenso_videos(items, {})
    assert consenso["estado_consenso"] == "confirmado"
    assert consenso["requiere_revision"] is True
    assert va.consenso_valido(consenso) is False


def test_n4_legado_2_de_2_migra_a_pendiente_sin_videos(tmp_path, monkeypatch):
    """N4 en la migración: 2/2 histórico baja a pendiente conservando medida."""
    backend = _backend()
    fila = {"nombre": "CASO22", "case_key": "CASO22", "videos": [],
            "campos": {"opening_time_ms": {"valor": "5.4", "estado": "extraido_video",
                                           "fuente": "video_airbag_consenso"}},
            "video_apertura": {"estado": "primera_aparicion", "tiempo_ms": 5.4,
                               "soporte": 2, "total_principales": 2,
                               "requiere_revision": False}}
    validacion = tmp_path / "validacion_resultados.json"
    validacion.write_text(json.dumps({"resultados": [fila]}), encoding="utf-8")
    monkeypatch.setattr(backend, "RUTA_VALIDACION", validacion)
    backend._cache.update({"mtime": None, "datos": None})
    try:
        f = backend._cargar_datos()["resultados"][0]
        assert f["campos"]["opening_time_ms"]["estado"] == "pendiente_revision"
        assert float(f["campos"]["opening_time_ms"]["valor"]) == pytest.approx(5.4)
        assert f["video_apertura_anterior"]["soporte"] == 2
    finally:
        backend._cache.update({"mtime": None, "datos": None})


# ---------------------------------------------------------------------------
# N5 · persistencia real de la migración
# ---------------------------------------------------------------------------


def _fila_legada_invalida():
    return {"nombre": "Caso 1", "case_key": "c1", "videos": [],
            "video_apertura": {"estado": "primera_aparicion", "tiempo_ms": 5.4,
                               "soporte": 1, "requiere_revision": False},
            "campos": {"opening_time_ms": {"valor": "5.4", "estado": "extraido_video",
                                           "fuente": "video_airbag_consenso"}}}


def test_n5_migracion_fallida_invalida_cache_y_reintenta(tmp_path, monkeypatch):
    backend = _backend()
    validacion = tmp_path / "validacion_resultados.json"
    validacion.write_text(json.dumps({"resultados": [_fila_legada_invalida()]}),
                          encoding="utf-8")
    monkeypatch.setattr(backend, "RUTA_VALIDACION", validacion)
    backend._cache.update({"mtime": None, "datos": None})
    monkeypatch.setattr(backend, "_advertencia_migracion_consensos", None)
    # R15.2 (CRÍTICO-02): la escritura del documento vive en la autoridad
    # compartida; el fallo intermitente se inyecta allí.
    from utilidades import persistencia_documento
    real = persistencia_documento.escribir_json_seguro
    llamadas = {"n": 0}

    def intermitente(*args, **kwargs):
        llamadas["n"] += 1
        if llamadas["n"] == 1:
            return ResultadoPersistencia(False, validacion, 1, "ERROR INYECTADO")
        return real(*args, **kwargs)

    monkeypatch.setattr(persistencia_documento, "escribir_json_seguro", intermitente)
    try:
        datos = backend._cargar_datos()
        # Fallo visible + caché invalidada para reintento.
        assert backend._advertencia_migracion_consensos
        assert backend._cache["mtime"] is None
        # La vista en memoria es fail-closed aunque el disco siga antiguo.
        assert datos["resultados"][0]["campos"]["opening_time_ms"]["estado"] \
            == "pendiente_revision"
        disco = json.loads(validacion.read_text(encoding="utf-8"))
        assert disco["resultados"][0]["campos"]["opening_time_ms"]["estado"] \
            == "extraido_video"  # disco intacto, nada declarado como hecho
        # Reintento en la siguiente carga: ahora sí persiste.
        backend._cargar_datos()
        disco2 = json.loads(validacion.read_text(encoding="utf-8"))
        assert disco2["resultados"][0]["campos"]["opening_time_ms"]["estado"] \
            == "pendiente_revision"
        assert disco2["resultados"][0]["video_apertura_anterior"]["soporte"] == 1
        assert backend._advertencia_migracion_consensos is None
    finally:
        backend._cache.update({"mtime": None, "datos": None})
        monkeypatch.setattr(backend, "_advertencia_migracion_consensos", None)


# ---------------------------------------------------------------------------
# N6 · publicación bloqueada: cola, borradores y descargas
# ---------------------------------------------------------------------------


def test_n6_cola_excel_conserva_estado_hasta_corregir_el_mapeo(tmp_path):
    from utilidades.cola_excel import ColaExcel
    from utilidades.identidad_excel import PublicacionExcelBloqueadaError
    intentos = {"n": 0}
    oficial = tmp_path / "oficial.xlsx"

    def generar():
        intentos["n"] += 1
        if intentos["n"] == 1:
            raise PublicacionExcelBloqueadaError(
                "Exportación Excel BLOQUEADA: identidad no verificable; "
                "borrador guardado aparte.")
        return oficial  # mapeo corregido: publicación oficial

    cola = ColaExcel(tmp_path / "cola.sqlite3", generar)
    assert cola.solicitar("correccion-1") == 1
    assert cola.procesar() is False
    estado = cola.estado()
    assert estado["pendientes"] == 1  # conserva el estado, no finge guardado
    assert "BLOQUEADA" in (estado["error"] or "")
    assert cola.procesar() is True
    assert cola.estado()["pendientes"] == 0


def test_n6_asociado_lanza_ante_borrador_en_vez_de_exito(tmp_path, monkeypatch):
    backend = _backend()
    from utilidades.identidad_excel import PublicacionExcelBloqueadaError
    borrador = tmp_path / "resultado.borrador_no_publicable.xlsx"
    borrador.write_bytes(b"PK\x03\x04borrador")
    datos = {"resultados": [{"nombre": "caso"}],
             "nombre_excel_base": str(tmp_path / "resultado.xlsx")}
    monkeypatch.setattr(backend, "_cargar_datos", lambda: datos)
    monkeypatch.setattr(backend, "_limites_apertura_efectivos",
                        lambda: (6, 7, "predeterminado", None))
    monkeypatch.setattr(backend, "_generar_excel_serializado",
                        lambda *a, **k: borrador)
    with pytest.raises(PublicacionExcelBloqueadaError):
        backend._generar_excel_asociado()


def test_n6_descarga_oficial_rechaza_borrador_y_borrador_va_por_su_ruta(
        tmp_path, monkeypatch):
    backend = _backend()
    from fastapi import HTTPException
    borrador = tmp_path / "resultado_ES.borrador_no_publicable.xlsx"
    borrador.write_bytes(b"PK\x03\x04borrador")
    datos = {"resultados": [{"nombre": "caso"}],
             "nombre_excel_base": str(tmp_path / "resultado.xlsx")}
    monkeypatch.setattr(backend, "_cargar_datos", lambda: datos)
    with pytest.raises(HTTPException) as err:
        backend.descargar_exportado(borrador.name)
    assert err.value.status_code == 404
    respuesta = backend.descargar_borrador(borrador.name)
    assert pathlib.Path(respuesta.path).name == borrador.name
    # La ruta de borrador no sirve archivos oficiales.
    oficial = tmp_path / "resultado_ES.xlsx"
    Workbook().save(oficial)
    with pytest.raises(HTTPException):
        backend.descargar_borrador(oficial.name)


def test_n6_permissionerror_en_reemplazo_atomico_no_toca_la_salida(
        tmp_path, monkeypatch):
    """Borrador/publicación con PermissionError inyectado en el reemplazo."""
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = tmp_path / "limpia.xlsx"
    _plantilla(plantilla)
    oficial = tmp_path / "salida.xlsx"
    libro_previo = Workbook()
    libro_previo.active["A1"] = "resultado anterior vigente"
    libro_previo.save(oficial)
    hash_antes = _hash(oficial)
    import os as _os

    def reemplazo_bloqueado(src, dst, *args, **kwargs):
        raise PermissionError("destino abierto en Excel")

    monkeypatch.setattr(_os, "replace", reemplazo_bloqueado)
    with pytest.raises(PermissionError):
        generar_desde_plantilla(plantilla, oficial, [], identidad_proyecto=None)
    assert _hash(oficial) == hash_antes  # salida anterior intacta
    temporales = [p for p in tmp_path.iterdir() if p.name.startswith(".")]
    assert temporales == []  # sin temporales huérfanos
