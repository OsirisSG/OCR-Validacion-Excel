"""Pruebas de auditoría V15.9.3: actualmente ROJAS con el parche recibido.

Ejecutar desde el repositorio completo, de forma aislada:
  PYTHONPATH=. python -m pytest /ruta/PRUEBAS_ROJAS_V15_9_3.py -q
No accede a datos reales: usa libros y registros temporales de pytest.
"""
import json
from pathlib import Path
from unittest.mock import patch
import pytest
from openpyxl import Workbook

import video_airbag
import utilidades.proyectos as proyectos
from utilidades.identidad_excel import aplicar_identidad_a_salida
from utilidades.persistencia import ResultadoPersistencia

IDENTIDAD_B = {'nombre_proyecto':'Proyecto B','project_id':'proyecto-b','origen':'declarado_usuario','sin_asignar':False}


def test_legacy_1_de_1_no_es_3_de_4():
    assert not video_airbag.consenso_valido({'estado':'primera_aparicion','tiempo_ms':5.4,'soporte':1,'total_principales':1,'requiere_revision':False})


def test_legacy_2_de_2_no_es_3_de_4():
    assert not video_airbag.consenso_valido({'estado':'primera_aparicion','tiempo_ms':5.4,'soporte':2,'total_principales':2,'requiere_revision':False})


def test_plantilla_titulo_viejo_no_publicable_aun_con_celda_mapeada():
    wb=Workbook(); ws=wb.active;ws.title='Resumen'
    ws['A1']='Reporte del Proyecto A'; ws['A2']='Proyecto:';ws['B2']='Proyecto A'
    meta=aplicar_identidad_a_salida(wb, IDENTIDAD_B, hojas_busqueda=['Resumen'], hoja_resumen='Resumen')
    assert not (meta['publicable'] and 'Proyecto A' in str(ws['A1'].value))


def test_sin_asignar_no_deja_proyecto_precargado_visible():
    wb=Workbook(); ws=wb.active;ws.title='Resumen';ws['A1']='Reporte del Proyecto A'
    identidad={'origen':'declarado_sin_asignar','sin_asignar':True,'project_id':'sin-asignar'}
    meta=aplicar_identidad_a_salida(wb, identidad, hojas_busqueda=['Resumen'], hoja_resumen='Resumen')
    assert not (meta['publicable'] and 'Proyecto A' in str(ws['A1'].value))


def test_registro_parcialmente_corrupto_no_pierde_entradas(tmp_path):
    ruta=tmp_path/'proyectos.json'
    contenido={'version':1,'proyectos':[{'proyecto_id':'proyecto-a','nombre_proyecto':'Proyecto A'},'entrada_anterior_no_parseable']}
    ruta.write_text(json.dumps(contenido), encoding='utf8')
    antes=ruta.read_bytes()
    with pytest.raises(proyectos.RegistroProyectosCorruptoError):
        proyectos.registrar_proyecto('Proyecto C',ruta=ruta)
    assert ruta.read_bytes()==antes


def test_migracion_no_reporta_exito_si_fallo_escritura(tmp_path, monkeypatch):
    import dashboard.backend.app as backend
    destino=tmp_path/'validacion_resultados.json'
    datos={'resultados':[{'nombre':'Caso 1','case_key':'c1','videos':[],
         'video_apertura':{'estado':'primera_aparicion','tiempo_ms':5.4,'soporte':1,'requiere_revision':False},
         'campos':{'opening_time_ms':{'valor':'5.4','estado':'extraido_video','fuente':'video_airbag_consenso'}}}]}
    destino.write_text(json.dumps(datos),encoding='utf8')
    monkeypatch.setattr(backend,'RUTA_VALIDACION',destino)
    # R15.2 (CRÍTICO-02): la escritura del documento vive en la autoridad
    # compartida; el fallo inyectado se parchea allí (contrato reemplazado,
    # documentado en CHANGELOG_R15_2).
    monkeypatch.setattr('utilidades.persistencia_documento.escribir_json_seguro',lambda *a,**k: ResultadoPersistencia(False,destino,1,'ERROR INYECTADO'))
    # Debe propagarse error o devolver 0 si el cambio no se hizo duradero;
    # no afirmar que '1' migración persistió cuando disco sigue en estado antiguo.
    resultado=backend._migrar_consensos_video(datos)
    assert resultado==0


def test_api_exportacion_no_ofrece_borrador_como_exito(tmp_path):
    import dashboard.backend.app as backend
    borrador=tmp_path/'resultado_ES.borrador_no_publicable.xlsx'
    borrador.write_bytes(b'existe pero no debe presentarse como oficial')
    datos={'resultados':[{'nombre':'caso'}],'nombre_excel_base':str(tmp_path/'resultado.xlsx')}
    with (patch.object(backend,'_cargar_datos',return_value=datos),
          patch.object(backend,'_limites_apertura_efectivos',return_value=(6,7,'predeterminado',None)),
          patch.object(backend,'_generar_excel_serializado',return_value=borrador)):
        result=backend.exportar_resultados(backend.SolicitudExportacion(idiomas=['es']))
    if result.get('generados'):
        assert all('borrador_no_publicable' not in d['nombre'] for d in result['generados']) or result.get('publicable') is False
