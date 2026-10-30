"""V15.9.9 · Ampliación de la corrección CRIT-N16.

Tablas de dos columnas VACÍAS (sólo cabecera «Proyecto | Etapa», sin filas
de datos) en hojas declaradas y complementarias: la cabecera reconocible
nunca se sobrescribe, no se registra como campo y, sin otro campo
inequívoco, el reporte identificado queda en borrador no publicable.
Cobertura exigida por `CORRECCIONES_OBLIGATORIAS_V_MATRIX_V15_9_9.md`:
ES/EN/DE, tablas vacías y con datos, combinadas, fórmulas, vocabulario de
cabeceras extendido, colisión semántica genuina (proyecto llamado «Etapa»),
exportadores empresarial y BMG reales con Excel temporales guardados y
releídos y hash de plantilla original, y paridad con el mapeo declarado.
"""
from __future__ import annotations

import pathlib
import sys

import pytest
from openpyxl import Workbook, load_workbook

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "pruebas"))

from test_plantilla_empresarial import _plantilla
from test_v15_9_6_ampliacion import _hash, _plantilla_publicable
from utilidades.identidad_excel import aplicar_identidad_a_salida

IDENTIDAD_B = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b",
               "origen": "declarado_usuario", "sin_asignar": False,
               "pendiente": False, "ejecucion_id": "ej-v1599"}
IDENTIDAD_ETAPA = {"nombre_proyecto": "Etapa", "project_id": "etapa",
                   "origen": "declarado_usuario", "sin_asignar": False}


def _libro_base():
    libro = Workbook()
    libro.active.title = "Resumen"
    return libro


def _aplicar(libro, identidad=None, idioma="es", **extra):
    return aplicar_identidad_a_salida(
        libro, identidad or IDENTIDAD_B, hojas_busqueda=["Resumen"],
        hoja_resumen="Resumen", idioma=idioma, **extra)


# ---------------------------------------------------------------------------
# Cabecera vacía: omisión sin contradicciones falsas (aplicador directo)
# ---------------------------------------------------------------------------


def test_n16_cabecera_vacia_complementaria_sola_no_verifica_sin_bloqueos():
    libro = _libro_base()
    ficha = libro.create_sheet("Ficha")
    ficha["A1"], ficha["B1"] = "Proyecto", "Etapa"     # tabla vacía
    meta = _aplicar(libro)
    assert (ficha["A1"].value, ficha["B1"].value) == ("Proyecto", "Etapa")
    assert meta["celdas"] == []
    assert meta["celdas_complementarias"] == []
    # Cabecera legítima: omisión silenciosa, sin contradicción estructural.
    assert not [c for c in meta["contradicciones"] if c.get("estructural")]
    assert meta["identidad_verificada"] is False
    assert meta["estado_identidad"] == "no_verificada"
    assert meta["publicable"] is False


def test_n16_cabecera_vacia_con_combinada_debajo_sin_excepcion():
    libro = _libro_base()
    resumen = libro["Resumen"]
    resumen["A2"], resumen["B2"] = "Proyecto", "Etapa"
    resumen.merge_cells("B3:C3")        # estructura combinada debajo
    meta = _aplicar(libro)              # sin excepción de bajo nivel
    assert (resumen["A2"].value, resumen["B2"].value) == ("Proyecto", "Etapa")
    assert meta["celdas"] == []
    assert meta["publicable"] is False


@pytest.mark.parametrize("idioma,etiqueta,cabecera", [
    ("es", "Proyecto", "Valor"),
    ("es", "Proyecto", "Estado"),
    ("en", "Project", "Case"),
    ("en", "Project", "Date"),
    ("de", "Projekt", "Fall"),
    ("de", "Projekt", "Wert"),
])
def test_n16_vocabulario_cabeceras_extendido_tres_idiomas(idioma, etiqueta,
                                                          cabecera):
    libro = _libro_base()
    resumen = libro["Resumen"]
    resumen["A2"], resumen["B2"] = etiqueta, cabecera
    meta = _aplicar(libro, idioma=idioma)
    assert resumen["B2"].value == cabecera, (idioma, cabecera)
    assert meta["celdas"] == []
    assert meta["publicable"] is False


def test_n16_colision_semantica_proyecto_llamado_etapa_exige_explicita():
    """Proyecto llamado exactamente «Etapa»: la cabecera posible no se toca.

    Con etiqueta desnuda no se sobrescribe jamás (fail-closed, borrador);
    la asignación EXPLÍCITA (contrato con separador) sí resuelve la identidad.
    """
    libro = _libro_base()
    libro["Resumen"]["A2"], libro["Resumen"]["B2"] = "Proyecto", "Etapa"
    meta = _aplicar(libro, identidad=IDENTIDAD_ETAPA)
    assert libro["Resumen"]["B2"].value == "Etapa"
    assert meta["celdas"] == []
    assert meta["publicable"] is False

    libro2 = _libro_base()
    libro2["Resumen"]["A2"], libro2["Resumen"]["B2"] = "Proyecto:", "Etapa"
    meta2 = _aplicar(libro2, identidad=IDENTIDAD_ETAPA)
    assert libro2["Resumen"]["B2"].value == "Etapa"
    assert [c["celda_valor"] for c in meta2["celdas"]] == ["B2"]
    assert meta2["publicable"] is True


def test_n16_cabecera_vacia_complementaria_no_afecta_campo_declarado():
    libro = _libro_base()
    libro["Resumen"]["A2"], libro["Resumen"]["B2"] = "Proyecto:", "Proyecto A"
    ficha = libro.create_sheet("Ficha")
    ficha["A1"], ficha["B1"] = "Proyecto", "Etapa"
    meta = _aplicar(libro)
    assert libro["Resumen"]["B2"].value == "Proyecto B"   # campo declarado
    assert (ficha["A1"].value, ficha["B1"].value) == ("Proyecto", "Etapa")
    assert meta["publicable"] is True
    assert meta["identidad_verificada"] is True


def test_n16_tabla_vacia_guardada_y_releida_hoja_complementaria(tmp_path):
    """Verificación con Excel temporal guardado y RELEÍDO (requisito 5)."""
    plantilla = tmp_path / "plantilla_n16.xlsx"
    libro = _libro_base()
    libro["Resumen"]["A2"], libro["Resumen"]["B2"] = "Proyecto:", "Proyecto A"
    ficha = libro.create_sheet("Ficha")
    ficha["A1"], ficha["B1"] = "Proyecto", "Etapa"
    libro.save(plantilla)
    hash_plantilla = _hash(plantilla)
    copia = tmp_path / "salida_n16.xlsx"
    copia.write_bytes(plantilla.read_bytes())
    releido = load_workbook(copia)
    meta = _aplicar(releido)
    releido.save(copia)
    final = load_workbook(copia)
    assert (final["Ficha"]["A1"].value, final["Ficha"]["B1"].value) \
        == ("Proyecto", "Etapa")
    assert final["Resumen"]["B2"].value == "Proyecto B"
    assert meta["publicable"] is True
    assert _hash(plantilla) == hash_plantilla   # original intacto


# ---------------------------------------------------------------------------
# Exportadores reales: empresarial y BMG (copias temporales + hashes)
# ---------------------------------------------------------------------------


def test_n16_empresarial_cabecera_vacia_declarada_genera_borrador(tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = tmp_path / "base_n16.xlsx"
    _plantilla(plantilla)
    libro = load_workbook(plantilla)
    resumen = libro["Resumen"]
    # Tabla vacía de dos columnas en zona limpia de la hoja DECLARADA.
    resumen["H1"], resumen["I1"] = "Proyecto", "Etapa"
    libro.save(plantilla)
    hash_plantilla = _hash(plantilla)
    oficial = tmp_path / "salida_n16.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla, oficial, [], identidad_proyecto=IDENTIDAD_B))
    assert not oficial.exists()
    assert "borrador_no_publicable" in resultado.name
    borrador = load_workbook(resultado)          # releer el Excel temporal
    assert (borrador["Resumen"]["H1"].value, borrador["Resumen"]["I1"].value) \
        == ("Proyecto", "Etapa")
    assert _hash(plantilla) == hash_plantilla    # plantilla original intacta


def test_n16_empresarial_cabecera_vacia_con_campo_explicito_publica(tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = tmp_path / "base_n16b.xlsx"
    _plantilla(plantilla)
    libro = load_workbook(plantilla)
    resumen = libro["Resumen"]
    resumen["H1"], resumen["I1"] = "Proyecto", "Etapa"     # tabla vacía
    resumen["A2"], resumen["B2"] = "Proyecto:", "Proyecto A"  # campo válido
    libro["Captura_pruebas"]["A1"] = "Reporte del Proyecto B"
    libro.save(plantilla)
    hash_plantilla = _hash(plantilla)
    oficial = tmp_path / "salida_n16b.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla, oficial, [], identidad_proyecto=IDENTIDAD_B))
    assert resultado == oficial and oficial.is_file()
    salida = load_workbook(oficial)
    assert (salida["Resumen"]["H1"].value, salida["Resumen"]["I1"].value) \
        == ("Proyecto", "Etapa")                 # cabecera intacta
    assert salida["Resumen"]["B2"].value == "Proyecto B"
    assert _hash(plantilla) == hash_plantilla


def test_n16_bmg_cabecera_vacia_no_bloquea_proyecto_explicito(tmp_path):
    from reporte_bmg import generar_reporte_bmg
    from configuracion import cargar_config, cargar_reglas
    from test_reporte_bmg_v7 import caso
    real = RAIZ / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    hash_real = _hash(real)
    copia = tmp_path / "bmg_n16.xlsx"
    copia.write_bytes(real.read_bytes())
    libro = load_workbook(copia)
    resumen = libro["Übersicht Standversuche"]
    # Cabecera «Projekt | Stufe» en zona vacía de la hoja declarada (la fila
    # ya contiene el título en A1: fila tabular ⇒ omisión sin bloqueo).
    resumen["Z1"], resumen["AA1"] = "Projekt", "Stufe"
    libro.save(copia)
    oficial = tmp_path / "bmg_oficial.xlsx"
    resultados = [caso(1001, "1ST", "HT", "UGL", 5.8)]
    salida = pathlib.Path(generar_reporte_bmg(
        copia, oficial, resultados, cargar_config(), cargar_reglas(),
        idioma="de", identidad_proyecto=IDENTIDAD_B))
    assert salida == oficial and oficial.exists()
    publicado = load_workbook(salida)
    hoja = publicado["Übersicht Standversuche"]
    assert (hoja["Z1"].value, hoja["AA1"].value) == ("Projekt", "Stufe")
    assert hoja["A1"].value == "Proyecto B"
    assert _hash(real) == hash_real              # plantilla del repo intacta
