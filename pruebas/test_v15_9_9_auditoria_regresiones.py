"""N16: tablas VACÍAS de dos columnas no son campos de identidad.

Ejecutar desde la raíz del repositorio (fuera de datos productivos):
  python -m pytest -q /ruta/PRUEBAS_ROJAS_V15_9_8_N16.py

El caso existe cuando el Excel preconfigurado todavía no contiene filas de
pruebas: «Proyecto | Etapa» sigue siendo cabecera, nunca un par editable.
"""
from openpyxl import Workbook, load_workbook
import pytest
from utilidades.identidad_excel import aplicar_identidad_a_salida

IDENTIDAD = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b",
             "origen": "declarado_usuario", "sin_asignar": False}


def aplicar(wb, idioma):
    return aplicar_identidad_a_salida(
        wb, IDENTIDAD, idioma=idioma,
        hojas_busqueda=["Resumen"], hoja_resumen="Resumen")


@pytest.mark.parametrize("idioma,etiqueta,cabecera", [
    ("es", "Proyecto", "Etapa"),
    ("en", "Project", "Stage"),
    ("de", "Projekt", "Stufe"),
    ("es", "Proyecto", "Caso"),
])
def test_n16_tabla_vacia_dos_columnas_preserva_cabecera_y_no_verifica(
        tmp_path, idioma, etiqueta, cabecera):
    """Cabecera sola, sin filas: ni edición, ni identidad verificada."""
    w = Workbook(); s = w.active; s.title = "Resumen"
    s["A2"] = etiqueta; s["B2"] = cabecera
    original = tmp_path / "plantilla.xlsx"; w.save(original)
    wb = load_workbook(original)
    meta = aplicar(wb, idioma)
    salida = tmp_path / "salida.xlsx"; wb.save(salida)
    after = load_workbook(salida)["Resumen"]
    assert after["A2"].value == etiqueta
    assert after["B2"].value == cabecera, "Se sobreescribió una cabecera de tabla vacía"
    assert not [x for x in meta["celdas"] if x["celda_valor"] == "B2"]
    assert meta["identidad_verificada"] is False
    assert meta["publicable"] is False
    assert load_workbook(original)["Resumen"]["B2"].value == cabecera


def test_n16_tabla_vacia_y_campo_explicito_preserva_tabla_y_publica():
    """Cabecera vacía + otro campo inequívoco: publicar, sin alterar cabecera."""
    w = Workbook(); s = w.active; s.title = "Resumen"
    s["A1"] = "Proyecto"; s["B1"] = "Etapa"
    s["A4"] = "Proyecto:"; s["B4"] = "Proyecto A"
    meta = aplicar(w, "es")
    assert s["B1"].value == "Etapa"
    assert s["B4"].value == "Proyecto B"
    assert meta["publicable"] is True
    assert not any(x["celda_valor"] == "B1" for x in meta["celdas"])


def test_n16_control_campo_unico_desnudo_funciona():
    w = Workbook(); s = w.active; s.title = "Resumen"
    s["A2"] = "Proyecto"; s["B2"] = "Proyecto A"
    meta = aplicar(w, "es")
    assert s["B2"].value == "Proyecto B" and meta["publicable"] is True


def test_n16_control_campo_explicito_funciona():
    w = Workbook(); s = w.active; s.title = "Resumen"
    s["A2"] = "Proyecto:"; s["B2"] = "Proyecto A"
    meta = aplicar(w, "es")
    assert s["B2"].value == "Proyecto B" and meta["publicable"] is True
