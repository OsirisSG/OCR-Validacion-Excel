from pathlib import Path

import pytest
from openpyxl import load_workbook

from configuracion import cargar_config, cargar_reglas
from flujo_empresarial import CLAVES_PLANTILLA
from reporte_bmg import generar_reporte_bmg


def _video(tiempo=20.5, *, estado="confirmado", frame=205, nombre="frontal.avi"):
    return {
        "id": nombre, "ruta": f"C:/Ensayos/{nombre}", "nombre": nombre, "tipo_camara": "FRONTAL",
        "fps_grabacion": 10000, "fps": 10000, "fps_fuente": "overlay_rec", "frame_cero": 0,
        "overlay_timestamp_inicial_ms": 0, "frame_evento": 62, "tiempo_ms": 6.2,
        "revision_manual": {"estado": "confirmado", "frame_index": 62, "tiempo_ms": 6.2},
        "zona_deteccion": {"bbox_norm": [.1, .2, .3, .4]},
        "despliegue_maximo": {
            "estado": estado, "frame_index": frame, "tiempo_ms": tiempo,
            "requiere_revision": estado != "confirmado", "fps": 10000,
            "resolucion_nativa_ms": .1, "fuente_tiempo": "reloj_overlay",
            "metodo": "confirmacion_manual", "usuario": "Persona revisora",
            "confirmado_en": "2026-10-09T13:00:00", "bbox_norm": [.0, .0, .8, .9],
            "advertencias": ["Revisar posibles oclusiones de la bolsa."],
        },
    }


def _fila(videos):
    return {
        "project_id": "P", "case_key": "P|1ST|123456", "tipo_st": "1ST",
        "test_number": "123456", "temperature_condition": "HT", "inflator_type": "UGL",
        "campos": {
            "test_number": {"valor": "123456"},
            "opening_time_ms": {"valor": "5,8 ms", "estado": "confirmado", "fuente": "manual"},
            "full_inflation_time_ms": {"valor": "18,5 ms", "estado": "confirmado"},
        },
        "videos": videos,
    }


def _exportar(tmp_path, fila, idioma="es"):
    plantilla = Path(__file__).resolve().parents[1] / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    salida = tmp_path / f"tiempos_{idioma}.xlsx"
    generar_reporte_bmg(plantilla, salida, [fila], cargar_config(), cargar_reglas(), idioma=idioma)
    return load_workbook(salida)


@pytest.mark.parametrize("idioma,hoja,estado", [
    ("es", "Evidencia de tiempos de vídeo", "Confirmado"),
    ("en", "Video timing evidence", "Confirmed"),
    ("de", "Video-Zeitnachweise", "Bestätigt"),
])
def test_exporta_maximo_separado_del_contrato_con_evidencia_multilingue(tmp_path, idioma, hoja, estado):
    libro = _exportar(tmp_path, _fila([_video()]), idioma)
    lista = libro["Versuchsmatrix"]
    assert len(CLAVES_PLANTILLA) == 36
    assert lista["X8"].value == 5.8
    assert lista["Z8"].value == 18.5
    assert lista["AK8"].value == 20.5
    assert lista["AL8"].value == estado
    assert lista["AK8"].data_type == lista["Z8"].data_type == lista["X8"].data_type == "n"
    evidencia = libro[hoja]
    assert evidencia["D2"].value == "C:/Ensayos/frontal.avi"
    assert evidencia["D2"].hyperlink is not None
    assert evidencia["F2"].value == 10000
    assert evidencia["G2"].value == "overlay_rec"
    assert evidencia["H2"].value == 0
    assert evidencia["I2"].value == 0
    assert evidencia["J2"].value == 62
    assert evidencia["K2"].value == 6.2
    assert evidencia["M2"].value == 205
    assert evidencia["N2"].value == 20.5
    assert evidencia["O2"].value == .1
    assert evidencia["P2"].value == "reloj_overlay"
    assert evidencia["Q2"].value == "[0.1, 0.2, 0.3, 0.4]"
    assert evidencia["R2"].value == "[0.0, 0.0, 0.8, 0.9]"
    assert evidencia["S2"].value == estado
    assert evidencia["T2"].value == "Persona revisora"
    assert evidencia["U2"].value == "2026-10-09T13:00:00"
    assert "oclusiones" in evidencia["W2"].value
    for celda in ("F2", "H2", "I2", "J2", "K2", "M2", "N2", "O2"):
        assert evidencia[celda].data_type == "n"


def test_candidato_maximo_solo_aparece_en_evidencia_y_conserva_manual(tmp_path):
    libro = _exportar(tmp_path, _fila([_video(estado="candidato")]))
    lista = libro["Versuchsmatrix"]
    assert lista["AK8"].value is None
    assert lista["AL8"].value == "Pendiente de revisión"
    assert lista["X8"].value == 5.8
    assert lista["Z8"].value == 18.5
    evidencia = libro["Evidencia de tiempos de vídeo"]
    assert evidencia["N2"].value == 20.5
    assert evidencia["S2"].value == "Pendiente de revisión"
    color = libro["Übersicht Standversuche"]["D16"].fill.fgColor.rgb[-6:].upper()
    assert color == cargar_config()["colores"]["verde"].lstrip("#").upper()


@pytest.mark.parametrize("segundo,esperado,estado", [
    (20.1, 20.1, "Confirmado"),
    (21.0, None, "Conflicto entre vídeos"),
])
def test_confirmaciones_coherentes_o_conflictivas_entre_camaras(tmp_path, segundo, esperado, estado):
    videos = [_video(20.0, frame=200), _video(segundo, frame=round(segundo * 10), nombre="lateral.avi")]
    libro = _exportar(tmp_path, _fila(videos))
    assert libro["Versuchsmatrix"]["AK8"].value == esperado
    assert libro["Versuchsmatrix"]["AL8"].value == estado
    assert libro["Evidencia de tiempos de vídeo"].max_row == 3


def test_no_hace_consenso_de_tiempos_distintos_sin_resolucion(tmp_path):
    videos = [_video(20), _video(20.01, nombre="lateral.avi")]
    for video in videos:
        video["despliegue_maximo"].update({"resolucion_nativa_ms": None, "fps": None})
    libro = _exportar(tmp_path, _fila(videos))
    assert libro["Versuchsmatrix"]["AK8"].value is None
    assert libro["Versuchsmatrix"]["AL8"].value == "Conflicto entre vídeos"


@pytest.mark.parametrize("estado_invalidado", [
    "revisar_apertura", "revisar_orden_temporal", "revisar_origen_temporal", "archivo_modificado",
])
def test_confirmacion_invalidada_en_otra_camara_bloquea_maximo_consolidado(tmp_path, estado_invalidado):
    videos = [_video(20.0, frame=200), _video(20.1, estado=estado_invalidado, frame=201, nombre="lateral.avi")]
    libro = _exportar(tmp_path, _fila(videos))
    assert libro["Versuchsmatrix"]["AK8"].value is None
    assert libro["Versuchsmatrix"]["AL8"].value == "Pendiente de revisión"
    evidencia = libro["Evidencia de tiempos de vídeo"]
    assert evidencia["S2"].value == "Confirmado"
    assert evidencia["S3"].value == "Evidencia no válida"
    assert evidencia["N3"].value == 20.1  # Conserva el tiempo previo como evidencia, no como resultado.
    assert "Evidencia no válida" in evidencia["W3"].value


def test_candidato_nunca_confirmado_no_bloquea_otro_maximo_humano_vigente(tmp_path):
    videos = [_video(20.0, frame=200), _video(30, estado="candidato", frame=300, nombre="lateral.avi")]
    videos[1]["despliegue_maximo"].pop("confirmado_en")
    videos[1]["despliegue_maximo"]["metodo"] = "propuesta_automatica"
    libro = _exportar(tmp_path, _fila(videos))
    assert libro["Versuchsmatrix"]["AK8"].value == 20.0
    assert libro["Versuchsmatrix"]["AL8"].value == "Confirmado"
    evidencia = libro["Evidencia de tiempos de vídeo"]
    assert evidencia["N3"].value == 30
    assert evidencia["S3"].value == "Pendiente de revisión"


def test_faltantes_y_tiempos_no_finitos_no_se_convierten_en_cero(tmp_path):
    video = _video(float("inf"), frame=None)
    video["despliegue_maximo"].update({"fps": float("nan"), "resolucion_nativa_ms": None})
    video.update({"fps": None, "fps_grabacion": None, "frame_cero": None,
                  "frame_evento": None, "tiempo_ms": float("nan"), "overlay_timestamp_inicial_ms": None})
    video["revision_manual"] = {}
    fila = _fila([video])
    fila["campos"]["opening_time_ms"]["valor"] = "NaN"
    fila["campos"]["full_inflation_time_ms"]["valor"] = None
    libro = _exportar(tmp_path, fila)
    for celda in ("X8", "Z8", "AK8"):
        assert libro["Versuchsmatrix"][celda].value is None
    evidencia = libro["Evidencia de tiempos de vídeo"]
    for celda in ("F2", "H2", "I2", "J2", "K2", "M2", "N2", "O2"):
        assert evidencia[celda].value is None
    assert evidencia["S2"].value == "Evidencia no válida"


def test_exportacion_incorpora_revision_persistida_sobre_copia(tmp_path, monkeypatch):
    import maximos_video
    fila = _fila([])

    def incorporar(copia):
        copia["videos"] = [_video()]

    monkeypatch.setattr(maximos_video, "incorporar_maximos", incorporar)
    libro = _exportar(tmp_path, fila)
    assert libro["Versuchsmatrix"]["AK8"].value == 20.5
    assert fila["videos"] == []


def test_evidencia_conserva_t0_negativo_y_revision_sin_apertura(tmp_path):
    video = _video()
    video.pop("overlay_timestamp_inicial_ms")
    video["despliegue_maximo"]["timestamp_inicial_ms"] = -2.5
    video["revision_manual"] = {"estado": "sin_evento"}
    libro = _exportar(tmp_path, _fila([video]))
    ws = libro["Evidencia de tiempos de vídeo"]
    assert ws["I2"].value == -2.5
    assert ws["J2"].value is None and ws["K2"].value is None
    assert ws["L2"].value == "Sin apertura (revisión manual)"
    assert ws["X2"].value == "revision_manual"


def test_guia_queda_documentada_y_maximo_no_es_inflado_completo(tmp_path):
    libro = _exportar(tmp_path, _fila([]))
    ws = libro["Diccionario de datos"]
    reglas = {ws.cell(fila, 1).value: ws.cell(fila, 3).value for fila in range(2, ws.max_row + 1)}
    assert "abombamiento" in reglas["opening_time_ms"]
    assert "EDAG" in reglas["week"] and "WWYYNN" in reglas["week"]
    assert "VW AG" in reglas["test_place"]
    assert "DMY" in reglas["module_date"] and "conflicto" in reglas["module_date"]
    assert "predeterminados" in reglas["engineering_level"]
    assert "máximo despliegue" in reglas["full_inflation_time_ms"]
    assert "maximum_visible_deployment_ms" in reglas
