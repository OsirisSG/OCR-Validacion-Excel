"""Tres escenarios nuevos de auditoría (pruebas ROJAS con V15.9.4).

Este archivo debe colocarse temporalmente bajo pruebas/ para ejecutarlo con
PYTHONPATH=. python -m pytest pruebas/test_... -q; no toca datos reales.
"""
from openpyxl import Workbook
from utilidades.identidad_excel import aplicar_identidad_a_salida

IDENTIDAD_B = {
    'nombre_proyecto': 'Proyecto B',
    'project_id': 'proyecto-b',
    'origen': 'declarado_usuario',
    'sin_asignar': False,
}


def test_n7_otro_proyecto_en_hoja_adicional_no_debe_publicarse():
    wb = Workbook()
    wb.active.title = 'Resumen'
    wb.create_sheet('Listado')
    otra = wb.create_sheet('Indicadores')
    wb['Resumen']['A2'] = 'Proyecto:'
    wb['Resumen']['B2'] = 'Proyecto A'
    otra['A1'] = 'Reporte del Proyecto A'  # hoja visible, no declarada como resumen/lista
    meta = aplicar_identidad_a_salida(
        wb, IDENTIDAD_B, hojas_busqueda=['Resumen', 'Listado'], hoja_resumen='Resumen'
    )
    assert meta['publicable'] is False, (
        'La identidad anterior aparece en una hoja visible de la salida pero se publicó')


def test_n8_titulo_del_proyecto_correcto_con_anio_no_debe_bloquearse():
    wb = Workbook()
    wb.active.title = 'Resumen'
    wb['Resumen']['A1'] = 'Reporte del Proyecto B — Año 2026'
    wb['Resumen']['A2'] = 'Proyecto:'
    wb['Resumen']['B2'] = 'Proyecto A'  # campo etiquetado se actualiza a B
    meta = aplicar_identidad_a_salida(
        wb, IDENTIDAD_B, hojas_busqueda=['Resumen'], hoja_resumen='Resumen'
    )
    assert meta['publicable'] is True, (
        'Título válido con proyecto correcto y fecha fue declarado contradictorio')


def test_n9_sin_asignar_no_equivale_a_identidad_verificada():
    wb = Workbook()
    wb.active.title = 'Resumen'
    wb['Resumen']['A1'] = 'Reporte de ensayos sin asignación'
    identidad = {'origen':'declarado_sin_asignar', 'sin_asignar':True,
                 'project_id':'sin-asignar'}
    meta = aplicar_identidad_a_salida(
        wb, identidad, hojas_busqueda=['Resumen'], hoja_resumen='Resumen'
    )
    assert meta['identidad_verificada'] is False, (
        'Sin asignar es pendiente, no una identidad confirmada')
