
from pathlib import Path
from openpyxl import load_workbook

from configuracion import cargar_config, cargar_reglas
from flujo_empresarial import CLAVES_PLANTILLA
from reporte_bmg import HEADERS, generar_reporte_bmg
from utilidades.consolidacion import consolidar_documentos
from utilidades.identidad import identidad_caso


def _fila(etapa="1ST", numero="232541", ejec="2026-09-21T10:00:00"):
    project="p"
    return {
        "project_id": project, "tipo_st": etapa, "test_number": numero,
        "case_key": identidad_caso(project, etapa, numero, numero, f"{etapa}/{numero}"),
        "ultima_ejecucion": ejec,
        "temperature_condition": "HT", "inflator_type": "UGL",
        "campos": {
            "test_number": {"valor": numero},
            "temperature_condition": {"valor": "HT"},
            "inflator_type": {"valor": "UGL"},
            "opening_time_ms": {"valor": 5.8},
            "pab_comments": {"valor": "base"},
        },
    }


def test_correccion_mas_nueva_gana_aunque_la_ejecucion_sea_mas_vieja():
    a=_fila(ejec="2026-09-21T12:00:00")
    a["correcciones_manual_campos"]={"pab_comments":{"valor":"CORRECCION NUEVA","corregido_en":"2026-09-21T13:00:00","corregido_por":"u"}}
    b=_fila(ejec="2026-09-21T14:00:00")
    b["correcciones_manual_campos"]={"pab_comments":{"valor":"CORRECCION VIEJA","corregido_en":"2026-09-21T11:00:00","corregido_por":"u"}}
    r=consolidar_documentos([{"resultados":[a]},{"resultados":[b]}])["resultados"][0]
    assert r["campos"]["pab_comments"]["valor"] == "CORRECCION NUEVA"
    assert r["correcciones_manual_campos"]["pab_comments"]["corregido_en"] == "2026-09-21T13:00:00"


def test_empate_distinto_se_marca_como_conflicto():
    a=_fila(); b=_fila(ejec="2026-09-21T11:00:00")
    marca="2026-09-21T15:00:00"
    a["correcciones_manual_campos"]={"pab_comments":{"valor":"A","corregido_en":marca}}
    b["correcciones_manual_campos"]={"pab_comments":{"valor":"B","corregido_en":marca}}
    r=consolidar_documentos([{"resultados":[a]},{"resultados":[b]}])["resultados"][0]
    assert r["conflictos_correcciones_manual"]
    assert r["campos"]["pab_comments"]["requiere_revision"] is True


def test_todos_los_encabezados_tienen_es_en_de():
    assert len(CLAVES_PLANTILLA) == 36
    assert set(CLAVES_PLANTILLA) <= set(HEADERS)
    for clave in CLAVES_PLANTILLA:
        assert all(HEADERS[clave].get(lang) for lang in ("es","en","de"))


def test_caso_fuera_de_las_siete_combinaciones_no_desaparece(tmp_path):
    raiz=Path(__file__).resolve().parents[1]
    f=_fila(numero="999999")
    f["temperature_condition"]="CUSTOM"; f["inflator_type"]="X1"
    f["campos"]["temperature_condition"]={"valor":"CUSTOM"}
    f["campos"]["inflator_type"]={"valor":"X1"}
    salida=tmp_path/"reporte.xlsx"
    generar_reporte_bmg(raiz/"plantillas"/"reporte_bmg_1st_2st.xlsx", salida, [f], cargar_config(), cargar_reglas(), idioma="es")
    wb=load_workbook(salida, data_only=False)
    ws=wb["Übersicht Standversuche"]
    valores=[ws.cell(r,c).value for r in range(1,ws.max_row+1) for c in range(1,11)]
    assert "Casos sin clasificación de matriz" in valores
    visibles={str(v) for v in valores if v is not None}
    assert "999999" in visibles
    assert "232550" not in visibles and "222510" not in visibles, "No deben quedar IDs de ejemplo de la plantilla"
