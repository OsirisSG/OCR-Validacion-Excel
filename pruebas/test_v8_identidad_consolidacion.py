from copy import deepcopy
from pathlib import Path

from configuracion import cargar_config, cargar_reglas
from flujo_empresarial import interpretar_fecha_impresa, normalizar_serial
from reporte_bmg import generar_reporte_bmg
from utilidades.consolidacion import consolidar_documentos
from utilidades.identidad import identidad_caso


def _fila(project, stage, num, valor="A"):
    return {
        "project_id": project,
        "tipo_st": stage,
        "test_number": str(num),
        "case_key": identidad_caso(project, stage, str(num), f"{num} caso", f"{stage}/{num}"),
        "ultima_ejecucion": "2026-09-20T10:00:00",
        "campos": {"test_number": {"valor": str(num)}, "opening_time_ms": {"valor": 5.5},
                   "pab_comments": {"valor": valor}},
    }


def test_identidad_no_mezcla_mismo_numero_entre_proyecto_y_etapa():
    ids = {
        identidad_caso("A", "1ST", "232541", "caso", "HT/caso"),
        identidad_caso("A", "2ST", "232541", "caso", "HT/caso"),
        identidad_caso("B", "1ST", "232541", "caso", "HT/caso"),
    }
    assert len(ids) == 3
    assert all("/" not in x and "\\" not in x for x in ids)


def test_consolidacion_es_idempotente_y_no_duplica_etapas():
    a = {"resultados": [_fila("P", "1ST", 232541)]}
    b = {"resultados": [_fila("P", "2ST", 232541)]}
    uno = consolidar_documentos([a, b])
    dos = consolidar_documentos([uno, deepcopy(uno)])
    assert len(uno["resultados"]) == 2
    assert len(dos["resultados"]) == 2
    assert set(dos["tipos_st"]) == {"1ST", "2ST"}


def test_serial_preserva_asteriscos_espacios_y_ceros():
    assert normalizar_serial("  *00 ab 12* ") == "*00 AB 12*"
    assert normalizar_serial("0000123") == "0000123"


def test_fecha_ambigua_no_se_inventa_y_fecha_iso_si_se_interpreta():
    amb = interpretar_fecha_impresa("25-06-19")
    iso = interpretar_fecha_impresa("2025-06-19")
    assert amb["estado"] == "ambigua" and amb["interpretacion"] is None
    assert iso["estado"] == "interpretada" and iso["interpretacion"] == "2025-06-19"


def test_reporte_no_recorta_mas_de_50_casos_por_etapa(tmp_path):
    raiz = Path(__file__).resolve().parents[1]
    plantilla = raiz / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    resultados = []
    for i in range(65):
        f = _fila("P", "1ST", 100000 + i)
        f.update({"temperature_condition": "HT", "inflator_type": "UGL"})
        f["campos"].update({"temperature_condition": {"valor": "HT"}, "inflator_type": {"valor": "UGL"}})
        resultados.append(f)
    for i in range(3):
        f = _fila("P", "2ST", 200000 + i)
        f.update({"temperature_condition": "RT", "inflator_type": "NOM"})
        f["campos"].update({"temperature_condition": {"valor": "RT"}, "inflator_type": {"valor": "NOM"}})
        resultados.append(f)
    salida = tmp_path / "dinamico.xlsx"
    generar_reporte_bmg(plantilla, salida, resultados, cargar_config(), cargar_reglas(), idioma="en")
    from openpyxl import load_workbook
    wb = load_workbook(salida, read_only=False, data_only=False)
    ws = wb["Versuchsmatrix"]
    # 65 1ST desde fila 8; 2ST debe empezar después, no truncarse en fila 58.
    assert ws["C72"].value is not None
    assert ws["C73"].value is not None
    tr = wb["Traceability"]
    assert tr["A1"].value == "Project"

def test_consolidacion_migra_clave_v7_sin_duplicar():
    viejo = {"raiz": "/datos/ProyectoP", "resultados": [{
        "case_key": "ProyectoP|1ST|HT/232541 RDW NOM HT", "tipo_st":"1ST",
        "test_number":"232541", "nombre":"232541 RDW NOM HT", "ruta_relativa":"HT/232541 RDW NOM HT",
        "campos":{"test_number":{"valor":"232541"}}}]}
    nuevo = {"raiz": "/otra/ProyectoP", "resultados": [{
        "project_id":"proyectop", "case_key": identidad_caso("proyectop","1ST","232541","232541 RDW NOM HT","HT/232541 RDW NOM HT"),
        "tipo_st":"1ST", "test_number":"232541", "nombre":"232541 RDW NOM HT", "ruta_relativa":"HT/232541 RDW NOM HT",
        "ultima_ejecucion":"2026-09-20T12:00:00", "campos":{"test_number":{"valor":"232541"}}}]}
    assert len(consolidar_documentos([viejo,nuevo])["resultados"]) == 1
