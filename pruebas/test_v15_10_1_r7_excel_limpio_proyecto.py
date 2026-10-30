from pathlib import Path

from openpyxl import load_workbook

from configuracion import cargar_config, cargar_reglas
from reporte_bmg import generar_reporte_bmg
from test_reporte_bmg_v7 import caso


def _valores(libro):
    out=[]
    for ws in libro.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out.append(str(c.value))
    return out


def _identidad():
    return {
        "nombre_proyecto": "Ejemplo_VW",
        "project_id": "ejemplo_vw",
        "origen": "declarado_usuario",
        "seleccion_explicita": True,
        "sin_asignar": False,
    }


def test_r7_excel_limpia_ids_plantilla_y_usa_proyecto_y_colores(tmp_path):
    """R15 (SUPERSEDE documentado): desde V15.10 sólo una confirmación humana
    (VERIFICADO_HUMANO con referencia temporal validada) colorea el resumen;
    un campo OCR ``confirmado`` sin confirmación humana queda gris por
    abstención (PLAN_CORRECCION_R13 §Semáforo; puertas de publicación R10).
    La intención R7 se conserva: proyecto declarado en títulos, IDs de
    plantilla limpiados y colores verde/amarillo/rojo por criterio elegido."""
    root=Path(__file__).resolve().parents[1]
    plantilla=root/'plantillas'/'reporte_bmg_1st_2st.xlsx'
    out=tmp_path/'r7.xlsx'
    resultados=[
        caso(900001,'1ST','HT','UGL',5.8,manual=True),
        caso(900002,'1ST','HT','OGL',6.4,manual=True),
        caso(900003,'1ST','HT','NOM',7.2,manual=True),
    ]
    criterios={"verde_max":6.0,"amarillo_max":7.0,"origen":"elegido"}
    generar_reporte_bmg(plantilla,out,resultados,cargar_config(),cargar_reglas(),
                        criterios_apertura=criterios, identidad_proyecto=_identidad())
    wb=load_workbook(out,data_only=False)
    vals=_valores(wb)
    assert '232541' not in vals
    assert wb['Übersicht Standversuche']['A1'].value == 'Ejemplo_VW'
    assert wb['Versuchsmatrix']['A2'].value == 'Ejemplo_VW'
    assert wb['Rangos']['A2'].value == 'Ejemplo_VW'
    cfg=cargar_config()
    def rgb(cell): return (cell.fill.fgColor.rgb or '')[-6:].upper()
    assert rgb(wb['Übersicht Standversuche']['D7']) == cfg['colores']['verde'].lstrip('#').upper()
    assert rgb(wb['Übersicht Standversuche']['E7']) == cfg['colores']['amarillo'].lstrip('#').upper()
    assert rgb(wb['Übersicht Standversuche']['F7']) == cfg['colores']['rojo'].lstrip('#').upper()


def test_r7_segunda_corrida_no_arrastra_casos_de_otra_corrida(tmp_path):
    root=Path(__file__).resolve().parents[1]
    plantilla=root/'plantillas'/'reporte_bmg_1st_2st.xlsx'
    primero=tmp_path/'primero.xlsx'
    segundo=tmp_path/'segundo.xlsx'
    criterios={"verde_max":6.0,"amarillo_max":7.0,"origen":"elegido"}
    generar_reporte_bmg(plantilla,primero,[caso(910001,'1ST','HT','UGL',5.8), caso(910002,'1ST','HT','OGL',6.4)],
                        cargar_config(),cargar_reglas(),criterios_apertura=criterios,identidad_proyecto=_identidad())
    # Segunda corrida independiente sobre la misma plantilla institucional: no debe
    # heredar los casos de la corrida anterior.
    generar_reporte_bmg(plantilla,segundo,[caso(920001,'1ST','NT','NOM',5.5)],
                        cargar_config(),cargar_reglas(),criterios_apertura=criterios,identidad_proyecto=_identidad())
    wb=load_workbook(segundo,data_only=False)
    vals=_valores(wb)
    assert '910001' not in vals and '910002' not in vals
    assert '920001' in vals
