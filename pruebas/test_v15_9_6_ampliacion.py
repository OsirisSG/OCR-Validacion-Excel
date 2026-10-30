"""V15.9.6 · Ampliación de pruebas exigida por la auditoría (N10-N12).

Complementa `test_v15_9_6_auditoria_regresiones.py` (los 5 casos literales)
con los escenarios de `CORRECCIONES_OBLIGATORIAS_V_MATRIX_V15_9_6.md`:

- N10: mapeo estructural de pares etiqueta/valor en hojas fuera de
  ``hojas_busqueda`` (corrección específica registrada, bloqueos sin
  reemplazo, exenciones, «Sin asignar»), con plantillas reales BMG y
  empresarial y hash de plantilla intacto.
- N11: comparación del nombre COMPLETO normalizado (Beta/Beta Sur/Beta
  Norte, A/AB), variantes triviales (mayúsculas/tildes/espaciado), alias
  registrados explícitamente e ID distinto del nombre visible.
- N12: etapas/fases contextuales (Etapa/Stage/Fase/Phase/Stufe + 1ST/2ST/3RD/
  NACH/VOR) y pruebas negativas donde la etapa forma parte del nombre de OTRO
  proyecto.

Todo sintético (tmp_path); no toca datos reales ni de red.
"""
from __future__ import annotations

import hashlib
import pathlib
import sys

import pytest
from openpyxl import Workbook, load_workbook

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "pruebas"))

from test_plantilla_empresarial import _plantilla
from utilidades.identidad_excel import aplicar_identidad_a_salida

IDENTIDAD_B = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b",
               "origen": "declarado_usuario", "sin_asignar": False,
               "pendiente": False, "ejecucion_id": "ej-v1596"}
IDENTIDAD_BETA_SUR = {"nombre_proyecto": "Proyecto Beta Sur",
                      "project_id": "proyecto-beta-sur",
                      "origen": "declarado_usuario", "sin_asignar": False}
SIN_ASIGNAR = {"nombre_proyecto": None, "project_id": "sin-asignar",
               "origen": "sin_asignar_expreso", "sin_asignar": True,
               "pendiente": True}


def _hash(ruta) -> str:
    return hashlib.sha256(pathlib.Path(ruta).read_bytes()).hexdigest()


def _resumen_texto(libro, hoja="Resumen"):
    return "\n".join(str(c.value) for fila in libro[hoja].iter_rows()
                     for c in fila if c.value)


def _libro_etiquetado(titulo=None):
    libro = Workbook()
    libro.active.title = "Resumen"
    hoja = libro["Resumen"]
    if titulo is not None:
        hoja["A1"] = titulo
    hoja["A2"] = "Proyecto:"
    hoja["B2"] = "Proyecto A"  # valor histórico que el mapeo sustituye
    return libro


def _plantilla_publicable(ruta, titulo="Reporte del Proyecto B"):
    _plantilla(ruta)
    libro = load_workbook(ruta)
    captura = libro["Captura_pruebas"]
    captura["A1"] = titulo
    captura["A2"] = "Proyecto:"
    captura["B2"] = "Proyecto A"
    libro.save(ruta)
    return ruta


# ---------------------------------------------------------------------------
# N10 · mapeo estructural etiqueta/valor en todas las hojas elegibles
# ---------------------------------------------------------------------------


def test_n10_par_en_hoja_adicional_se_corrige_y_se_registra():
    libro = _libro_etiquetado("Reporte del Proyecto B")
    otra = libro.create_sheet("Indicadores")
    otra["A1"] = "Proyecto:"
    otra["B1"] = "A"
    meta = aplicar_identidad_a_salida(
        libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert otra["B1"].value == "Proyecto B"
    assert any(c["hoja"] == "Indicadores" and c["celda_valor"] == "B1"
               and c["valor_anterior"] == "A"
               for c in meta["celdas_complementarias"])
    assert meta["publicable"] is True
    assert meta["identidad_verificada"] is True
    # El bloque de identidad registra la celda corregida (trazabilidad).
    assert "Indicadores!B1" in _resumen_texto(libro)


def test_n10_par_con_valor_vacio_se_rellena():
    libro = _libro_etiquetado("Reporte del Proyecto B")
    otra = libro.create_sheet("Indicadores")
    otra["A1"] = "Proyecto:"  # campo de proyecto sin rellenar
    meta = aplicar_identidad_a_salida(
        libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert otra["B1"].value == "Proyecto B"
    assert meta["celdas_complementarias"][0]["valor_anterior"] is None


def test_n10_par_con_valor_numerico_bloquea_sin_reemplazar():
    libro = _libro_etiquetado("Reporte del Proyecto B")
    otra = libro.create_sheet("Indicadores")
    otra["A1"] = "Proyecto:"
    otra["B1"] = 2024  # estructura ambigua: no se reemplaza
    meta = aplicar_identidad_a_salida(
        libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is False
    assert otra["B1"].value == 2024
    assert any(c.get("estructural") and c["celda"] == "B1"
               for c in meta["contradicciones"])


def test_n10_fila_tabular_con_etiqueta_explicita_bloquea():
    libro = _libro_etiquetado("Reporte del Proyecto B")
    tabla = libro.create_sheet("Listado")
    tabla["A1"] = "Proyecto:"
    tabla["B1"] = "A"
    tabla["C1"] = "Etapa"  # fila con más contenido: estructura ambigua
    tabla["D1"] = "1ST"
    meta = aplicar_identidad_a_salida(
        libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is False
    assert tabla["B1"].value == "A"  # nada reemplazado en fila ambigua
    assert any(c.get("estructural") for c in meta["contradicciones"])
    # Si el valor YA coincide con la identidad, la fila tabular no bloquea.
    libro2 = _libro_etiquetado("Reporte del Proyecto B")
    tabla2 = libro2.create_sheet("Listado")
    tabla2["A1"] = "Proyecto:"
    tabla2["B1"] = "Proyecto B"
    tabla2["C1"] = "Etapa"
    meta2 = aplicar_identidad_a_salida(
        libro2, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta2["publicable"] is True


def test_n10_hoja_exenta_conserva_su_par_historico_intacto():
    libro = _libro_etiquetado("Reporte del Proyecto B")
    historico = libro.create_sheet("Historico")
    historico["A1"] = "Proyecto:"
    historico["B1"] = "Proyecto A"
    historico["A2"] = "Reporte del Proyecto A"
    meta = aplicar_identidad_a_salida(
        libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen",
        hojas_exentas=["Historico"])
    # La exención documentada cubre el par estructural Y el texto libre.
    assert historico["B1"].value == "Proyecto A"
    assert historico["A2"].value == "Reporte del Proyecto A"
    assert meta["publicable"] is True
    assert "Hojas exentas" in _resumen_texto(libro)


def test_n10_sin_asignar_rellena_el_par_complementario_como_pendiente():
    libro = _libro_etiquetado("Reporte de ensayos")
    otra = libro.create_sheet("Indicadores")
    otra["A1"] = "Proyecto:"
    otra["B1"] = "Proyecto A"
    meta = aplicar_identidad_a_salida(
        libro, dict(SIN_ASIGNAR), hojas_busqueda=["Resumen"],
        hoja_resumen="Resumen")
    assert otra["B1"].value == "Sin asignar (pendiente)"
    assert meta["estado_identidad"] == "pendiente_sin_asignar"
    assert meta["identidad_verificada"] is False
    assert meta["publicable"] is True  # entregable pendiente (P02)


def test_n10_empresarial_par_en_hoja_adicional_publica_y_plantilla_intacta(tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = _plantilla_publicable(tmp_path / "base.xlsx")
    libro = load_workbook(plantilla)
    extra = libro.create_sheet("Indicadores")
    extra["A1"] = "Proyecto:"
    extra["B1"] = "A"
    libro.save(plantilla)
    hash_plantilla = _hash(plantilla)
    oficial = tmp_path / "salida.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla, oficial, [], identidad_proyecto=IDENTIDAD_B))
    assert resultado == oficial and oficial.is_file()
    salida = load_workbook(oficial)
    assert salida["Indicadores"]["B1"].value == "Proyecto B"
    assert salida["Captura_pruebas"]["B2"].value == "Proyecto B"
    assert _hash(plantilla) == hash_plantilla


def test_n10_bmg_real_par_en_rangos_publica_con_identidad_coherente(tmp_path):
    from reporte_bmg import generar_reporte_bmg
    from configuracion import cargar_config, cargar_reglas
    from test_reporte_bmg_v7 import caso
    real = RAIZ / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    hash_real = _hash(real)
    resultados = [caso(1001, "1ST", "HT", "UGL", 5.8)]
    copia = tmp_path / "bmg.xlsx"
    copia.write_bytes(real.read_bytes())
    libro = load_workbook(copia)
    # Campo de proyecto explícito en «Rangos» (fuera de hojas_busqueda) y
    # título coherente con etapa 1ST en el resumen.
    libro["Rangos"]["A1"] = "Projekt:"
    libro["Übersicht Standversuche"]["Z1"] = "Bericht Projekt B — Stufe 1ST"
    libro.save(copia)
    oficial = tmp_path / "bmg_oficial.xlsx"
    salida = pathlib.Path(generar_reporte_bmg(
        copia, oficial, resultados, cargar_config(), cargar_reglas(),
        idioma="de", identidad_proyecto=IDENTIDAD_B))
    assert salida == oficial and oficial.is_file()
    resultado = load_workbook(oficial)
    assert resultado["Rangos"]["A1"].value == "Projekt:"  # etiqueta intacta
    assert resultado["Übersicht Standversuche"]["Z1"].value \
        == "Bericht Projekt B — Stufe 1ST"  # título intacto
    assert _hash(real) == hash_real  # plantilla del repo intacta


# ---------------------------------------------------------------------------
# N11 · nombre completo normalizado, alias explícitos, sin coincidencia parcial
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("titulo,debe_publicar", [
    ("Reporte del Proyecto Beta Sur", True),
    ("Reporte del proyecto beta sur", True),      # mayúsculas/espaciado
    ("REPORTE DEL PROYECTO BETA   SUR", True),    # espaciado múltiple
    ("Reporte del Proyecto Beta", False),         # coincidencia parcial
    ("Reporte del Proyecto Sur", False),          # coincidencia parcial
    ("Reporte del Proyecto Beta Norte", False),   # otro proyecto próximo
])
def test_n11_nombre_compuesto_exige_coincidencia_completa(titulo, debe_publicar):
    libro = _libro_etiquetado(titulo)
    meta = aplicar_identidad_a_salida(
        libro, IDENTIDAD_BETA_SUR, hojas_busqueda=["Resumen"],
        hoja_resumen="Resumen")
    assert meta["publicable"] is debe_publicar, titulo
    if not debe_publicar:
        assert meta["contradicciones"]
        assert "coincidencia parcial" in meta["contradicciones"][0]["detalle"]


def test_n11_variantes_triviales_de_tildes_se_aceptan():
    identidad = dict(IDENTIDAD_B, nombre_proyecto="Proyecto Ñu Ägil",
                     project_id="proyecto-nu-agil")
    for titulo in ("Reporte del Proyecto Ñu Ägil",
                   "Reporte del Proyecto NU AGIL",
                   "Reporte del Proyecto ñu ägil"):
        libro = _libro_etiquetado(titulo)
        meta = aplicar_identidad_a_salida(
            libro, identidad, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
        assert meta["publicable"] is True, titulo


def test_n11_alias_registrado_explicitamente_valida_y_no_registrado_no():
    identidad = dict(IDENTIDAD_B, aliases=["PB-7", "Beta Siete"])
    libro = _libro_etiquetado("Reporte del Proyecto PB-7")
    meta = aplicar_identidad_a_salida(
        libro, identidad, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is True
    libro2 = _libro_etiquetado("Reporte del Proyecto PB-8")  # alias NO registrado
    meta2 = aplicar_identidad_a_salida(
        libro2, identidad, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta2["publicable"] is False


def test_n11_id_tecnico_distinto_del_nombre_visible_valida_completo():
    identidad = dict(IDENTIDAD_B, nombre_proyecto="Proyecto B",
                     project_id="PRJ-2026-07")
    for titulo in ("Reporte del Proyecto PRJ-2026-07",
                   "Reporte del Proyecto B — Año 2026"):
        libro = _libro_etiquetado(titulo)
        meta = aplicar_identidad_a_salida(
            libro, identidad, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
        assert meta["publicable"] is True, titulo
    # Un ID parecido pero distinto no valida.
    libro = _libro_etiquetado("Reporte del Proyecto PRJ-2026-70")
    meta = aplicar_identidad_a_salida(
        libro, identidad, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is False


# ---------------------------------------------------------------------------
# N12 · etapas/fases contextuales del flujo V-Matrix
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("titulo", [
    "Reporte del Proyecto B — Etapa 1ST",
    "Reporte del Proyecto B — Etapa 2ST",
    "Reporte del Proyecto B | Fase NACH",
    "Reporte del Proyecto B (Fase VOR)",
    "Report of Project B — Stage 1ST",
    "Report of Project B — Phase NACH",
    "Bericht Projekt B — Stufe 2ST",
    "Bericht Projekt B — Phase VOR",
    "Reporte del Proyecto B - Etapa 3RD - Año 2026",
    "Reporte del Proyecto B — Etapa 1ST, revisión 2",
])
def test_n12_etapa_o_fase_contextual_no_bloquea(titulo):
    libro = _libro_etiquetado(titulo)
    meta = aplicar_identidad_a_salida(
        libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["contradicciones"] == [], titulo
    assert meta["publicable"] is True


@pytest.mark.parametrize("titulo", [
    "Reporte del Proyecto 1ST",   # la etapa como nombre de OTRO proyecto
    "Reporte del Proyecto NACH",
    "Reporte del Proyecto VOR",
])
def test_n12_valor_de_etapa_suelto_no_es_permiso_global(titulo):
    libro = _libro_etiquetado(titulo)
    meta = aplicar_identidad_a_salida(
        libro, IDENTIDAD_B, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is False, titulo
    assert meta["contradicciones"]


def test_n12_identidad_con_etapa_legitima_en_su_nombre_valida():
    identidad = dict(IDENTIDAD_B, nombre_proyecto="Proyecto B 1ST",
                     project_id="proyecto-b-1st")
    libro = _libro_etiquetado("Reporte del Proyecto B 1ST")
    meta = aplicar_identidad_a_salida(
        libro, identidad, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta["publicable"] is True
    # Y el proyecto homónimo sin etapa NO valida contra «Proyecto B 1ST».
    libro2 = _libro_etiquetado("Reporte del Proyecto B")
    meta2 = aplicar_identidad_a_salida(
        libro2, identidad, hojas_busqueda=["Resumen"], hoja_resumen="Resumen")
    assert meta2["publicable"] is False


# ---------------------------------------------------------------------------
# Integración exigida: BMG/empresarial ES/EN/DE, histórico A + usuario B
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("idioma,titulo", [
    ("es", "Reporte del Proyecto Beta Sur — Etapa 1ST"),
    ("en", "Report of Project Beta Sur — Stage 1ST"),
    ("de", "Bericht Projekt Beta Sur — Stufe 1ST"),
])
def test_integracion_empresarial_compuesto_etapa_tres_idiomas(
        idioma, titulo, tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = _plantilla_publicable(tmp_path / f"base_{idioma}.xlsx",
                                      titulo=titulo)
    oficial = tmp_path / f"salida_{idioma}.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla, oficial, [], identidad_proyecto=IDENTIDAD_BETA_SUR,
        idioma=idioma))
    assert resultado == oficial and oficial.is_file(), idioma
    assert load_workbook(oficial)["Captura_pruebas"]["B2"].value \
        == "Proyecto Beta Sur"


def test_integracion_empresarial_historico_a_con_usuario_b_bloquea(tmp_path):
    """Plantilla histórica del proyecto A + usuario declara B ⇒ borrador."""
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = _plantilla_publicable(tmp_path / "historica_a.xlsx",
                                      titulo="Reporte del Proyecto A — Etapa 1ST")
    hash_plantilla = _hash(plantilla)
    oficial = tmp_path / "salida_b.xlsx"
    resultado = pathlib.Path(generar_desde_plantilla(
        plantilla, oficial, [], identidad_proyecto=IDENTIDAD_B))
    assert not oficial.exists()
    assert "borrador_no_publicable" in resultado.name
    libro = load_workbook(resultado)
    assert libro["Captura_pruebas"]["A1"].value \
        == "Reporte del Proyecto A — Etapa 1ST"  # intacto
    assert libro["Captura_pruebas"]["B2"].value == "Proyecto B"  # mapeada
    assert "REFERENCIA CONTRADICTORIA" in _resumen_texto(libro)
    assert _hash(plantilla) == hash_plantilla
