"""V15.9.2 · C7: proyecto elegido explícitamente por el usuario (P01-P12).

Cubre la identidad autoritativa (declarado > «Sin asignar» expreso > config
legacy > carpeta legacy-pendiente), el registro de proyectos con detección de
colisiones, la reanudación sin cambio silencioso, la independencia de
``test_reason``/``tipo_st`` y la identificación de los Excel de SALIDA en
celdas de metadatos identificadas, sin tocar jamás la plantilla original.

Ejecutar: python -m pytest pruebas/test_v15_9_2_proyecto_c7.py -q
"""
from __future__ import annotations

import hashlib
import importlib
import pathlib
import sys
import threading

import pytest

RAIZ = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "pruebas"))

from openpyxl import load_workbook

import flujo_empresarial as fe
import utilidades.proyectos as registro
from estructura import mapear_estructura
from utilidades.identidad import resolver_identidad_proyecto
from test_plantilla_empresarial import _plantilla


def _backend():
    return importlib.import_module("dashboard.backend.app")


def _hash(ruta: pathlib.Path) -> str:
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


IDENTIDAD_B = {"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b",
               "origen": "declarado_usuario", "sin_asignar": False,
               "pendiente": False, "responsable": "ana",
               "ejecucion_id": "ej-001", "fecha": "2026-10-12T10:00:00"}


# ---------------------------------------------------------------------------
# Identidad autoritativa y precedencia
# ---------------------------------------------------------------------------


def test_c7_precedencia_declarado_sobre_config_y_carpeta(tmp_path):
    identidad = resolver_identidad_proyecto(
        declarado={"nombre_proyecto": "Proyecto B", "usuario": "ana"},
        raiz=tmp_path / "Proyecto_A", configurado="proyecto-viejo")
    assert identidad["project_id"] == "proyecto-b"
    assert identidad["nombre_proyecto"] == "Proyecto B"
    assert identidad["origen"] == "declarado_usuario"
    assert identidad["pendiente"] is False


def test_c7_sin_asignar_expreso_queda_pendiente(tmp_path):
    identidad = resolver_identidad_proyecto(
        declarado={"sin_asignar": True, "usuario": "ana"},
        raiz=tmp_path / "Proyecto_A", configurado="proyecto-viejo")
    assert identidad["sin_asignar"] is True
    assert identidad["project_id"] == "sin-asignar"
    assert identidad["nombre_proyecto"] is None
    assert identidad["pendiente"] is True
    assert identidad["origen"] == "sin_asignar_expreso"


def test_c7_checkpoint_manda_sobre_declarado():
    identidad = resolver_identidad_proyecto(
        declarado={"nombre_proyecto": "Otro"},
        checkpoint={"nombre_proyecto": "Proyecto B", "project_id": "proyecto-b"})
    assert identidad["project_id"] == "proyecto-b"
    assert identidad["origen"] == "checkpoint"


def test_c7_legacy_carpeta_conserva_id_pero_queda_pendiente(tmp_path):
    """P11: cargas legacy conservan su identidad técnica sin adopción visible."""
    identidad = resolver_identidad_proyecto(raiz=tmp_path / "Proyecto_A")
    assert identidad["project_id"] == "proyecto_a"
    assert identidad["origen"] == "nombre_carpeta_legacy"
    assert identidad["nombre_proyecto"] is None
    assert identidad["pendiente"] is True


def test_c7_mapear_estructura_usa_proyecto_declarado(tmp_path):
    lote = tmp_path / "CarpetaDelLote"
    lote.mkdir()
    config = {"empresarial": {"proyecto": {"nombre_proyecto": "Proyecto B",
                                           "usuario": "ana"}}}
    estructura = mapear_estructura(lote, config)
    assert estructura["proyecto"]["project_id"] == "proyecto-b"
    assert estructura["proyecto"]["nombre_proyecto"] == "Proyecto B"
    # Sin declaración: legacy por carpeta, pendiente de confirmación.
    estructura2 = mapear_estructura(lote, {})
    assert estructura2["proyecto"]["project_id"] == "carpetadellote"
    assert estructura2["proyecto"]["origen"] == "nombre_carpeta_legacy"
    assert estructura2["proyecto"]["pendiente"] is True


def test_c7_p03_case_key_separado_por_proyecto(tmp_path):
    """P03: el mismo árbol con proyectos distintos no comparte case_key."""
    raiz = tmp_path / "lote"
    caso = raiz / "1ST" / "TOR1_HT"
    caso.mkdir(parents=True)
    key_a = fe.case_key(raiz, "1ST", caso, "T1", "proyecto-a")
    key_b = fe.case_key(raiz, "1ST", caso, "T1", "proyecto-b")
    assert key_a != key_b
    assert key_a.startswith("proyecto-a|") and key_b.startswith("proyecto-b|")


def test_c7_p05_tipo_st_independiente_del_proyecto(tmp_path):
    """P05: proyecto y etapa 1ST/2ST son dimensiones separadas."""
    raiz = tmp_path / "lote"
    caso = raiz / "1ST" / "TOR1_HT"
    caso.mkdir(parents=True)
    key_1 = fe.case_key(raiz, "1ST", caso, "T1", "proyecto-x")
    key_2 = fe.case_key(raiz, "2ST", caso, "T1", "proyecto-x")
    assert key_1 != key_2
    identidad = resolver_identidad_proyecto(declarado={"nombre_proyecto": "X"})
    assert "tipo_st" not in identidad  # la identidad no decide la etapa


def test_c7_p06_test_reason_no_copia_el_proyecto():
    """P06: test_reason conserva SU evidencia OCR; no copia nombre_proyecto."""
    imagen = {"ruta": "foto.jpg", "fuente": "ocr", "confianza_ocr": 0.9}
    contexto = {"project_id": "proyecto-b", "nombre_proyecto": "Proyecto B",
                "tipo_st": "1ST"}
    candidatos = fe.extraer_candidatos_texto(
        "PROYECTO: CrashTestReason\nTEXTO ADICIONAL", imagen, contexto)
    reasons = [c for c in candidatos if c["clave"] == "test_reason"]
    assert reasons and reasons[0]["valor_normalizado"] == "CrashTestReason"
    assert not any("Proyecto B" in str(c.get("valor_normalizado"))
                   for c in candidatos if c["clave"] == "test_reason")


# ---------------------------------------------------------------------------
# Registro de proyectos y colisiones (P12)
# ---------------------------------------------------------------------------


def test_c7_registro_crea_lista_y_persiste(tmp_path):
    ruta = tmp_path / "proyectos.json"
    creado = registro.registrar_proyecto("Proyecto B", responsable="ana", ruta=ruta)
    assert creado["proyecto_id"] == "proyecto-b"
    assert ruta.is_file()
    assert registro.listar_proyectos(ruta)[0]["nombre_proyecto"] == "Proyecto B"
    # Re-registrar mismo id+nombre no duplica.
    registro.registrar_proyecto("Proyecto B", ruta=ruta)
    assert len(registro.listar_proyectos(ruta)) == 1


def test_c7_p12_colision_id_nombre_distinto(tmp_path):
    ruta = tmp_path / "proyectos.json"
    registro.registrar_proyecto("Proyecto A", proyecto_id="proyecto-x", ruta=ruta)
    with pytest.raises(registro.ColisionProyectoError) as err:
        registro.registrar_proyecto("Proyecto B", proyecto_id="proyecto-x", ruta=ruta)
    assert "Proyecto A" in str(err.value)
    assert err.value.existente.get("proyecto_id") == "proyecto-x"
    # Nada se fusionó en silencio: sigue habiendo un solo proyecto, el original.
    proyectos = registro.listar_proyectos(ruta)
    assert len(proyectos) == 1 and proyectos[0]["nombre_proyecto"] == "Proyecto A"


def test_c7_sin_asignar_no_se_registra_como_proyecto(tmp_path):
    ruta = tmp_path / "proyectos.json"
    with pytest.raises(ValueError):
        registro.registrar_proyecto("Sin asignar", proyecto_id="sin-asignar", ruta=ruta)


def test_c7_registro_con_escritura_fallida_no_declara_exito(tmp_path, monkeypatch):
    monkeypatch.setattr(registro, "escribir_json_seguro", lambda *a, **k: False)
    with pytest.raises(registro.ErrorPersistenciaProyectos):
        registro.registrar_proyecto("Proyecto B", ruta=tmp_path / "proyectos.json")


def test_c7_registro_de_ejecucion_anota_trazabilidad(tmp_path):
    ruta = tmp_path / "proyectos.json"
    registro.registrar_proyecto("Proyecto B", ruta=ruta)
    actualizado = registro.registrar_ejecucion("proyecto-b", "ej-42",
                                               responsable="ana", ruta=ruta)
    assert actualizado["ejecuciones"] == 1
    assert actualizado["ultima_ejecucion"]["ejecucion_id"] == "ej-42"
    # «Sin asignar» no alimenta el registro.
    assert registro.registrar_ejecucion("sin-asignar", ruta=ruta) is None


# ---------------------------------------------------------------------------
# Excel de SALIDA: identidad autoritativa, plantilla intacta (P01/P02/P07/P08/P10)
# ---------------------------------------------------------------------------


@pytest.fixture
def plantilla_con_celda_proyecto(tmp_path):
    """Plantilla cuyo único campo de proyecto está IDENTIFICADO por etiqueta."""
    from test_plantilla_empresarial import _plantilla as crear
    ruta = tmp_path / "plantilla.xlsx"
    crear(ruta)
    libro = load_workbook(ruta)
    captura = libro["Captura_pruebas"]
    captura["A1"] = "Proyecto:"
    captura["B1"] = "Proyecto A"   # texto precargado histórico
    libro.save(ruta)
    return ruta


def test_c7_p01_salida_usa_proyecto_declarado_y_plantilla_intacta(
        plantilla_con_celda_proyecto, tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    plantilla = plantilla_con_celda_proyecto
    hash_antes = _hash(plantilla)
    salida = tmp_path / "salida.xlsx"
    generar_desde_plantilla(plantilla, salida, [], identidad_proyecto=IDENTIDAD_B)
    libro = load_workbook(salida)
    # P01: la celda oficial de proyecto dice B en la SALIDA.
    assert libro["Captura_pruebas"]["B1"].value == "Proyecto B"
    # Bloque de identidad verificable en Resumen.
    resumen = "\n".join(str(c.value) for fila in libro["Resumen"].iter_rows()
                        for c in fila if c.value)
    assert "proyecto-b" in resumen and "Proyecto B" in resumen
    assert "ej-001" in resumen  # ejecucion_id de la instantánea
    # La plantilla original NO fue editada.
    assert _hash(plantilla) == hash_antes
    assert load_workbook(plantilla)["Captura_pruebas"]["B1"].value == "Proyecto A"


def test_c7_p02_sin_asignar_muestra_pendiente_sin_identidad_inventada(
        plantilla_con_celda_proyecto, tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    salida = tmp_path / "salida_sin_asignar.xlsx"
    identidad = {"nombre_proyecto": None, "project_id": "sin-asignar",
                 "origen": "sin_asignar_expreso", "sin_asignar": True,
                 "pendiente": True, "ejecucion_id": "ej-002"}
    generar_desde_plantilla(plantilla_con_celda_proyecto, salida, [],
                            identidad_proyecto=identidad)
    libro = load_workbook(salida)
    assert libro["Captura_pruebas"]["B1"].value == "Sin asignar (pendiente)"
    resumen = "\n".join(str(c.value) for fila in libro["Resumen"].iter_rows()
                        for c in fila if c.value)
    assert "sin-asignar" in resumen
    assert "Proyecto A" not in str(libro["Captura_pruebas"]["B1"].value)


def test_c7_p07_identidad_coherente_en_los_tres_idiomas(
        plantilla_con_celda_proyecto, tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    for idioma, esperado in (("es", "Sin asignar (pendiente)"),
                             ("en", "Unassigned (pending)"),
                             ("de", "Nicht zugeordnet (ausstehend)")):
        identidad = {"nombre_proyecto": None, "project_id": "sin-asignar",
                     "sin_asignar": True, "pendiente": True, "origen": "x"}
        salida = tmp_path / f"salida_{idioma}.xlsx"
        generar_desde_plantilla(plantilla_con_celda_proyecto, salida, [],
                                identidad_proyecto=identidad, idioma=idioma)
        assert load_workbook(salida)["Captura_pruebas"]["B1"].value == esperado
    # Con proyecto declarado, el nombre visible es idéntico en los 3 idiomas.
    nombres = set()
    for idioma in ("es", "en", "de"):
        salida = tmp_path / f"salida_b_{idioma}.xlsx"
        generar_desde_plantilla(plantilla_con_celda_proyecto, salida, [],
                                identidad_proyecto=IDENTIDAD_B, idioma=idioma)
        nombres.add(load_workbook(salida)["Captura_pruebas"]["B1"].value)
    assert nombres == {"Proyecto B"}


def test_c7_p08_plantilla_sin_campo_identificable_bloquea_publicacion(tmp_path):
    """V15.9.3 (CRIT-03): la prueba certifica que la publicación fue EVITADA.

    La versión V15.9.2 publicaba la salida oficial con sólo una advertencia;
    ahora el destino oficial NO debe crearse ni reemplazarse: se guarda un
    borrador no publicable en ubicación diferenciada.
    """
    from plantilla_empresarial import generar_desde_plantilla
    from test_plantilla_empresarial import _plantilla as crear
    plantilla = tmp_path / "externa.xlsx"
    crear(plantilla)
    libro = load_workbook(plantilla)
    # Título histórico SIN etiqueta identificable de proyecto.
    libro["Captura_pruebas"]["A1"] = "Proyecto A"
    libro.save(plantilla)
    hash_antes = _hash(plantilla)
    salida = tmp_path / "salida_p08.xlsx"
    resultado = pathlib.Path(
        generar_desde_plantilla(plantilla, salida, [], identidad_proyecto=IDENTIDAD_B))
    # La salida oficial NO se generó; quedó un borrador diferenciado.
    assert not salida.exists()
    assert resultado != salida and "borrador_no_publicable" in resultado.name
    assert resultado.is_file()
    libro_out = load_workbook(resultado)
    # No se suplanta el título antiguo: sigue diciendo A en el borrador.
    assert libro_out["Captura_pruebas"]["A1"].value == "Proyecto A"
    # Advertencia verificable + identidad técnica disponible en el bloque.
    resumen = "\n".join(str(c.value) for fila in libro_out["Resumen"].iter_rows()
                        for c in fila if c.value)
    assert "ADVERTENCIA" in resumen and "bloquea" in resumen
    assert "BORRADOR NO PUBLICABLE" in resumen.upper()
    assert "proyecto-b" in resumen  # la identidad NO se pierde, se expone
    assert _hash(plantilla) == hash_antes


def test_c7_p10_fallo_de_escritura_no_declara_exito(
        plantilla_con_celda_proyecto, tmp_path):
    from plantilla_empresarial import generar_desde_plantilla
    hash_antes = _hash(plantilla_con_celda_proyecto)
    destino_invalido = tmp_path / "no_existe_dir"
    destino_invalido.mkdir()  # directorio: guardar encima debe fallar
    with pytest.raises(Exception):
        generar_desde_plantilla(plantilla_con_celda_proyecto, destino_invalido,
                                [], identidad_proyecto=IDENTIDAD_B)
    assert _hash(plantilla_con_celda_proyecto) == hash_antes


def test_c7_generar_excel_propaga_identidad_desde_datos(
        plantilla_con_celda_proyecto, tmp_path):
    from generar_excel import generar_excel
    from configuracion import cargar_config
    datos = {"resultados": [], "perfil": "empresarial",
             "proyecto_ejecucion": IDENTIDAD_B}
    salida = tmp_path / "maestro.xlsx"
    generar_excel(None, salida, cargar_config(),
                  ruta_plantilla=plantilla_con_celda_proyecto,
                  datos_validacion=datos, idioma="es")
    assert load_workbook(salida)["Captura_pruebas"]["B1"].value == "Proyecto B"


# ---------------------------------------------------------------------------
# Backend: validación de la solicitud, checkpoint y registro (P04/P09/P12)
# ---------------------------------------------------------------------------


@pytest.fixture
def entorno_api(tmp_path, monkeypatch):
    backend = _backend()
    monkeypatch.setattr(registro, "RUTA_PROYECTOS", tmp_path / "proyectos.json")
    lote = tmp_path / "Proyecto_A"
    lote.mkdir()
    estado_original = dict(backend._pipeline_estado)
    llamadas = {}

    def fondo_falso(**kwargs):
        llamadas.update(kwargs)

    monkeypatch.setattr(backend, "_capacidad_pipeline",
                        lambda: {"listo": True, "recursos": {}})
    monkeypatch.setattr(backend, "_ejecutar_pipeline_fondo", fondo_falso)
    monkeypatch.setattr(backend, "_persistir_checkpoint_pipeline", lambda: None)
    yield backend, lote, llamadas
    hilo = backend._pipeline_hilo
    if isinstance(hilo, threading.Thread) and hilo.is_alive():
        hilo.join(timeout=5)
    backend._pipeline_hilo = None
    backend._pipeline_estado.clear()
    backend._pipeline_estado.update(estado_original)


def test_c7_backend_exige_nombre_o_sin_asignar_expreso(entorno_api):
    """P09: sin decisión explícita no se procesa (400), nada se hereda."""
    backend, lote, _ = entorno_api
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as err:
        backend.iniciar_pipeline(backend.SolicitudPipeline(
            ruta=str(lote), proyecto={"nombre_proyecto": "   ", "sin_asignar": False}))
    assert err.value.status_code == 400
    assert "Sin asignar" in str(err.value.detail)


def test_c7_backend_p12_colision_devuelve_409(entorno_api):
    backend, lote, _ = entorno_api
    from fastapi import HTTPException
    registro.registrar_proyecto("Proyecto A", proyecto_id="proyecto-x")
    with pytest.raises(HTTPException) as err:
        backend.iniciar_pipeline(backend.SolicitudPipeline(
            ruta=str(lote),
            proyecto={"nombre_proyecto": "Proyecto B", "proyecto_id": "proyecto-x"}))
    assert err.value.status_code == 409
    assert "Proyecto A" in str(err.value.detail)


def test_c7_backend_acepta_declaracion_y_snapshot(entorno_api):
    backend, lote, llamadas = entorno_api
    respuesta = backend.iniciar_pipeline(backend.SolicitudPipeline(
        ruta=str(lote),
        proyecto={"nombre_proyecto": "Proyecto B", "usuario": "ana"}))
    assert respuesta["proyecto_ejecucion"]["project_id"] == "proyecto-b"
    assert respuesta["proyecto_ejecucion"]["nombre_proyecto"] == "Proyecto B"
    assert respuesta["proyecto_ejecucion"]["ejecucion_id"]
    # El hilo recibió el proyecto declarado (no la carpeta «Proyecto_A»).
    assert llamadas["proyecto"]["nombre_proyecto"] == "Proyecto B"
    # El checkpoint guarda la identidad para reanudaciones.
    guardada = backend._pipeline_estado["solicitud"]
    assert guardada["proyecto_ejecucion"]["project_id"] == "proyecto-b"
    # El registro quedó anotado (trazabilidad de la ejecución).
    proyectos = registro.listar_proyectos()
    assert any(p["proyecto_id"] == "proyecto-b" and p["ejecuciones"] == 1
               for p in proyectos)


def test_c7_backend_sin_asignar_expreso_pasa_con_estado_pendiente(entorno_api):
    backend, lote, llamadas = entorno_api
    respuesta = backend.iniciar_pipeline(backend.SolicitudPipeline(
        ruta=str(lote), proyecto={"sin_asignar": True, "usuario": "ana"}))
    assert respuesta["proyecto_ejecucion"]["sin_asignar"] is True
    assert respuesta["proyecto_ejecucion"]["pendiente"] is True
    assert llamadas["proyecto"]["sin_asignar"] is True
    # «Sin asignar» no crea un proyecto fantasma en el registro.
    assert registro.listar_proyectos() == []


def test_c7_backend_p04_reanudar_con_otro_proyecto_exige_migracion(entorno_api):
    backend, lote, _ = entorno_api
    from fastapi import HTTPException
    backend._pipeline_estado["solicitud"] = {
        "proyecto_ejecucion": {"project_id": "proyecto-b",
                               "nombre_proyecto": "Proyecto B"}}
    with pytest.raises(HTTPException) as err:
        backend.iniciar_pipeline(backend.SolicitudPipeline(
            ruta=str(lote), modo_ejecucion="reanudar",
            proyecto={"nombre_proyecto": "Proyecto C"}))
    assert err.value.status_code == 409
    assert "migración explícita" in str(err.value.detail)
    # Con migración explícita sí se acepta el cambio.
    respuesta = backend.iniciar_pipeline(backend.SolicitudPipeline(
        ruta=str(lote), modo_ejecucion="reanudar",
        proyecto={"nombre_proyecto": "Proyecto C", "migracion_explicita": True}))
    assert respuesta["proyecto_ejecucion"]["project_id"] == "proyecto-c"


def test_c7_backend_sin_proyecto_mantiene_compatibilidad_legacy(entorno_api):
    """Clientes antiguos: identidad legacy por carpeta, visible como pendiente."""
    backend, lote, _ = entorno_api
    respuesta = backend.iniciar_pipeline(backend.SolicitudPipeline(ruta=str(lote)))
    assert respuesta["proyecto_ejecucion"]["project_id"] == "proyecto_a"
    assert respuesta["proyecto_ejecucion"]["origen"] == "nombre_carpeta_legacy"
    assert respuesta["proyecto_ejecucion"]["pendiente"] is True


def test_c7_api_proyectos_expone_registrados_y_sugerencias(entorno_api, monkeypatch):
    backend, lote, _ = entorno_api
    registro.registrar_proyecto("Proyecto B")
    datos = {"raiz": str(lote),
             "resultados": [{"project_id": "proyecto-legacy", "nombre": "x"}]}
    monkeypatch.setattr(backend, "_cargar_datos", lambda: datos)
    respuesta = backend.listar_proyectos()
    ids = {p["proyecto_id"] for p in respuesta["proyectos"]}
    assert "proyecto-b" in ids
    sugeridos = {s["proyecto_id"] for s in respuesta["sugerencias"]}
    # Carpeta actual y resultados anteriores: SUGERENCIAS pendientes, jamás
    # proyectos registrados.
    assert "proyecto_a" in sugeridos and "proyecto-legacy" in sugeridos
    assert all(s["pendiente_aceptacion"] is True for s in respuesta["sugerencias"])


def test_c7_p04_pipeline_reanudar_guarda_proyecto_del_checkpoint(tmp_path):
    """Nivel pipeline: cambiar el proyecto al reanudar exige migración."""
    from pipeline import ejecutar_pipeline
    from inventario import VERSION_INVENTARIO
    lote = tmp_path / "lote"
    lote.mkdir()
    inv = tmp_path / "inventario.json"
    inv.write_text(__import__("json").dumps({
        "version": VERSION_INVENTARIO, "raiz": str(lote), "casos": [],
        "proyecto": {"project_id": "proyecto-b", "nombre_proyecto": "Proyecto B"},
    }), encoding="utf-8")
    with pytest.raises(ValueError) as err:
        ejecutar_pipeline(lote, modo_ejecucion="reanudar", ruta_inventario=inv,
                          proyecto={"nombre_proyecto": "Proyecto C"})
    assert "migración explícita" in str(err.value)
