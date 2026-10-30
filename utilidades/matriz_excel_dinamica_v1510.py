"""Salida OPCIONAL de matriz dinámica en hoja separada; contrato BMG intacto.

Permite usar la vista dinámica sin tocar fórmulas/estilos/tabla de plantilla.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path


def escribir_hoja_dinamica(libro, matriz: dict, *, nombre: str='Matriz dinámica V15.10'):
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    if nombre in libro.sheetnames:
        # No reemplazar hojas anteriores de manera implícita.
        raise ValueError('La hoja ya existe; no sobrescribir resultados')
    ws=libro.create_sheet(nombre)
    ws.append(['V-Matrix • Matriz dinámica', 'Solo casos reales, sin cupos artificiales'])
    ws.merge_cells(start_row=1,start_column=2,end_row=1,end_column=8)
    ws.cell(1,1).font=Font(bold=True,size=15)
    colores={'verde':'DCFCE7','amarillo':'FEF3C7','rojo':'FEE2E2','gris':'E5E7EB'}
    row=3
    for etapa in matriz.get('etapas',[]):
        ws.cell(row,1,str(etapa.get('tipo_st') or 'Sin etapa')).font=Font(bold=True,size=13)
        row+=1
        columnas=etapa.get('columnas',[])
        for j,c in enumerate(columnas,1):
            ws.cell(row,j,f"{c.get('temperatura','')} {c.get('generador','')}").font=Font(bold=True)
        fila_datos=row+1
        max_rows=max((len(c.get('celdas',[])) for c in columnas),default=0)
        for j,col in enumerate(columnas,1):
            for k,item in enumerate(col.get('celdas',[])):
                if item.get('tipo')!='caso':continue
                x=ws.cell(fila_datos+k,j,item.get('test_number') or item.get('id'))
                estado=item.get('estado') or 'gris'
                x.fill=PatternFill('solid',fgColor=colores.get(estado,'E5E7EB'))
                from openpyxl.comments import Comment
                motivo=item.get('motivo_estado') or 'sin_diagnostico'
                tiempo=item.get('opening_time_ms')
                x.comment=Comment(f"Apertura: {tiempo if tiempo is not None else 'sin publicar'} ms\nEstado: {estado}\nMotivo: {motivo}",'V-Matrix')
                x.alignment=Alignment(horizontal='center')
        row=fila_datos+max_rows+2
    ws.cell(row,1,'Sin clasificar').font=Font(bold=True)
    row+=1
    for item in matriz.get('sin_clasificar',[]):
        ws.cell(row,1,item.get('test_number') or item.get('id'))
        ws.cell(row,2,item.get('causa') or item.get('motivo_estado'))
        row+=1
    for j in range(1,9):ws.column_dimensions[get_column_letter(j)].width=21 if j!=2 else 38
    ws.freeze_panes='B4'
    return ws


def main():
    p=argparse.ArgumentParser(description='Generar nueva matriz dinámica sin modificar plantilla original')
    p.add_argument('--casos',required=True,help='JSON con filas publicadas del dashboard')
    p.add_argument('--salida',required=True,help='Nuevo archivo xlsx (no plantilla original)')
    p.add_argument('--plantilla',help='Opcional libro xlsx a copiar (nunca se sobrescribe)')
    a=p.parse_args()
    from openpyxl import load_workbook,Workbook
    from utilidades.matriz_apertura import construir_matriz_dashboard,criterio_desde_limites
    rows=json.loads(Path(a.casos).read_text(encoding='utf-8'))
    if isinstance(rows,dict):rows=rows.get('filas',[])
    if not isinstance(rows,list):p.error('El JSON debe contener lista de filas')
    m=construir_matriz_dashboard(rows,criterio_desde_limites(6,7))
    if m is None:p.error('No hay casos; no se generará matriz ficticia')
    origen=Path(a.plantilla).resolve() if a.plantilla else None
    destino=Path(a.salida).resolve()
    if origen and destino==origen:p.error('No sobrescribir plantilla')
    if destino.exists():p.error('El archivo de salida existe; elige un nombre nuevo')
    wb=load_workbook(origen) if origen else Workbook()
    escribir_hoja_dinamica(wb,m)
    destino.parent.mkdir(parents=True,exist_ok=True)
    wb.save(destino)
    print(f'Matriz dinámica guardada: {destino}')

if __name__=='__main__':main()
