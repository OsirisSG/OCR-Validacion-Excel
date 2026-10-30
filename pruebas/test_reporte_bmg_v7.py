from pathlib import Path

from openpyxl import load_workbook

from configuracion import cargar_config, cargar_reglas
from reporte_bmg import generar_reporte_bmg


def caso(numero, tipo_st, temp, inflador, apertura, manual=False):
    """Caso sintético para el reporte BMG.

    R15 (SUPERSEDE documentado): desde V15.10 el contrato canónico sólo
    publica/colorea un tiempo de apertura con confirmación humana
    (VERIFICADO_HUMANO) o AUTO_VERIFICABLE con referencia temporal validada;
    las propuestas y los campos sin evidencia suficiente quedan GRISES
    (PLAN_CORRECCION_R13 §Semáforo: «gris: caso ya evaluado pero sin tiempo
    defendible, en conflicto o con referencia temporal insuficiente»;
    RESULTADOS_PRUEBAS_V15_9: «la confirmación humana nunca se reemplaza por
    un consenso automático»). ``manual=True`` genera una confirmación humana
    con referencia válida; el valor antiguo de ``estado: confirmado`` sin
    confirmación humana ya no colorea (abstención, no borrado).
    """
    if apertura is None:
        campo_apertura = {"valor": None, "estado": "sin_dato"}
    elif manual:
        campo_apertura = {"valor": apertura, "estado": "confirmado",
                          "confirmacion_manual": True,
                          "referencia_tiempo_valida": True}
    else:
        campo_apertura = {"valor": apertura, "estado": "confirmado"}
    campos = {
        "test_number": {"valor": str(numero), "estado": "confirmado"},
        "temperature_condition": {"valor": temp, "estado": "confirmado"},
        "inflator_type": {"valor": inflador, "estado": "confirmado"},
        "opening_time_ms": campo_apertura,
    }
    return {
        "id": f"id-{numero}", "case_key": f"case-{numero}", "nombre": str(numero),
        "identificador": str(numero), "test_number": str(numero), "tipo_st": tipo_st,
        "temperature_condition": temp, "inflator_type": inflador, "campos": campos,
    }


def test_reporte_bmg_combina_1st_2st_y_colorea_apertura(tmp_path):
    """R15 (SUPERSEDE documentado del contrato V7).

    Dos cambios de doctrina V15.10 (R9-R13) reemplazan lo que esta prueba
    fijaba en V7; el código no cambió en R15, se alinea la prueba:

    1. Resumen DINÁMICO: ``_poblar_matriz`` dibuja los casos existentes
       empezando en la fila 7 (1ST) y en el bloque dinámico 2ST (fila 33),
       ya no en las posiciones fijas de plantilla D16/G42, y la altura
       depende de los casos presentes, no del cupo histórico.
    2. Abstención: un ``opening_time_ms`` con ``estado: confirmado`` pero SIN
       confirmación humana ya no se publica ni se colorea; queda gris
       (PLAN_CORRECCION_R13 §Semáforo). Para ejercer verde/amarillo/rojo la
       prueba usa ``manual=True`` (VERIFICADO_HUMANO con referencia válida).

    Se retiran las aserciones de celdas fijas (D16/E16/F16/G42, K17) que
    describían el resumen estático anterior; se conservan las intenciones
    originales: combinación 1ST/2ST, colores por criterio, sin dato => gris,
    detalle por bloque e hipervínculo resumen->detalle.
    """
    raiz = Path(__file__).resolve().parents[1]
    plantilla = raiz / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    salida = tmp_path / "resultado.xlsx"
    resultados = [
        caso(1001, "1ST", "HT", "UGL", 5.8, manual=True),
        caso(1002, "1ST", "HT", "OGL", 6.5, manual=True),
        caso(1003, "1ST", "HT", "NOM", 7.2, manual=True),
        caso(2001, "2ST", "RT", "NOM", None),
    ]
    config, reglas = cargar_config(), cargar_reglas()
    generar_reporte_bmg(plantilla, salida, resultados, config, reglas)

    libro = load_workbook(salida, data_only=False)
    resumen = libro["Übersicht Standversuche"]
    lista = libro["Versuchsmatrix"]

    # Resumen dinámico: 1ST arranca en la fila 7, 2ST en su bloque propio.
    assert resumen["D7"].value == "1001"
    assert resumen["E7"].value == "1002"
    assert resumen["F7"].value == "1003"
    bloque_2st = [celda for fila in resumen.iter_rows(min_row=20, max_row=45, min_col=2, max_col=10)
                  for celda in fila if str(celda.value) == "2001"]
    assert len(bloque_2st) == 1, "el caso 2ST debe dibujarse una sola vez"

    def rgb(celda):
        return (celda.fill.fgColor.rgb or "")[-6:].upper()

    assert rgb(resumen["D7"]) == config["colores"]["verde"].lstrip("#").upper()
    assert rgb(resumen["E7"]) == config["colores"]["amarillo"].lstrip("#").upper()
    assert rgb(resumen["F7"]) == config["colores"]["rojo"].lstrip("#").upper()
    assert rgb(bloque_2st[0]) == config["colores"]["gris"].lstrip("#").upper()

    # El detalle 1ST ocupa su bloque y 2ST inicia en el bloque siguiente.
    assert lista["C8"].value == "1003"
    assert lista["C9"].value == "1002"
    assert lista["C10"].value == "1001"
    assert lista["C58"].value == "2001"
    assert resumen["D7"].hyperlink is not None
    assert "Versuchsmatrix" in resumen["D7"].hyperlink.target


def test_tiempo_en_conflicto_se_muestra_gris(tmp_path):
    """R15 (SUPERSEDE documentado): el conflicto se dibuja en el resumen
    dinámico (fila 7), no en la posición fija D16 del resumen estátrico V7.
    La intención se conserva: tiempo marcado ``requiere_revision`` => gris."""
    raiz = Path(__file__).resolve().parents[1]
    plantilla = raiz / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    salida = tmp_path / "conflicto.xlsx"
    item = caso(3001, "1ST", "HT", "UGL", 5.5, manual=True)
    item["campos"]["opening_time_ms"]["requiere_revision"] = True
    generar_reporte_bmg(plantilla, salida, [item], cargar_config(), cargar_reglas())
    libro = load_workbook(salida)
    celda = libro["Übersicht Standversuche"]["D7"]
    assert celda.value == "3001"
    assert (celda.fill.fgColor.rgb or "")[-6:].upper() == cargar_config()["colores"]["gris"].lstrip("#").upper()
