"""Adaptador de la plantilla de captura; las claves estables son el contrato."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import uuid
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


HOJA_CAPTURA = "Captura_pruebas"
HOJA_DICCIONARIO = "Diccionario_campos"
HOJA_CATALOGOS = "Catalogos"
HOJA_TRAZABILIDAD = "Trazabilidad_OCR"
FILA_CLAVES = 5
FILA_DATOS = 6

COLUMNAS_TRAZABILIDAD = [
    "test_number", "tipo_st", "temperature_condition", "module_version",
    "inflator_type", "fase", "tor", "ruta_absoluta", "ruta_relativa",
    "estado_imagen", "texto_crudo", "confianza_ocr", "coordenadas_roi",
    "zoom_aplicado", "modo_recorte", "dispositivo", "tiempo_deteccion",
    "tiempo_ocr", "campos_detectados", "requiere_revision", "mensaje_error",
    "qr_detectado", "payload_qr", "poligono_qr", "fuente_campo",
    "modelo_ocr", "correccion_confirmada", "dataset_entrenamiento", "estado_cache",
]


def _coincide_estrategia(estrategia: dict, estructura: dict, ruta_raiz: Path) -> bool:
    """Evalúa marcadores declarativos para añadir estilos sin cambiar Python."""
    perfil = "empresarial" if estructura.get("casos_empresariales") else "legacy"
    perfiles = {str(valor).lower() for valor in estrategia.get("perfiles", [])}
    if perfiles and perfil not in perfiles:
        return False
    tipos = {str(valor).upper() for valor in estrategia.get("tipos_st", [])}
    detectados = {str(valor).upper() for valor in estructura.get("tipos_st", [])}
    if estructura.get("tipo_st"):
        detectados.add(str(estructura["tipo_st"]).upper())
    if tipos and not tipos.intersection(detectados):
        return False
    texto_ruta = "/".join(ruta_raiz.parts).upper()
    marcadores = [str(valor).upper() for valor in estrategia.get("marcadores_ruta", [])]
    return not marcadores or any(marcador in texto_ruta for marcador in marcadores)


def _raices_plantillas(ruta_raiz: Path, directorios: Iterable[str]) -> list[Path]:
    raiz_codigo = Path(__file__).resolve().parent
    candidatas = [ruta_raiz, ruta_raiz / "plantillas", ruta_raiz.parent,
                  raiz_codigo / "plantillas"]
    for declarada in directorios:
        ruta = Path(str(declarada)).expanduser()
        candidatas.append((raiz_codigo / ruta).resolve() if not ruta.is_absolute()
                          else ruta.resolve())
    salida = []
    for ruta in candidatas:
        ruta = ruta.resolve()
        if ruta.is_dir() and ruta not in salida:
            salida.append(ruta)
    return salida


def detectar_plantilla_automatica(ruta_raiz: str | Path, estructura: dict,
                                  config: dict, explicita: str | Path | None = None,
                                  plantilla_id: str | None = None) -> dict:
    """Selecciona una plantilla compatible o el generador integrado.

    Las estrategias viven en ``fase3.deteccion_plantillas.estrategias``. Un
    estilo futuro sólo necesita declarar perfil/tipo, marcadores de ruta y
    patrones de archivo; cada XLSX se valida antes de ser seleccionado.
    """
    raiz = Path(ruta_raiz).expanduser().resolve()
    cfg = config.get("fase3", {}).get("deteccion_plantillas", {})
    registro = RegistroPlantillas(config)
    if plantilla_id:
        registrada = registro.obtener(plantilla_id)
        if not registrada:
            raise ValueError("La plantilla registrada seleccionada ya no existe.")
        cargar_contrato(registrada["ruta"])
        return {**registrada, "estilo": "plantilla_registrada",
                "fuente": "seleccionada", "estrategia": "registro_local"}
    configurada = config.get("fase3", {}).get("plantilla_empresarial")
    for ruta, fuente in ((explicita, "explícita"), (configurada, "configuración")):
        if not ruta:
            continue
        candidata = Path(str(ruta).strip().strip('"\'')).expanduser()
        candidata = ((Path(__file__).resolve().parent / candidata).resolve()
                     if not candidata.is_absolute() else candidata.resolve())
        if not candidata.is_file():
            if fuente == "explícita":
                raise ValueError("La plantilla indicada no existe o no es accesible.")
            continue
        cargar_contrato(candidata)
        return {"ruta": str(candidata), "estilo": "plantilla_empresarial",
                "fuente": fuente, "estrategia": "seleccion_directa"}

    if not cfg.get("activar", True) or not estructura.get("casos_empresariales"):
        return {"ruta": None, "estilo": "generador_estandar", "fuente": "integrada",
                "estrategia": "sin_plantilla_externa"}

    perfil = "empresarial" if estructura.get("casos_empresariales") else "legacy"
    tipos_detectados = {str(valor).upper() for valor in estructura.get("tipos_st", [])}
    if estructura.get("tipo_st"):
        tipos_detectados.add(str(estructura["tipo_st"]).upper())
    texto_ruta = "/".join(raiz.parts).upper()
    compatibles_registro = []
    for registrada in registro.listar():
        tipos_registrados = {str(tipo).upper() for tipo in registrada.get("tipos_st", [])}
        perfiles = {str(valor).lower() for valor in registrada.get("perfiles", [])}
        marcadores = [str(valor).upper() for valor in registrada.get("marcadores_ruta", [])]
        if perfiles and perfil not in perfiles:
            continue
        if tipos_registrados and not tipos_registrados.intersection(tipos_detectados):
            continue
        if marcadores and not any(marcador in texto_ruta for marcador in marcadores):
            continue
        compatibles_registro.append(registrada)
    # Un marcador de ruta es una selección inequívoca. Sin marcadores se usa
    # automáticamente sólo cuando queda un contrato compatible; ante empate no
    # se adivina ni se mezclan contratos.
    con_marcador = [item for item in compatibles_registro if item.get("marcadores_ruta")]
    candidatas_registro = con_marcador or compatibles_registro
    if len(candidatas_registro) == 1:
        registrada = candidatas_registro[0]
        cargar_contrato(registrada["ruta"])
        return {**registrada, "estilo": "plantilla_registrada",
                "fuente": "registro", "estrategia": "estructura_registrada"}

    estrategias = cfg.get("estrategias") or [{
        "nombre": "empresarial_1st_2st", "perfiles": ["empresarial"],
        "tipos_st": ["1ST", "2ST"],
        "patrones_archivo": ["*plantilla*.xlsx", "*captura*.xlsx", "*template*.xlsx"],
    }]
    raices = _raices_plantillas(raiz, cfg.get("directorios", []))
    rechazadas = []
    for estrategia in estrategias:
        if not _coincide_estrategia(estrategia, estructura, raiz):
            continue
        patrones = estrategia.get("patrones_archivo", ["*.xlsx"])
        vistas = set()
        for carpeta in raices:
            for patron in patrones:
                for candidata in sorted(carpeta.glob(str(patron))):
                    resuelta = candidata.resolve()
                    if resuelta in vistas or not resuelta.is_file():
                        continue
                    vistas.add(resuelta)
                    try:
                        cargar_contrato(resuelta)
                    except (OSError, ValueError, KeyError) as exc:
                        rechazadas.append({"ruta": str(resuelta), "motivo": str(exc)})
                        continue
                    return {"ruta": str(resuelta), "estilo": "plantilla_empresarial",
                            "fuente": "estructura", "estrategia": estrategia.get(
                                "nombre", "configurada"), "rechazadas": rechazadas}
    return {"ruta": None, "estilo": "generador_estandar", "fuente": "integrada",
            "estrategia": "estructura_sin_xlsx_compatible", "rechazadas": rechazadas}


def cargar_contrato(ruta_plantilla: str | Path) -> dict:
    libro = load_workbook(ruta_plantilla, data_only=False)
    faltantes = {HOJA_CAPTURA, HOJA_DICCIONARIO, HOJA_CATALOGOS, "Resumen"} - set(libro.sheetnames)
    if faltantes:
        raise ValueError(f"La plantilla no contiene las hojas requeridas: {sorted(faltantes)}")
    captura = libro[HOJA_CAPTURA]
    claves = {}
    for celda in captura[FILA_CLAVES]:
        if celda.value:
            clave = str(celda.value).strip()
            if clave in claves:
                raise ValueError(f"Clave estable duplicada en la plantilla: {clave}")
            claves[clave] = celda.column

    diccionario = libro[HOJA_DICCIONARIO]
    encabezados = {str(c.value).strip(): c.column for c in diccionario[4] if c.value}
    esquema = {}
    for fila in range(5, diccionario.max_row + 1):
        registro = {nombre: diccionario.cell(fila, columna).value
                    for nombre, columna in encabezados.items()}
        clave = str(registro.get("clave_estable") or "").strip()
        if clave:
            esquema[clave] = registro

    catalogos = {}
    hoja_cat = libro[HOJA_CATALOGOS]
    for columna in range(1, hoja_cat.max_column + 1):
        nombre = hoja_cat.cell(1, columna).value
        if isinstance(nombre, str) and nombre.startswith("CAT_"):
            valores = []
            for fila in range(3, hoja_cat.max_row + 1):
                valor = hoja_cat.cell(fila, columna).value
                if valor not in (None, ""):
                    valores.append(str(valor))
            catalogos[nombre] = valores
    return {"claves": claves, "esquema": esquema, "catalogos": catalogos}


class RegistroPlantillas:
    """Biblioteca local de contratos Excel independientes y reutilizables."""

    def __init__(self, config: dict | None = None, directorio: str | Path | None = None):
        cfg = (config or {}).get("fase3", {}).get("registro_plantillas", {})
        base = Path(directorio or cfg.get("directorio", ".plantillas"))
        if not base.is_absolute():
            base = Path(__file__).resolve().parent / base
        self.directorio = base.resolve()
        self.archivos = self.directorio / "archivos"
        self.indice = self.directorio / "registro.json"

    def _leer(self) -> dict:
        if not self.indice.is_file():
            return {"version": 1, "plantillas": []}
        try:
            datos = json.loads(self.indice.read_text(encoding="utf-8"))
            return datos if isinstance(datos.get("plantillas"), list) else {
                "version": 1, "plantillas": []}
        except (OSError, json.JSONDecodeError, AttributeError):
            return {"version": 1, "plantillas": []}

    def _guardar(self, datos: dict) -> None:
        self.directorio.mkdir(parents=True, exist_ok=True)
        temporal = self.indice.with_suffix(".tmp")
        temporal.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
        temporal.replace(self.indice)

    def listar(self) -> list[dict]:
        salida = []
        for item in self._leer()["plantillas"]:
            actual = dict(item)
            actual["disponible"] = Path(actual.get("ruta", "")).is_file()
            salida.append(actual)
        return salida

    def obtener(self, plantilla_id: str) -> dict | None:
        return next((item for item in self.listar()
                     if item.get("id") == plantilla_id and item.get("disponible")), None)

    def registrar(self, ruta: str | Path, nombre: str | None = None,
                  tipos_st: Iterable[str] | None = None,
                  marcadores_ruta: Iterable[str] | None = None) -> dict:
        origen = Path(str(ruta).strip().strip('"\'')).expanduser().resolve()
        if not origen.is_file() or origen.suffix.lower() != ".xlsx":
            raise ValueError("La plantilla debe ser un archivo .xlsx accesible.")
        contrato = cargar_contrato(origen)
        digest = hashlib.sha256(origen.read_bytes()).hexdigest()
        datos = self._leer()
        existente = next((p for p in datos["plantillas"] if p.get("sha256") == digest), None)
        if existente:
            if nombre and str(nombre).strip():
                existente["nombre"] = str(nombre).strip()
            existente["tipos_st"] = sorted({
                *existente.get("tipos_st", []),
                *(str(tipo).upper() for tipo in (tipos_st or [])
                  if str(tipo).upper() in {"1ST", "2ST"}),
            })
            existente["marcadores_ruta"] = sorted({
                *existente.get("marcadores_ruta", []),
                *(str(valor).strip() for valor in (marcadores_ruta or [])
                  if str(valor).strip()),
            })
            self._guardar(datos)
            return {**existente, "disponible": Path(existente["ruta"]).is_file(),
                    "duplicada": True}
        plantilla_id = f"tpl-{digest[:12]}-{uuid.uuid4().hex[:4]}"
        self.archivos.mkdir(parents=True, exist_ok=True)
        destino = self.archivos / f"{plantilla_id}.xlsx"
        shutil.copy2(origen, destino)
        registro = {
            "id": plantilla_id, "nombre": str(nombre or origen.stem).strip() or origen.stem,
            "ruta": str(destino), "ruta_origen": str(origen), "sha256": digest,
            "tipos_st": sorted({str(tipo).upper() for tipo in (tipos_st or [])
                                 if str(tipo).upper() in {"1ST", "2ST"}}),
            "perfiles": ["empresarial"],
            "marcadores_ruta": sorted({str(valor).strip() for valor in
                                         (marcadores_ruta or []) if str(valor).strip()}),
            "hojas": [HOJA_CAPTURA, HOJA_DICCIONARIO, HOJA_CATALOGOS, "Resumen"],
            "claves": len(contrato["claves"]),
            "claves_estables": sorted(contrato["claves"]),
            "catalogos": sorted(contrato["catalogos"]),
            "registrada_en": datetime.now().isoformat(timespec="seconds"),
        }
        datos["plantillas"].append(registro)
        self._guardar(datos)
        return {**registro, "disponible": True, "duplicada": False}


def _limites_y_regex(valor) -> tuple[float | None, str | None]:
    if valor in (None, ""):
        return None, None
    texto = str(valor).strip()
    # La celda combina máximo y regex con `` espacio|espacio ``. No dividir
    # alternancias legítimas como ``^(NT|RT|HT)$``.
    partes = [p.strip() for p in re.split(r"\s+\|\s+", texto, maxsplit=1)]
    maximo = None
    patron = None
    try:
        maximo = float(partes[0])
    except ValueError:
        patron = partes[0] if partes[0].startswith("^") else None
    if len(partes) == 2:
        patron = partes[1]
    return maximo, patron


def validar_valor(clave: str, valor, contrato: dict) -> dict:
    regla = contrato["esquema"].get(clave, {})
    if valor in (None, ""):
        return {"valido": str(regla.get("requerido") or "").lower() not in {"sí", "si"},
                "valor": None, "errores": ["campo requerido"] if
                str(regla.get("requerido") or "").lower() in {"sí", "si"} else []}
    tipo = str(regla.get("tipo_esperado") or "string").lower()
    errores = []
    convertido = valor
    try:
        if tipo == "integer":
            convertido = int(float(valor))
        elif tipo == "decimal":
            convertido = float(str(valor).replace(",", "."))
        elif tipo == "date" and isinstance(valor, str):
            convertido = datetime.strptime(valor, "%Y-%m-%d").date()
        elif tipo in {"identifier", "string", "enum"}:
            convertido = str(valor)
    except (TypeError, ValueError):
        errores.append(f"tipo inválido: se esperaba {tipo}")
    catalogo = str(regla.get("valores_o_catalogo") or "")
    if catalogo.startswith("CAT_") and str(convertido) not in contrato["catalogos"].get(catalogo, []):
        errores.append(f"fuera de catálogo {catalogo}")
    minimo = regla.get("minimo")
    maximo, patron = _limites_y_regex(regla.get("maximo_regex"))
    if isinstance(convertido, (int, float)):
        if minimo not in (None, "") and convertido < float(minimo):
            errores.append(f"menor que {minimo}")
        if maximo is not None and convertido > maximo:
            errores.append(f"mayor que {maximo:g}")
    if patron and not re.fullmatch(patron, str(valor)):
        errores.append("no cumple la expresión regular")
    return {"valido": not errores, "valor": convertido, "errores": errores}


def _copiar_fila_modelo(hoja, origen: int, destino: int) -> None:
    hoja.row_dimensions[destino].height = hoja.row_dimensions[origen].height
    for columna in range(1, hoja.max_column + 1):
        src, dst = hoja.cell(origen, columna), hoja.cell(destino, columna)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)


def _extender_validaciones(hoja, ultima_fila: int) -> None:
    for validacion in hoja.data_validations.dataValidation:
        nuevos = []
        for rango in validacion.ranges.ranges:
            if rango.min_row <= FILA_DATOS and rango.max_row >= FILA_DATOS:
                nuevos.append(f"{get_column_letter(rango.min_col)}{rango.min_row}:"
                              f"{get_column_letter(rango.max_col)}{max(rango.max_row, ultima_fila)}")
            else:
                nuevos.append(str(rango))
        validacion.sqref = " ".join(nuevos)


def _actualizar_resumen(libro, ultima_fila: int) -> None:
    hoja = libro["Resumen"]
    for fila in hoja.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.startswith("="):
                celda.value = re.sub(
                    r"(Captura_pruebas!\$[A-Z]+\$6:\$[A-Z]+\$)\d+\b",
                    rf"\g<1>{ultima_fila}", celda.value)


def generar_desde_plantilla(ruta_plantilla: str | Path, ruta_salida: str | Path,
                            resultados: list[dict]) -> Path:
    contrato = cargar_contrato(ruta_plantilla)
    libro = load_workbook(ruta_plantilla)
    hoja = libro[HOJA_CAPTURA]
    ultima = max(FILA_DATOS, FILA_DATOS + len(resultados) - 1)
    if ultima > hoja.max_row:
        for fila in range(hoja.max_row + 1, ultima + 1):
            _copiar_fila_modelo(hoja, FILA_DATOS, fila)
    # La salida corresponde a esta ejecución; se conservan estilos, no datos anteriores.
    for fila in range(FILA_DATOS, max(hoja.max_row, ultima) + 1):
        for columna in contrato["claves"].values():
            hoja.cell(fila, columna).value = None

    for numero, resultado in enumerate(resultados, start=FILA_DATOS):
        campos = resultado.get("campos") or resultado.get("consolidado", {}).get("campos", {})
        for clave, columna in contrato["claves"].items():
            valor = campos.get(clave, {})
            valor = valor.get("valor") if isinstance(valor, dict) else valor
            validacion = validar_valor(clave, valor, contrato)
            # Se conserva vacío cuando un candidato viola el contrato; la evidencia queda en trazabilidad.
            celda = hoja.cell(numero, columna)
            celda.value = validacion["valor"] if validacion["valido"] else None
            tipo = str(contrato["esquema"].get(clave, {}).get("tipo_esperado") or "")
            if tipo in {"identifier", "string", "enum"}:
                celda.number_format = "@"
            elif tipo == "date":
                celda.number_format = "yyyy-mm-dd"

    _extender_validaciones(hoja, ultima)
    ultima_columna = get_column_letter(max(contrato["claves"].values()))
    for tabla in hoja.tables.values():
        inicio = tabla.ref.split(":")[0]
        tabla.ref = f"{inicio}:{ultima_columna}{ultima}"
    _actualizar_resumen(libro, ultima)

    if HOJA_TRAZABILIDAD in libro.sheetnames:
        del libro[HOJA_TRAZABILIDAD]
    traza = libro.create_sheet(HOJA_TRAZABILIDAD)
    fill = PatternFill("solid", fgColor="1F4E78")
    for columna, nombre in enumerate(COLUMNAS_TRAZABILIDAD, start=1):
        celda = traza.cell(1, columna, nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = fill
        traza.column_dimensions[get_column_letter(columna)].width = (
            46 if "ruta" in nombre else 50 if nombre == "texto_crudo" else 20)
    for resultado in resultados:
        for evidencia in resultado.get("trazabilidad", []):
            traza.append([
                ", ".join(evidencia.get(nombre, [])) if isinstance(evidencia.get(nombre), list)
                else str(evidencia.get(nombre)) if isinstance(evidencia.get(nombre), tuple)
                else evidencia.get(nombre)
                for nombre in COLUMNAS_TRAZABILIDAD
            ])
    traza.freeze_panes = "A2"
    ultima_traza = get_column_letter(len(COLUMNAS_TRAZABILIDAD))
    traza.auto_filter.ref = f"A1:{ultima_traza}{max(traza.max_row, 1)}"
    if traza.max_row >= 2:
        tabla = Table(displayName="TablaTrazabilidadOCR", ref=f"A1:{ultima_traza}{traza.max_row}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        traza.add_table(tabla)
    for fila in traza.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)

    # Excel/LibreOffice deben recalcular el resumen al abrir la salida.
    libro.calculation.fullCalcOnLoad = True
    libro.calculation.forceFullCalc = True
    libro.calculation.calcMode = "auto"

    ruta_salida = Path(ruta_salida)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta_salida)
    return ruta_salida
