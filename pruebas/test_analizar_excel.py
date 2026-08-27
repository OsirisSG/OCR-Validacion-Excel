"""Pruebas del inventario y perfilado semántico de libros Excel."""

from datetime import date

import pytest
from openpyxl import Workbook

from analizar_excel import analizar_libro, normalizar_nombre


def _crear_libro(ruta):
    wb = Workbook()
    ventas = wb.active
    ventas.title = "Ventas 2026"
    ventas.merge_cells("A1:E1")
    ventas["A1"] = "Reporte mensual"
    ventas.append([])
    ventas.append(["ID Cliente", "Estado", "Fecha alta", "Monto", "Total calculado"])
    ventas.append(["C-001", "Aprobado", date(2026, 1, 4), 100.0, "=D4*1.16"])
    ventas.append(["C-002", "Pendiente", date(2026, 1, 5), 250.5, "=D5*1.16"])
    ventas.append(["C-003", "Aprobado", date(2026, 1, 6), 80, "=D6*1.16"])
    ventas.append(["C-004", "Rechazado", None, 40, "=D7*1.16"])

    catalogo = wb.create_sheet("Catálogo")
    catalogo.append(["Código", "Categoría", "Categoría", "Activo"])
    catalogo.append(["A-1", "Herramienta", "Manual", "Sí"])
    catalogo.append(["A-2", "Equipo", "Automático", "No"])

    oculta = wb.create_sheet("Auxiliar")
    oculta.sheet_state = "hidden"
    wb.save(ruta)


def _hoja(resultado, nombre):
    return next(h for h in resultado["hojas"] if h["nombre"] == nombre)


def _columna(hoja, nombre_normalizado):
    return next(c for c in hoja["columnas"] if c["nombre_normalizado"] == nombre_normalizado)


def test_analiza_todas_las_hojas_columnas_categorias_y_formulas(tmp_path):
    ruta = tmp_path / "libro.xlsx"
    _crear_libro(ruta)

    resultado = analizar_libro(ruta, max_categorias=10)

    assert resultado["total_hojas"] == 3
    assert resultado["hojas_ocultas"] == ["Auxiliar"]
    assert resultado["total_formulas"] == 4

    ventas = _hoja(resultado, "Ventas 2026")
    assert ventas["fila_encabezado"] == 3
    assert ventas["columnas_con_datos"] == 5
    assert ventas["filas_de_datos_no_vacias"] == 4

    estado = _columna(ventas, "estado")
    assert estado["es_posible_categoria"] is True
    assert set(estado["categorias_posibles"]) == {"Aprobado", "Pendiente", "Rechazado"}
    assert estado["posibles_significados"][0]["categoria"] == "estado_resultado"

    monto = _columna(ventas, "monto")
    assert monto["tipo_dominante"] in {"entero", "decimal"}
    assert monto["resumen_numerico"]["maximo"] == 250.5
    assert any(p["categoria"] == "importe_moneda" for p in monto["posibles_significados"])

    total = _columna(ventas, "total_calculado")
    assert total["formulas"] == 4
    assert total["tipo_dominante"] == "formula"

    catalogo = _hoja(resultado, "Catálogo")
    assert [c["nombre_normalizado"] for c in catalogo["columnas"]] == [
        "codigo", "categoria", "categoria_2", "activo"
    ]
    assert _hoja(resultado, "Auxiliar")["vacia"] is True


def test_rechaza_formatos_heredados_y_normaliza_encabezados(tmp_path):
    ruta_xls = tmp_path / "legado.xls"
    ruta_xls.write_bytes(b"no es un xls real")
    with pytest.raises(ValueError, match="convierte primero"):
        analizar_libro(ruta_xls)

    assert normalizar_nombre("Confianza OCR (%)") == "confianza_ocr"
    assert normalizar_nombre("Número de prueba") == "numero_de_prueba"
