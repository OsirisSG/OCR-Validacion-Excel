"""Trazabilidad H4 en matriz dashboard/Excel para tiempos finales no publicables."""
from pathlib import Path

from openpyxl import load_workbook

from configuracion import cargar_config, cargar_reglas
from reporte_bmg import generar_reporte_bmg


def test_excel_conserva_formula_y_muestra_rechazo_raw_y_fuente_en_matriz(tmp_path):
    raiz = Path(__file__).resolve().parents[1]
    plantilla = raiz / "plantillas" / "reporte_bmg_1st_2st.xlsx"
    salida = tmp_path / "apertura_invalida.xlsx"
    resultado = {
        "id": "invalid-42", "case_key": "proyecto|1ST|HT|UGL|42",
        "project_id": "proyecto", "nombre": "Proyecto de prueba",
        "tipo_st": "1ST", "test_number": "42", "temperature_condition": "HT",
        "inflator_type": "UGL",
        "campos": {
            "test_number": {"valor": "42", "estado": "confirmado"},
            "temperature_condition": {"valor": "HT", "estado": "confirmado"},
            "inflator_type": {"valor": "UGL", "estado": "confirmado"},
            "opening_time_ms": {
                "valor": -122.8, "raw_time_ms": -122.8,
                "estado": "confirmado_manual", "fuente_preferida": "revision_manual",
                "confirmacion_manual": True,
            },
        },
    }
    identidad = {
        "nombre_proyecto": "Proyecto de prueba", "project_id": "proyecto",
        "origen": "declarado_usuario", "sin_asignar": False,
        "pendiente": False, "ejecucion_id": "prueba-h4",
    }
    generado = generar_reporte_bmg(plantilla, salida, [resultado], cargar_config(), cargar_reglas(),
                                   identidad_proyecto=identidad)

    libro = load_workbook(generado, data_only=False)
    resumen = libro["Übersicht Standversuche"]
    assert resumen["D7"].value == "42"
    assert resumen["D7"].comment is not None
    comentario_matriz = resumen["D7"].comment.text
    assert "NO_POSITIVO" in comentario_matriz
    assert "-122.8" in comentario_matriz
    assert "revision_manual" in comentario_matriz
    assert resumen["K17"].value == "=SUM(D17:J17)"

    apertura = libro["Apertura V15.10"]
    encabezados = {celda.value: celda.column for celda in apertura[1]}
    assert apertura.cell(2, encabezados["Decisión V15.10"]).value == "PROPUESTA"
    assert apertura.cell(2, encabezados["Apertura confirmada ms"]).value is None
    assert apertura.cell(2, encabezados["Tiempo propuesto ms"]).value is None
    assert apertura.cell(2, encabezados["Estado tiempo"]).value == "NO_POSITIVO"
    assert apertura.cell(2, encabezados["Tiempo crudo ms"]).value == -122.8
    assert "estrictamente mayor que cero" in apertura.cell(2, encabezados["Motivo rechazo"]).value
    assert apertura.cell(2, encabezados["Ir al detalle"]).hyperlink is not None
    assert apertura.cell(2, encabezados["Decisión V15.10"]).comment is not None
    assert "Tiempo crudo: -122.8 ms" in apertura.cell(2, encabezados["Decisión V15.10"]).comment.text
